#!/usr/bin/env python3
from pathlib import Path
import re, sys
from docx import Document

if len(sys.argv) != 4:
    raise SystemExit('uso: compilar_puia_org.py fonte.org referencia.docx saida.docx')
org_path, ref_path, out_path = map(Path, sys.argv[1:4])
text = org_path.read_text(encoding='utf-8')
lines = text.splitlines()

# First two bold lines
bold_lines = [re.sub(r'^\*|\*$', '', l.strip()).strip() for l in lines if l.strip().startswith('*') and l.strip().endswith('*') and not re.match(r'^\*[0-9]', l.strip())]
if len(bold_lines) < 2:
    raise SystemExit('ERRO: título/tema não encontrados no Org')
title, theme = bold_lines[0], bold_lines[1]

# Parse org tables
def parse_tables(lines):
    tables=[]; cur=[]
    for l in lines:
        if l.strip().startswith('|'):
            if '+' in l and set(l.replace('|','').replace('+','').replace('-','').strip()) == set():
                continue
            cells=[c.strip() for c in l.strip().strip('|').split('|')]
            cur.append(cells)
        elif cur:
            tables.append(cur); cur=[]
    if cur: tables.append(cur)
    return tables

tables=parse_tables(lines)
if len(tables)<2:
    raise SystemExit('ERRO: duas tabelas esperadas no Org')
field_table, resp_table=tables[0], tables[1]

# Parse sections *1..*5
sections={}
for n in range(1,6):
    pat=re.compile(rf'^\*{n}\.\s*(.+?)\*$')
    idx=None; heading=None
    for i,l in enumerate(lines):
        m=pat.match(l.strip())
        if m: idx=i; heading=f'{n}. {m.group(1)}'; break
    if idx is None: raise SystemExit(f'ERRO: seção {n} ausente')
    end=len(lines)
    for j in range(idx+1,len(lines)):
        if re.match(r'^\*[1-5]\.\s*.+\*$', lines[j].strip()): end=j; break
    body=[l.strip() for l in lines[idx+1:end] if l.strip() and not l.strip().startswith('|') and not ('+' in l and l.strip().startswith('|'))]
    sections[n]=(heading,body)

# clean bullet markers
sec2_bullets=[re.sub(r'^-\s*','',x) for x in sections[2][1] if x.startswith('-')]
sec2_bullets=[x for x in sec2_bullets if x]
if len(sec2_bullets)!=5:
    raise SystemExit(f'ERRO: esperados 5 bullets na seção 2; encontrados {len(sec2_bullets)}')

def first_nonbullet(body):
    return ' '.join(x for x in body if not x.startswith('-'))

sec1=first_nonbullet(sections[1][1])
sec3=first_nonbullet(sections[3][1])
sec5_lines=[x for x in sections[5][1] if not x.startswith('-')]
# final signature lines are kept separate by two spaces in Org but splitlines already separate
sec5_body=[]; signature=''; date=''
for x in sec5_lines:
    if x.startswith('Assinatura:'): signature=x
    elif x.startswith('Data:'): date=x
    else: sec5_body.append(x)
if len(sec5_body)<2:
    raise SystemExit('ERRO: seção 5 deve conter corpo + declaração final')
sec5_main=sec5_body[0]
sec5_decl=' '.join(sec5_body[1:])

doc=Document(ref_path)
# preserve table 0 (institutional header), fill metadata table 1
p=doc.paragraphs
p[0].text=title
p[1].text=theme
p[2].text=sections[1][0]
p[3].text=sec1
p[4].text=sections[2][0]
for i,b in enumerate(sec2_bullets, start=5):
    p[i].text='• '+b
p[10].text=sections[3][0]
p[11].text=sec3
p[12].text=sections[4][0]
p[13].text=sections[5][0]
p[14].text=sec5_main
p[15].text=sec5_decl
p[16].text=signature
p[17].text=date

# Fill tables from org while preserving reference table geometry/styles.
def fill_table(tbl, data):
    if len(tbl.rows)!=len(data):
        raise SystemExit(f'ERRO: número de linhas divergente: docx {len(tbl.rows)} x org {len(data)}')
    for r,row in enumerate(data):
        if len(row)!=len(tbl.columns):
            raise SystemExit('ERRO: número de colunas divergente')
        for c,val in enumerate(row):
            tbl.cell(r,c).text=val

fill_table(doc.tables[1], field_table)
fill_table(doc.tables[2], resp_table)

# Restore compact font sizes after replacing text.
from docx.shared import Pt
for i,par in enumerate(doc.paragraphs):
    for run in par.runs:
        if i==0: run.font.size=Pt(12); run.bold=True
        elif i==1: run.font.size=Pt(9); run.bold=True
        elif i in [2,4,10,12,13]: run.font.size=Pt(9); run.bold=True
        else: run.font.size=Pt(7.8)
for ti,tbl in enumerate(doc.tables):
    for ri,row in enumerate(tbl.rows):
        for cell in row.cells:
            for par in cell.paragraphs:
                for run in par.runs:
                    run.font.size=Pt(6.9 if ti>0 else 7.2)
                    if ri==0: run.bold=True

doc.save(out_path)
print(f'PASS: DOCX gerado a partir do Org: {out_path}')
