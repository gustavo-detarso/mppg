import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from academic_pipeline import minuta_runtime as mod


class TestMinutaDocxHarmonizacaoFGV(unittest.TestCase):

    def test_01_corpo_normal(self):
        doc = Document()
        p = doc.add_paragraph("Primeiro paragrafo", style="Normal")

        mod._mppg_docx_harmonize_text(doc)

        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.assertAlmostEqual(
            p.paragraph_format.first_line_indent.cm, 1.25, places=2
        )
        self.assertEqual(p.paragraph_format.line_spacing.pt, 18)
        self.assertEqual(p.paragraph_format.space_after.pt, 0)

    def test_02_corpo_body_text(self):
        doc = Document()
        p = doc.add_paragraph("Segundo paragrafo", style="Body Text")

        mod._mppg_docx_harmonize_text(doc)

        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.assertEqual(p.paragraph_format.line_spacing.pt, 18)

    def test_03_referencias(self):
        doc = Document()
        st = doc.styles.add_style(
            "Bibliography", WD_STYLE_TYPE.PARAGRAPH
        )
        p = doc.add_paragraph("Pollock et al. 2026", style=st)

        mod._mppg_docx_harmonize_text(doc)

        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(p.paragraph_format.line_spacing.pt, 15)
        self.assertEqual(p.paragraph_format.space_after.pt, 4)
        self.assertEqual(p.paragraph_format.first_line_indent.cm, 0)

    def test_04_titulos(self):
        doc = Document()
        heading = doc.add_heading("1. Tema", level=1)

        mod._mppg_docx_harmonize_text(doc)

        self.assertEqual(
            heading.alignment, WD_ALIGN_PARAGRAPH.LEFT
        )
        self.assertEqual(
            str(doc.styles["Heading 1"].font.color.rgb), "000000"
        )
        self.assertEqual(
            doc.styles["Heading 1"].font.name, "Times New Roman"
        )
        self.assertEqual(
            heading.paragraph_format.first_line_indent.cm, 0
        )

    def test_05_preserva_ficha_e_links(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "FGV"

        tabela_antes = table._tbl.xml

        url = "https://doi.org/10.7326/M18-0850"
        rid = doc.part.relate_to(
            url, RT.HYPERLINK, is_external=True
        )

        p = doc.add_paragraph()
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), rid)

        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "DOI"
        run.append(text)
        hyperlink.append(run)
        p._p.append(hyperlink)

        def links():
            return {
                key: rel.target_ref
                for key, rel in doc.part.rels.items()
                if rel.reltype == RT.HYPERLINK
            }

        links_antes = links()
        header_antes = doc.sections[0].header._element.xml
        footer_antes = doc.sections[0].footer._element.xml

        mod._mppg_docx_harmonize_text(doc)

        self.assertEqual(tabela_antes, table._tbl.xml)
        self.assertEqual(links_antes, links())
        self.assertEqual(
            header_antes, doc.sections[0].header._element.xml
        )
        self.assertEqual(
            footer_antes, doc.sections[0].footer._element.xml
        )

    def test_06_idempotencia_e_estilos_alheios(self):
        doc = Document()
        quote = doc.add_paragraph("Uma citacao", style="Quote")
        body = doc.add_paragraph(
            "Texto principal", style="Normal"
        )
        blank = doc.add_paragraph("", style="Normal")

        mod._mppg_docx_harmonize_text(doc)
        primeiro = doc._element.xml

        mod._mppg_docx_harmonize_text(doc)

        self.assertEqual(primeiro, doc._element.xml)
        self.assertIsNone(quote.paragraph_format.line_spacing)
        self.assertIsNone(blank.paragraph_format.line_spacing)
        self.assertEqual(body.paragraph_format.line_spacing.pt, 18)

    def test_07_persistencia_docx(self):
        doc = Document()
        doc.add_paragraph("Testando", style="Body Text")

        mod._mppg_docx_harmonize_text(doc)

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "teste.docx"
            doc.save(path)

            reaberto = Document(path)
            p = reaberto.paragraphs[0]

            self.assertEqual(p.paragraph_format.line_spacing.pt, 18)
            self.assertAlmostEqual(
                p.paragraph_format.first_line_indent.cm,
                1.25,
                places=2,
            )
            self.assertEqual(
                str(reaberto.styles["Heading 1"].font.color.rgb),
                "000000",
            )


if __name__ == "__main__":
    unittest.main()
