import unittest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from academic_pipeline import minuta_runtime as mod


def documento_com_referencia():
    doc = Document()
    paragraph = doc.add_paragraph("Referencia de teste")

    ppr = paragraph._p.get_or_add_pPr()
    style = OxmlElement("w:pStyle")
    style.set(qn("w:val"), "Bibliography")
    ppr.insert(0, style)

    return doc


class TestMinutaDocxHyperlinksFGV(unittest.TestCase):

    def test_01_cria_estilo_ausente(self):
        doc = documento_com_referencia()

        with self.assertRaises(KeyError):
            doc.styles["Bibliography"]

        mod._mppg_docx_ensure_bibliography_style(doc)

        style = doc.styles["Bibliography"]

        self.assertEqual(
            style.type, WD_STYLE_TYPE.PARAGRAPH
        )
        self.assertEqual(
            style.base_style.style_id, "Normal"
        )

    def test_02_nao_duplica_estilo(self):
        doc = documento_com_referencia()

        mod._mppg_docx_ensure_bibliography_style(doc)
        mod._mppg_docx_ensure_bibliography_style(doc)

        matching = [
            s for s in doc.styles
            if s.style_id == "Bibliography"
        ]

        self.assertEqual(len(matching), 1)

    def test_03_nao_altera_minuta_sem_bibliografia(self):
        doc = Document()
        doc.add_paragraph("Texto sem bibliografia.")

        antes = len(doc.styles)

        mod._mppg_docx_ensure_bibliography_style(doc)

        self.assertEqual(len(doc.styles), antes)

        with self.assertRaises(KeyError):
            doc.styles["Bibliography"]

    def test_04_preserva_estilo_preexistente(self):
        doc = documento_com_referencia()

        style = doc.styles.add_style(
            "Bibliography", WD_STYLE_TYPE.PARAGRAPH
        )
        style.font.bold = True

        mod._mppg_docx_ensure_bibliography_style(doc)

        self.assertTrue(
            doc.styles["Bibliography"].font.bold
        )


if __name__ == "__main__":
    unittest.main()
