import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from academic_pipeline import minuta_runtime as mod


class TestMinutaDocxPaginacaoFGV(unittest.TestCase):

    def test_01_campo_page_criado(self):
        doc = Document()

        mod._mppg_docx_page_numbers(doc)

        footer = doc.sections[0].footer

        instr = [
            n.text
            for n in footer._element.xpath(".//w:instrText")
        ]

        self.assertIn(" PAGE ", instr)

    def test_02_rodape_centralizado(self):
        doc = Document()

        mod._mppg_docx_page_numbers(doc)

        paragraph = doc.sections[0].footer.paragraphs[0]

        self.assertEqual(
            paragraph.alignment,
            WD_ALIGN_PARAGRAPH.CENTER,
        )

    def test_03_nao_duplica_page(self):
        doc = Document()

        mod._mppg_docx_page_numbers(doc)
        mod._mppg_docx_page_numbers(doc)

        footer = doc.sections[0].footer

        campos = footer._element.xpath(".//w:instrText")

        self.assertEqual(len(campos), 1)

    def test_04_preserva_rodape_existente(self):
        doc = Document()
        doc.sections[0].footer.paragraphs[0].text = (
            "Rodape institucional"
        )

        mod._mppg_docx_page_numbers(doc)

        footer = doc.sections[0].footer

        self.assertEqual(
            footer.paragraphs[0].text,
            "Rodape institucional",
        )

        self.assertGreaterEqual(len(footer.paragraphs), 2)


if __name__ == "__main__":
    unittest.main()
