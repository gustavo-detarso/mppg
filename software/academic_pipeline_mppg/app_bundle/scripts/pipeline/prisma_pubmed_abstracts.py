#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""NCBI PubMed EFetch/XML: resumos por PMID; CLI de enriquecimento PRISMA existente.

Sem dependências externas para o enriquecimento JSON/CSV; usa openpyxl instalado
no ambiente do MPPG para conservar planilhas XLSX e decisões humanas.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
import time
from typing import Any, Callable
import xml.etree.ElementTree as ET
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

BASE = 'https://eutils.ncbi.nlm.nih.gov/entrez/eutils/'
MAX_BATCH = 80
LAST_NCBI_REQUEST = 0.0


def ncbi_wait(has_key: bool = False) -> None:
    """Serializa requisições deste processo, abaixo dos limites NCBI (3/10 req/s)."""
    global LAST_NCBI_REQUEST
    minimum = 0.36 if not has_key else 0.12
    delay = minimum - (time.monotonic() - LAST_NCBI_REQUEST)
    if delay > 0:
        time.sleep(delay)
    LAST_NCBI_REQUEST = time.monotonic()


def _clean(s: Any) -> str:
    return re.sub(r'\s+', ' ', str(s or '')).strip()


def extract_pubmed_abstracts(xml_data: str) -> dict[str, str]:
    """Reconstrói AbstractText com Label, incluindo nós XML internos/itálicos."""
    root = ET.fromstring(xml_data)
    out: dict[str, str] = {}
    articles = list(root.iter('PubmedArticle')) + list(root.iter('PubmedBookArticle'))
    for article in articles:
        pmid = _clean(article.findtext('./MedlineCitation/PMID'))
        if not pmid:
            pmid = _clean(article.findtext('./BookDocument/PMID'))
        if not pmid or not pmid.isdecimal():
            continue
        sections = article.findall('./MedlineCitation/Article/Abstract/AbstractText')
        if not sections:
            sections = article.findall('./MedlineCitation/OtherAbstract/AbstractText')
        if not sections:
            sections = article.findall('./BookDocument/Abstract/AbstractText')
        parts = []
        for part in sections:
            body = _clean(''.join(part.itertext()))
            if not body:
                continue
            label = _clean(part.attrib.get('Label') or '')
            if not label:
                nlm = _clean(part.attrib.get('NlmCategory'))
                label = nlm if nlm.upper() not in ('UNASSIGNED', 'UNLABELLED') else ''
            parts.append(f'{label}: {body}' if label and not body.lower().startswith(label.lower() + ':') else body)
        if parts:
            out[pmid] = '\n'.join(parts)
    return out


def _default_request_text(url: str, **kwargs: Any) -> str:
    """Cliente stdlib com retry somente em falhas transitórias, sem expor API key."""
    before = kwargs.get('before_attempt')
    retries = max(0, int(kwargs.get('http_retries', 2)))
    for attempt in range(retries + 1):
        if before is not None:
            before()
        try:
            req = Request(url, headers={'User-Agent': 'AcademicPipeline-PRISMA/1.1', 'Accept': 'application/xml, application/json'})
            with urlopen(req, timeout=40) as response:
                return response.read().decode('utf-8', errors='replace')
        except HTTPError as exc:
            status = exc.code
            exc.close()
            if status not in (429, 500, 502, 503, 504) or attempt >= retries:
                raise RuntimeError(f'NCBI: HTTP {status}') from exc
        except (URLError, TimeoutError) as exc:
            if attempt >= retries:
                raise RuntimeError('NCBI: falha de rede/timeout') from exc
        time.sleep(min(12, 1.5 * (2 ** attempt)))
    raise RuntimeError('NCBI: erro desconhecido')


def fetch_pubmed_abstracts(
    ids: list[str], *, request_text: Callable[..., str] | None = None,
    email: str = '', api_key: str = '', batch_size: int = MAX_BATCH,
) -> tuple[dict[str, str], list[str]]:
    """EFetch batelado; erros por lote não impedem resultados dos demais lotes."""
    requester = request_text or _default_request_text
    valid = list(dict.fromkeys(str(v).strip() for v in ids if str(v).strip().isdecimal()))
    results: dict[str, str] = {}
    warnings: list[str] = []
    for start in range(0, len(valid), max(1, min(batch_size, MAX_BATCH))):
        batch = valid[start:start + max(1, min(batch_size, MAX_BATCH))]
        params = {'db': 'pubmed', 'id': ','.join(batch), 'retmode': 'xml', 'tool': 'AcademicPipelinePRISMA'}
        if email:
            params['email'] = email
        if api_key:
            params['api_key'] = api_key
        url = BASE + 'efetch.fcgi?' + urlencode(params)
        try:
            content = requester(
                url, before_attempt=lambda: ncbi_wait(bool(api_key)),
                retryable_http_statuses=(429, 500, 502, 503, 504),
                http_retries=2, retry_base_seconds=1.5, retry_max_seconds=12,
                retry_provider='pubmed',
            )
            parsed = extract_pubmed_abstracts(content)
            results.update({k: v for k, v in parsed.items() if k in batch})
        except (RuntimeError, ValueError, ET.ParseError, OSError) as exc:
            warnings.append(f'Lote EFetch {start // MAX_BATCH + 1}: {type(exc).__name__} (sem resumo recuperado).')
    return results, warnings


def _pmid_from_record(row: dict[str, Any]) -> str:
    rid = str(row.get('id_registro') or '')
    for pattern in (r'^pubmed:(\d+)$', r'^europe_pmc:(\d+)$'):
        m = re.match(pattern, rid, flags=re.IGNORECASE)
        if m:
            return m.group(1)
    for key in ('pmid', 'pubmed_id'):
        val = str(row.get(key) or '').strip()
        if val.isdecimal():
            return val
    url = str(row.get('url') or '')
    m = re.search(r'pubmed\.ncbi\.nlm\.nih\.gov/(\d+)', url, flags=re.IGNORECASE)
    return m.group(1) if m else ''


def _has_pubmed_source(row: dict[str, Any]) -> bool:
    src = row.get('fontes') or row.get('fonte') or []
    if isinstance(src, str):
        src = [src]
    return any(str(v).strip().lower() == 'pubmed' for v in src)


def _doi(row: dict[str, Any]) -> str:
    value = str(row.get('doi') or '').strip().lower()
    value = re.sub(r'^https?://(?:dx\.)?doi\.org/', '', value)
    value = re.sub(r'^doi:\s*', '', value)
    return value if value.startswith('10.') else ''


def lookup_pmid_by_doi(doi: str, *, email: str = '', api_key: str = '', request_text=None) -> str:
    """Busca identificador por DOI em [AID]; exige exatamente um PMID retornado."""
    params = {'db': 'pubmed', 'term': f'{doi}[AID]', 'retmax': 3, 'retmode': 'json', 'tool': 'AcademicPipelinePRISMA'}
    if email:
        params['email'] = email
    if api_key:
        params['api_key'] = api_key
    url = BASE + 'esearch.fcgi?' + urlencode(params)
    getter = request_text or _default_request_text
    raw = getter(url, before_attempt=lambda: ncbi_wait(bool(api_key)),
                 retryable_http_statuses=(429, 500, 502, 503, 504), http_retries=2,
                 retry_base_seconds=1.5, retry_max_seconds=12, retry_provider='pubmed')
    ids = (json.loads(raw).get('esearchresult') or {}).get('idlist') or []
    return str(ids[0]) if len(ids) == 1 and str(ids[0]).isdecimal() else ''


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _csv_update(content: bytes, changes: dict[str, str]) -> bytes:
    text = content.decode('utf-8-sig')
    rdr = csv.DictReader(io.StringIO(text, newline=''))
    if not rdr.fieldnames or 'id_registro' not in rdr.fieldnames or 'resumo' not in rdr.fieldnames:
        raise RuntimeError('CSV de triagem não contém colunas id_registro/resumo.')
    rows = list(rdr)
    ids = {str(r['id_registro']) for r in rows}
    if not changes.keys() <= ids:
        raise RuntimeError('CSV de triagem não contém todos os IDs alterados; operação recusada.')
    for row in rows:
        if row['id_registro'] in changes and not _clean(row['resumo']):
            row['resumo'] = changes[row['id_registro']]
    output = io.StringIO(newline='')
    writer = csv.DictWriter(output, fieldnames=rdr.fieldnames, lineterminator='\n', extrasaction='ignore')
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode('utf-8-sig' if content.startswith(b'\xef\xbb\xbf') else 'utf-8')


def _xlsx_update(content: bytes, changes: dict[str, str]) -> bytes:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError('openpyxl não disponível: execute este comando via pipenv run python.') from exc
    wb = openpyxl.load_workbook(io.BytesIO(content))
    matched: set[str] = set()
    for sheet in wb:
        header = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value is not None}
        if not {'id_registro', 'resumo'} <= header.keys():
            continue
        for row in sheet.iter_rows(min_row=2):
            rid = str(sheet.cell(row=row[0].row, column=header['id_registro']).value or '')
            if rid in changes:
                matched.add(rid)
                cell = sheet.cell(row=row[0].row, column=header['resumo'])
                if not _clean(cell.value):
                    cell.value = changes[rid]
    if not changes.keys() <= matched:
        raise RuntimeError('XLSX de triagem não contém todos os IDs alterados; operação recusada.')
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _atomic_bytes(path: Path, data: bytes) -> None:
    with tempfile.NamedTemporaryFile(dir=path.parent, prefix='.prisma_ncbi_', suffix=path.suffix, delete=False) as f:
        temp = Path(f.name)
        f.write(data)
        f.flush()
        os.fsync(f.fileno())
    try:
        if path.exists():
            os.chmod(temp, path.stat().st_mode & 0o777)
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


def enrich_existing(input_path: Path, *, apply: bool = False, resolve_doi: bool = True,
                    request_text=None, project_root: Path | None = None) -> dict[str, Any]:
    """Atualiza apenas resumos faltantes e mantém triagem XLSX/CSV sincronizada.

    Sem --apply não usa rede, não cria arquivo e apenas conta candidatos.
    """
    input_path = input_path.expanduser().resolve()
    original = input_path.read_bytes()
    records = json.loads(original)
    if not isinstance(records, list) or not all(isinstance(r, dict) for r in records):
        raise ValueError('Esperada lista JSON de registros PRISMA.')
    pending = [r for r in records if _has_pubmed_source(r) and not _clean(r.get('resumo'))]
    direct = [r for r in pending if _pmid_from_record(r)]
    uncertain = [r for r in pending if not _pmid_from_record(r) and _doi(r)]
    report: dict[str, Any] = {
        'schema_version': '1.0', 'registros': len(records),
        'pubmed_sem_resumo': len(pending), 'com_pmid_direto': len(direct),
        'necessitam_resolver_doi': len(uncertain), 'alterados': 0,
        'status': 'simulacao' if not apply else 'iniciado', 'sem_busca_tematica': True,
    }
    if not apply:
        return report
    from copy import deepcopy
    changed = deepcopy(records)
    email = os.getenv('NCBI_EMAIL', '').strip()
    api_key = os.getenv('NCBI_API_KEY', '').strip()
    candidates: dict[str, list[dict[str, Any]]] = {}
    warnings: list[str] = []
    for row in changed:
        if not _has_pubmed_source(row) or _clean(row.get('resumo')):
            continue
        pmid = _pmid_from_record(row)
        if not pmid and resolve_doi and _doi(row):
            try:
                pmid = lookup_pmid_by_doi(_doi(row), email=email, api_key=api_key, request_text=request_text)
            except (RuntimeError, OSError, ValueError) as exc:
                warnings.append(f'Falha ao resolver DOI de {row.get("id_registro", "")}: {type(exc).__name__}')
        if pmid:
            candidates.setdefault(pmid, []).append(row)
    abstracts, fetch_warns = fetch_pubmed_abstracts(list(candidates), email=email, api_key=api_key,
                                                    request_text=request_text)
    warnings.extend(fetch_warns)
    changes: dict[str, str] = {}
    for pmid, matches in candidates.items():
        value = abstracts.get(pmid, '')
        for row in matches:
            if value and not _clean(row.get('resumo')):
                row['resumo'] = value
                changes[str(row.get('id_registro') or '')] = value
    report.update({'status': 'sem_alteracoes' if not changes else 'concluido',
                   'alterados': len(changes), 'sem_resumo_apos': len(pending) - len(changes),
                   'resumos_efetch': len(abstracts), 'avisos': warnings})
    if not changes:
        return report

    # Todos os dados são preparados em memória ANTES de modificar qualquer arquivo.
    stem = input_path.name.removesuffix('.candidatos_deduplicados.json')
    if stem == input_path.name:
        raise ValueError('Arquivo não corresponde ao padrão .candidatos_deduplicados.json')
    related = [input_path.parent / (stem + ext) for ext in ('.triagem_titulo_resumo.csv', '.triagem_titulo_resumo.xlsx')]
    originals: dict[Path, bytes] = {input_path: original}
    new_content: dict[Path, bytes] = {input_path: (json.dumps(changed, ensure_ascii=False, indent=2) + '\n').encode('utf-8')}
    for path in related:
        if path.exists():
            old = path.read_bytes()
            originals[path] = old
            new_content[path] = _csv_update(old, changes) if path.suffix == '.csv' else _xlsx_update(old, changes)
    if not any(p.exists() for p in related):
        raise RuntimeError('Nenhuma planilha de triagem encontrada; sincronização recusada.')

    root = (project_root or Path(__file__).resolve().parents[3]).resolve()
    backup = root / '.academic_pipeline_backups' / ('prisma_pubmed_efetch_' + datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S_%f'))
    backup.mkdir(parents=True, exist_ok=False)
    manifest: dict[str, Any] = {'created_at': datetime.now(timezone.utc).isoformat(), 'files': []}
    for idx, (path, old) in enumerate(originals.items(), 1):
        copy_path = backup / f'{idx}_{path.name}'
        copy_path.write_bytes(old)
        manifest['files'].append({'original': str(path), 'backup': str(copy_path),
                                  'original_sha256': _sha(old), 'enriched_sha256': _sha(new_content[path])})
    (backup / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    written: list[Path] = []
    try:
        for path, new_data in new_content.items():
            if _sha(path.read_bytes()) != _sha(originals[path]):
                raise RuntimeError('Arquivo alterado simultaneamente; escrita cancelada.')
            _atomic_bytes(path, new_data)
            written.append(path)
    except Exception:
        for path in reversed(written):
            _atomic_bytes(path, originals[path])
        raise
    report['backup'] = str(backup)
    report_path = input_path.parent / f'{stem}.pubmed_enriquecimento_{datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")}.json'
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    report['relatorio'] = str(report_path)
    return report


def revert_enrichment(backup_dir: Path, *, project_root: Path | None = None) -> dict[str, Any]:
    """Restaura somente se arquivos atuais ainda possuem o hash da atualização."""
    root = (project_root or Path(__file__).resolve().parents[3]).resolve()
    backup_dir = backup_dir.expanduser().resolve()
    allowed = (root / '.academic_pipeline_backups').resolve()
    if backup_dir.parent != allowed or not backup_dir.name.startswith('prisma_pubmed_efetch_'):
        raise RuntimeError('Backup não pertence ao enriquecimento deste MPPG.')
    manifest = json.loads((backup_dir / 'manifest.json').read_text(encoding='utf-8'))
    restore: dict[Path, bytes] = {}
    new_bytes: dict[Path, bytes] = {}
    for entry in manifest['files']:
        target = Path(entry['original']).resolve()
        backup = Path(entry['backup']).resolve()
        if backup.parent != backup_dir or not backup.is_file() or not target.is_file():
            raise RuntimeError('Manifest de backup contém arquivo ausente ou fora do diretório.')
        saved = backup.read_bytes()
        current = target.read_bytes()
        if _sha(saved) != entry['original_sha256'] or _sha(current) != entry['enriched_sha256']:
            raise RuntimeError('Arquivo alterado após enriquecimento; reversão automática recusada.')
        restore[target] = saved
        new_bytes[target] = current
    restored: list[Path] = []
    try:
        for target, saved in restore.items():
            _atomic_bytes(target, saved)
            restored.append(target)
    except Exception:
        for target in reversed(restored):
            _atomic_bytes(target, new_bytes[target])
        raise
    return {'status': 'reversao_concluida', 'arquivos_restaurados': len(restored), 'backup': str(backup_dir)}

def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description='Enriquece resumos PubMed sem refazer buscas temáticas nem decisões de triagem.')
    parser.add_argument('--input', type=Path, help='Arquivo *.candidatos_deduplicados.json')
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument('--dry-run', action='store_true', help='Conta candidatos, sem rede e sem escrita')
    group.add_argument('--apply', action='store_true', help='Consulta NCBI, faz backup e atualiza JSON/CSV/XLSX')
    group.add_argument('--reverter-enriquecimento', type=Path, help='Restaura um backup verificado de JSON/CSV/XLSX')
    parser.add_argument('--sem-resolver-doi', action='store_true', help='Evita consultas para localizar PMIDs por DOI')
    args = parser.parse_args(argv)
    if not args.reverter_enriquecimento and args.input is None:
        parser.error('Informe --input nos modos --dry-run ou --apply.')
    try:
        if args.reverter_enriquecimento:
            result = revert_enrichment(args.reverter_enriquecimento)
        else:
            result = enrich_existing(args.input, apply=args.apply, resolve_doi=not args.sem_resolver_doi)
    except Exception as exc:
        print(f'ERRO_PUBMED_EFETCH: {type(exc).__name__}: {exc}', file=sys.stderr)
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
