#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Pipeline integrado de pesquisa + documento acadêmico para o projeto MPPG, com suporte a mock_run.
Versão v0.2.6: inclui enriquecimento bibliográfico de documentos locais por bases externas,
deduplicação final do .bib e suporte a DOI manual via CSV/TOML.

Objetivo:
1. Ler um TOML unificado.
2. Extrair dele um TOML compatível com o gerador_pesquisa_rc_2.py.
3. Executar opcionalmente o gerador_pesquisa_rc_2.py (PRISMA ou empírico).
4. Montar um bundle de handoff com os artefatos gerados.
5. Gerar opcionalmente um documento acadêmico em Org-mode a partir do contexto e dos artefatos da pesquisa.
6. Preservar o .org da pesquisa e o .org do documento acadêmico como artefatos distintos.

Este script não tenta reimplementar o motor metodológico do gerador_pesquisa_rc_2.py.
Ele o reutiliza como etapa de pesquisa e assume a etapa redacional do documento acadêmico, consumindo explicitamente os PDFs selecionados na pesquisa quando disponíveis.
"""
from __future__ import annotations

import argparse
import json
import os
import random
import re
import shutil
import shlex
import subprocess
import sys
import textwrap
import tomllib
import traceback
import zipfile
import csv
import difflib
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from dotenv import load_dotenv
from openai import OpenAI
from pydantic import BaseModel, Field
from pypdf import PdfReader

try:
    import docx  # type: ignore
except Exception:  # pragma: no cover
    docx = None

# ---------------------------------------------------------------------------
# Defaults
# ---------------------------------------------------------------------------
DEFAULT_MODEL = os.getenv("OPENAI_MODEL", "gpt-5.4")
DEFAULT_RESEARCH_SCRIPT = "./gerador_pesquisa_rc_2.py"
DEFAULT_AUTHOR = "Gustavo M. Mendes de Tarso"
DEFAULT_INSTITUTION = "Fundação Getúlio Vargas"
DEFAULT_SCHOOL = ""
DEFAULT_PAPER_TEMPLATE = "template_paper.org"
DEFAULT_DISSERTATION_TEMPLATE = "template_dissertacao.org"
DEFAULT_FALLBACK_TEMPLATE = "template_research.org"
RESEARCH_BRIDGE_OUT_ORIENTACAO = "_".join(["arquivo", "orientacao"])
RESEARCH_BRIDGE_TRIAGEM_PROMPT = "_".join(["triagem", "prompt", "path"])
RESEARCH_BRIDGE_TRIAGEM_EXTRAS = "_".join(["diretivas", "extras"])
DEFAULT_STYLE = "apa"
DEBUG = False

DEFAULT_DOC_TYPE = "paper"
DEFAULT_DOC_PREFIX_BY_TYPE = {
    "paper": "paper",
    "dissertacao": "dissertacao",
    "atividade": "atividade",
    "resumo": "resumo",
    "resumo_expandido": "resumo_expandido",
    "fichamento": "fichamento",
    "resposta_discursiva": "resposta_discursiva",
    "ensaio": "ensaio",
    "ensaio_curto": "ensaio_curto",
}

def normalize_document_type(raw: str | None) -> str:
    value = (raw or DEFAULT_DOC_TYPE).strip().lower()
    aliases = {
        "paper": "paper",
        "artigo": "paper",
        "papel": "paper",
        "documento": "paper",
        "dissertacao": "dissertacao",
        "dissertação": "dissertacao",
        "thesis": "dissertacao",
        "atividade": "atividade",
        "atividade_fgv": "atividade",
        "resumo": "resumo",
        "resumo_expandido": "resumo_expandido",
        "fichamento": "fichamento",
        "resposta": "resposta_discursiva",
        "resposta_discursiva": "resposta_discursiva",
        "ensaio": "ensaio",
        "ensaio_curto": "ensaio_curto",
    }
    return aliases.get(value, value)


def document_type_label(raw: str | None, *, article: bool = False) -> str:
    """Rótulo amigável para mensagens de log conforme tipo_documento do TOML."""
    doc_type = normalize_document_type(raw)
    labels = {
        "paper": "paper",
        "dissertacao": "dissertação",
        "atividade": "atividade",
        "resumo": "resumo",
        "resumo_expandido": "resumo expandido",
        "fichamento": "fichamento",
        "resposta_discursiva": "resposta discursiva",
        "ensaio": "ensaio",
        "ensaio_curto": "ensaio curto",
    }
    label = labels.get(doc_type, doc_type.replace("_", " "))
    if not article:
        return label
    feminine = {"dissertacao", "atividade", "resposta_discursiva"}
    return ("a " if doc_type in feminine else "o ") + label


RESEARCH_SECTIONS = {
    "atividade",
    "pesquisa",
    "bibliografia",
    "busca",
    "triagem",
    "queries",
    "saida",
    "latex",
    "openai",
    "controle",
}

TEXT_SUFFIXES = {".txt", ".md", ".org", ".rst", ".tex", ".json", ".csv", ".yaml", ".yml", ".xml"}
BINARY_SUFFIXES = {".pdf", ".docx"}
READABLE_SUFFIXES = TEXT_SUFFIXES | BINARY_SUFFIXES
ORIENTATION_ARCHIVE_SUFFIXES = {".zip"}
MAX_ORIENTATION_FILES_FROM_ARCHIVE = 80


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------
@dataclass
class ResearchPaths:
    root_dir: Path
    org_path: Path | None = None
    bib_path: Path | None = None
    debug_path: Path | None = None
    pdf_path: Path | None = None
    prisma_svg_path: Path | None = None
    prisma_pdf_path: Path | None = None
    config_path: Path | None = None
    fulltext_cache_dir: Path | None = None
    selected_entries: list[dict[str, Any]] = field(default_factory=list)
    selected_fulltext_paths: list[Path] = field(default_factory=list)
    fulltext_paths: list[Path] = field(default_factory=list)

@dataclass
class SourceDoc:
    path: str
    kind: str
    label: str
    extracted_text: str
    summary: str | None = None
    bib_key: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)

@dataclass
class DocumentContext:
    tema: str
    recorte: str
    objetivo: str
    pergunta_pesquisa: str | None = None
    hipotese: str | None = None
    palavras_chave: list[str] = field(default_factory=list)
    titulo_sugerido: str | None = None
    tipo_estudo: str | None = None
    idiomas: list[str] = field(default_factory=lambda: ["português", "inglês"])
    modo_origem: str = "revisao_sistematica"
    titulo_trabalho_base: str | None = None

class FinalFrontMatterOutput(BaseModel):
    title: str
    documento_type: str
    cover_note: str

class InferredBibMetadata(BaseModel):
    entry_type: str = "article"
    title: str
    authors: list[str] = Field(default_factory=list)
    editors: list[str] = Field(default_factory=list)
    year: str | None = None
    booktitle: str | None = None
    journaltitle: str | None = None
    publisher: str | None = None
    location: str | None = None
    pages: str | None = None
    volume: str | None = None
    number: str | None = None
    edition: str | None = None
    doi: str | None = None
    isbn: str | None = None
    url: str | None = None
    note: str | None = None

class RewrittenDocumentContextOutput(BaseModel):
    tema: str
    recorte: str
    objetivo: str
    pergunta_pesquisa: str | None = None
    hipotese: str | None = None
    titulo_sugerido: str | None = None
    rationale: str | None = None

class InferredDocumentContextOutput(BaseModel):
    titulo_sugerido: str
    tema: str
    recorte: str
    objetivo: str
    pergunta_pesquisa: str | None = None
    hipotese: str | None = None
    palavras_chave: list[str] = Field(default_factory=list)
    rationale: str | None = None

class DissertationFrontSectionsOutput(BaseModel):
    dedicatoria: str = ""
    agradecimentos: str = ""
    epigrafe_texto: str = ""
    epigrafe_autor: str = ""
    resumo: str
    palavras_chave: list[str] = Field(default_factory=list)
    abstract: str
    keywords: list[str] = Field(default_factory=list)
    lista_ilustracoes: list[str] = Field(default_factory=list)
    lista_tabelas: list[str] = Field(default_factory=list)
    siglas: list[str] = Field(default_factory=list)
    simbolos: list[str] = Field(default_factory=list)
    glossario: list[str] = Field(default_factory=list)

class DissertationStageBodyOutput(BaseModel):
    org_body: str

# ---------------------------------------------------------------------------
# Utils
# ---------------------------------------------------------------------------
def debug_print(*parts: object) -> None:
    if DEBUG:
        print("[DEBUG]", *parts, file=sys.stderr)

def safe_resolve_user_path(raw: Any, base_dir: Path | None = None) -> Path | None:
    text = str(raw).strip()
    if not text:
        return None
    looks_like_path = (
        text.startswith(("./", "../", "~/", "/"))
        or os.sep in text
        or (os.altsep is not None and os.altsep in text)
        or bool(Path(text).suffix)
    )
    if ("\n" in text) or (len(text) > 240 and not looks_like_path):
        return None
    try:
        path = Path(os.path.expanduser(text))
        if not path.is_absolute():
            path = ((base_dir or Path.cwd()) / path).resolve()
        else:
            path = path.resolve()
        return path
    except OSError:
        return None
    except Exception:
        return None


def get_config_base_dir(cfg: dict[str, Any]) -> Path:
    raw = cfg.get("__config_dir__")
    if raw:
        return Path(str(raw)).expanduser().resolve()
    return Path.cwd().resolve()


def resolve_configured_path(raw: Any, cfg: dict[str, Any], *, base_dir: Path | None = None) -> Path | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    bases: list[Path] = []
    if base_dir is not None:
        bases.append(base_dir.resolve())
    cfg_base = get_config_base_dir(cfg)
    if cfg_base not in bases:
        bases.append(cfg_base)
    cwd = Path.cwd().resolve()
    if cwd not in bases:
        bases.append(cwd)
    cfg_parent = cfg_base.parent.resolve()
    if cfg_parent not in bases:
        bases.append(cfg_parent)
    for candidate_base in bases:
        resolved = safe_resolve_user_path(text, base_dir=candidate_base)
        if resolved is not None and resolved.exists():
            return resolved
    return safe_resolve_user_path(text, base_dir=bases[0] if bases else None)

def safe_path_exists_file(path: Path | None) -> bool:
    if path is None:
        return False
    try:
        return path.exists() and path.is_file()
    except OSError:
        return False
    except Exception:
        return False

def load_env() -> None:
    load_dotenv(override=False)

def make_client(model_override: str | None = None) -> tuple[OpenAI, str]:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY não encontrado no .env ou no ambiente.")
    return OpenAI(api_key=api_key), (model_override or os.getenv("OPENAI_MODEL") or DEFAULT_MODEL)

def slugify(text: str) -> str:
    text = (text or "").strip().lower()
    repl = {
        "á": "a", "à": "a", "â": "a", "ã": "a", "é": "e", "ê": "e", "í": "i",
        "ó": "o", "ô": "o", "õ": "o", "ú": "u", "ç": "c",
    }
    for k, v in repl.items():
        text = text.replace(k, v)
    text = re.sub(r"[^a-z0-9_\-\s]", "", text)
    text = re.sub(r"[\s\-]+", "_", text)
    return text.strip("_") or "item"

def ensure_command(name: str) -> str:
    found = shutil.which(name)
    if not found:
        raise RuntimeError(f"Comando não encontrado no PATH: {name}")
    return found

def shorten_text(text: str, limit: int = 8000) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    if len(text) <= limit:
        return text
    return text[:limit].rstrip() + "…"

def write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def json_safe(value: Any) -> Any:
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, dict):
        return {str(k): json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [json_safe(v) for v in value]
    return value

def read_text_file(path: Path, max_chars: int = 40000) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return shorten_text(path.read_text(encoding="utf-8", errors="ignore"), max_chars)
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        chunks: list[str] = []
        total = 0
        for page in reader.pages:
            text = page.extract_text() or ""
            if text:
                chunks.append(text)
                total += len(text)
                if total >= max_chars:
                    break
        return shorten_text("\n".join(chunks), max_chars)
    if suffix == ".docx":
        if docx is not None:
            d = docx.Document(str(path))
            return shorten_text("\n".join(p.text for p in d.paragraphs if p.text), max_chars)
        raise RuntimeError("python-docx não está disponível para ler .docx")
    raise RuntimeError(f"Arquivo não suportado para leitura textual: {path}")

def add_guidance_from_value(docs: list[SourceDoc], raw_value: Any, kind: str, label: str, max_chars: int = 20000) -> None:
    if raw_value is None:
        return
    raw = str(raw_value).strip()
    if not raw:
        return

    maybe_path = safe_resolve_user_path(raw)
    if safe_path_exists_file(maybe_path):
        try:
            if maybe_path.suffix.lower() in READABLE_SUFFIXES:
                text = read_text_file(maybe_path, max_chars=max_chars)
            else:
                text = shorten_text(maybe_path.read_text(encoding="utf-8", errors="ignore"), max_chars)
            docs.append(SourceDoc(
                path=str(maybe_path),
                kind=kind,
                label=label or maybe_path.name,
                extracted_text=text,
            ))
            return
        except Exception as exc:
            debug_print(f"Falha ao ler orientação {maybe_path}: {exc}")

    # fallback: trata o valor como texto inline
    docs.append(SourceDoc(
        path=f"inline:{kind}",
        kind=kind,
        label=label,
        extracted_text=shorten_text(raw, max_chars),
    ))

def read_template_raw(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    if raw.lstrip().startswith("# -*- mode: snippet -*-"):
        marker = "\n# --\n"
        pos = raw.find(marker)
        if pos != -1:
            return raw[pos + len(marker):]
    return raw

def replace_org_header_line(org_text: str, prefix: str, new_value: str) -> str:
    pattern = re.compile(rf'^(\s*{re.escape(prefix)}\s*).*$' , re.MULTILINE)
    if pattern.search(org_text):
        return pattern.sub(lambda m: f"{m.group(1)}{new_value}", org_text, count=1)
    return f"{prefix} {new_value}\n" + org_text

def replace_latex_header_macro(org_text: str, macro: str, new_value: str) -> str:
    pattern = re.compile(rf'^(\s*#\+LATEX_HEADER:\s*\\{re.escape(macro)}\{{).*(\}}\s*)$', re.MULTILINE)
    replacement = rf'\1{new_value}\2'
    if pattern.search(org_text):
        return pattern.sub(replacement, org_text, count=1)
    return org_text

def remove_latex_header_macro(org_text: str, macro: str) -> str:
    """Remove uma macro #+LATEX_HEADER de linha única, se existir.

    Usada para campos opcionais: se o TOML deixar o campo vazio, a macro
    não deve sequer aparecer no .org gerado.
    """
    pattern = re.compile(rf'^\s*#\+LATEX_HEADER:\s*\\{re.escape(macro)}\{{.*\}}\s*$', re.MULTILINE)
    return pattern.sub('', org_text)

def replace_or_insert_latex_header_macro(org_text: str, macro: str, new_value: str) -> str:
    updated = replace_latex_header_macro(org_text, macro, new_value)
    if updated != org_text:
        return updated
    lines = org_text.splitlines()
    insert_at = 0
    for i, line in enumerate(lines):
        if line.startswith("#+LATEX_HEADER:"):
            insert_at = i + 1
    lines.insert(insert_at, f"#+LATEX_HEADER: \\{macro}{{{new_value}}}")
    return "\n".join(lines)

def set_required_latex_header_macro(org_text: str, macro: str, value: str) -> str:
    return replace_or_insert_latex_header_macro(org_text, macro, str(value or '').strip())

def set_optional_latex_header_macro(org_text: str, macro: str, value: str | None) -> str:
    clean = re.sub(r"\s+", " ", str(value or "").strip())
    if not clean or is_placeholder_value(clean):
        return remove_latex_header_macro(org_text, macro)
    return replace_or_insert_latex_header_macro(org_text, macro, clean)

def set_optional_date_latex_header_macro(org_text: str, macro: str, value: str | None) -> str:
    clean = re.sub(r"\s+", " ", str(value or "").strip())
    # Data vazia ou placeholder de formulário deve desaparecer do .org.
    # Se o usuário escrever explicitamente "A definir", o valor é mantido.
    if not clean or re.search(r"[_]{2,}|/_[^ ]|^__/.+", clean):
        return remove_latex_header_macro(org_text, macro)
    return replace_or_insert_latex_header_macro(org_text, macro, clean)

def insert_latex_header_line(org_text: str, line: str) -> str:
    lines = org_text.splitlines()
    insert_at = 0
    for i, current in enumerate(lines):
        if current.startswith("#+LATEX_HEADER:"):
            insert_at = i + 1
    lines.insert(insert_at, line)
    return "\n".join(lines)

def split_title_and_subtitle(title: str) -> tuple[str, str]:
    raw = re.sub(r"\s+", " ", str(title or "").strip())
    if not raw:
        return "TÍTULO DA DISSERTAÇÃO", ""
    for sep in (" — ", " – ", ": "):
        if sep in raw:
            a, b = raw.split(sep, 1)
            return a.strip(), b.strip()
    return raw, ""


def derive_area_from_course(course_name: str) -> str:
    text = re.sub(r"\s+", " ", str(course_name or "").strip())
    if not text:
        return "Políticas Públicas e Governo"
    for prefix in ("Mestrado Profissional em ", "Mestrado em ", "Master in "):
        if text.lower().startswith(prefix.lower()):
            remainder = text[len(prefix):].strip()
            if remainder:
                return remainder
    return text



def is_placeholder_value(value: str | None) -> bool:
    text = re.sub(r"\s+", " ", str(value or "")).strip().lower()
    if not text:
        return True
    placeholders = {
        "seu nome",
        "nome completo do aluno",
        "nome do aluno",
        "autor da epígrafe",
        "orientador(a) a definir",
        "a definir",
        "[inserir]",
        "título do apêndice",
        "titulo do apendice",
        "título do anexo",
        "titulo do anexo",
    }
    if text in placeholders:
        return True
    return text.startswith("[inserir")

def normalize_banca_member(raw: Any) -> dict[str, str] | None:
    """Normaliza um item da banca vindo do TOML.

    Formatos aceitos:
      [[documento.banca]]
      nome = "Prof. Dr. Nome"
      funcao = "Orientador"
      instituicao = "Fundação Getúlio Vargas"

    Também aceita string no formato:
      "Orientador | Prof. Dr. Nome | Fundação Getúlio Vargas"
    """
    if isinstance(raw, dict):
        nome = str(raw.get("nome") or raw.get("name") or "").strip()
        funcao = str(raw.get("funcao") or raw.get("função") or raw.get("papel") or raw.get("tipo") or raw.get("role") or "").strip()
        instituicao = str(raw.get("instituicao") or raw.get("instituição") or raw.get("vinculo") or raw.get("vínculo") or raw.get("affiliation") or "").strip()
    else:
        parts = [part.strip() for part in str(raw or "").split("|")]
        if len(parts) >= 3:
            funcao, nome, instituicao = parts[0], parts[1], " | ".join(parts[2:])
        elif len(parts) == 2:
            funcao, nome, instituicao = parts[0], parts[1], ""
        else:
            funcao, nome, instituicao = "", str(raw or "").strip(), ""

    if is_placeholder_value(nome):
        return None
    if is_placeholder_value(funcao):
        funcao = ""
    if is_placeholder_value(instituicao):
        instituicao = ""
    return {"funcao": funcao, "nome": nome, "instituicao": instituicao}


def collect_banca_members(documento: dict[str, Any]) -> list[dict[str, str]]:
    """Coleta a banca examinadora dinâmica definida no TOML.

    A fonte preferencial é [[documento.banca]] ou [[documento.banca_examinadora]].
    Mantém compatibilidade com membro_banca_1, membro_banca_2 etc., mas sem
    preencher placeholders automáticos.
    """
    raw_members = documento.get("banca") or documento.get("banca_examinadora") or []
    members: list[dict[str, str]] = []
    if isinstance(raw_members, list):
        for raw in raw_members:
            parsed = normalize_banca_member(raw)
            if parsed:
                members.append(parsed)
    elif raw_members:
        parsed = normalize_banca_member(raw_members)
        if parsed:
            members.append(parsed)

    if members:
        return members

    # Compatibilidade com campos antigos. A função/cargo pode ser informada em
    # funcao_banca_1, papel_banca_1, tipo_banca_1 etc.
    for idx in range(1, 11):
        raw_name = documento.get(f"membro_banca_{idx}")
        if raw_name is None:
            continue
        nome = str(raw_name or "").strip()
        if is_placeholder_value(nome):
            continue
        funcao = str(
            documento.get(f"funcao_banca_{idx}")
            or documento.get(f"função_banca_{idx}")
            or documento.get(f"papel_banca_{idx}")
            or documento.get(f"tipo_banca_{idx}")
            or ""
        ).strip()
        instituicao = str(
            documento.get(f"instituicao_banca_{idx}")
            or documento.get(f"instituição_banca_{idx}")
            or documento.get(f"vinculo_banca_{idx}")
            or documento.get(f"vínculo_banca_{idx}")
            or ""
        ).strip()
        parsed = normalize_banca_member({"nome": nome, "funcao": funcao, "instituicao": instituicao})
        if parsed:
            members.append(parsed)
    return members


def clear_banca_headers(org_text: str) -> str:
    for macro in ("membrobanca", "membrobancaum", "membrobancadois", "membrobancatres"):
        # remove também a macro dinâmica com três argumentos.
        if macro == "membrobanca":
            org_text = re.sub(r'^\s*#\+LATEX_HEADER:\s*\\membrobanca\{.*\}\{.*\}\{.*\}\s*$', '', org_text, flags=re.MULTILINE)
        else:
            org_text = remove_latex_header_macro(org_text, macro)
    return org_text


def apply_banca_headers(org_text: str, documento: dict[str, Any]) -> str:
    org_text = clear_banca_headers(org_text)
    members = collect_banca_members(documento)
    for member in members:
        funcao = member.get("funcao", "")
        nome = member.get("nome", "")
        instituicao = member.get("instituicao", "")
        org_text = insert_latex_header_line(
            org_text,
            f"#+LATEX_HEADER: \\membrobanca{{{funcao}}}{{{nome}}}{{{instituicao}}}",
        )
    return org_text


def apply_dissertation_template_metadata(org_text: str, cfg: dict[str, Any], final_title: str) -> str:
    atividade = cfg.get("atividade", {})
    documento = cfg.get("documento", {})
    institution = str(documento.get("institution_name") or DEFAULT_INSTITUTION or "Fundação Getúlio Vargas").strip()
    institution_upper = institution.upper()
    raw_author = str(atividade.get("aluno") or DEFAULT_AUTHOR).strip()
    author = DEFAULT_AUTHOR if is_placeholder_value(raw_author) else raw_author
    raw_course_name = str(atividade.get("curso") or "Mestrado Profissional em Políticas Públicas e Governo").strip()
    course_name = "Mestrado Profissional em Políticas Públicas e Governo" if is_placeholder_value(raw_course_name) else raw_course_name
    raw_professor_name = str(atividade.get("professor") or "").strip()
    professor_name = "" if is_placeholder_value(raw_professor_name) else raw_professor_name
    raw_city_name = str(atividade.get("polo") or "Brasília").strip()
    city_name = "Brasília" if is_placeholder_value(raw_city_name) else raw_city_name
    raw_school_name = str(documento.get("school_name") or DEFAULT_SCHOOL).strip()
    school_name = "" if is_placeholder_value(raw_school_name) else raw_school_name
    raw_program_name = str(documento.get("program_name") or course_name).strip()
    program_name = course_name if is_placeholder_value(raw_program_name) else raw_program_name
    raw_area_name = str(documento.get("area_de_concentracao") or derive_area_from_course(course_name)).strip()
    area_name = derive_area_from_course(course_name) if is_placeholder_value(raw_area_name) else raw_area_name
    linha_pesquisa = str(documento.get("linha_pesquisa") or documento.get("linha_de_pesquisa") or "").strip()
    coorientador = str(documento.get("coorientador") or documento.get("co_orientador") or documento.get("coorientadora") or "").strip()
    dissertation_year = str(documento.get("ano") or datetime.now().year)
    title_main, title_sub = split_title_and_subtitle(final_title)
    default_nature_target = school_name if school_name else institution
    connector = " da Fundação Getúlio Vargas" if school_name and school_name.strip().lower() != institution.strip().lower() else ""
    nature_text = str(documento.get("natureza_trabalho") or f"Dissertação apresentada à {default_nature_target}{connector}, como requisito para obtenção do título de Mestre em {area_name}.").strip()

    required_replacements = {
        "autor": author,
        "titulo": title_main,
        "cidade": city_name,
        "ano": dissertation_year,
        "instituicao": institution_upper,
        "programa": program_name,
        "curso": course_name,
        "areadeconcentracao": area_name,
        "naturezatrabalho": nature_text,
    }
    optional_replacements = {
        "subtitulo": title_sub,
        "escola": school_name,
        "linhapesquisa": linha_pesquisa,
        "orientador": professor_name,
        "coorientador": coorientador,
    }

    for macro, value in required_replacements.items():
        org_text = set_required_latex_header_macro(org_text, macro, value)
    for macro, value in optional_replacements.items():
        org_text = set_optional_latex_header_macro(org_text, macro, value)

    org_text = set_optional_date_latex_header_macro(org_text, "dataaprovacao", documento.get("data_aprovacao"))
    org_text = apply_banca_headers(org_text, documento)

    # Remove linhas em branco excessivas criadas pela retirada de macros opcionais.
    org_text = re.sub(r"\n{3,}", "\n\n", org_text)
    return org_text

def apply_citation_style(org_text: str, bib_filename: str, style: str) -> str:
    style = (style or DEFAULT_STYLE).strip().lower() or DEFAULT_STYLE
    bib_name = Path(str(bib_filename)).name
    cite_line = f"#+CITE_EXPORT: biblatex {style}"
    org_text = re.sub(r"(?im)^\s*#\+CITE_EXPORT:.*\n?", "", org_text)
    org_text = re.sub(r"(?im)^\s*#\+BIBLIOGRAPHY:\s+.*\n?", "", org_text)
    org_text = re.sub(r"(?im)^\s*#\+PRINT_BIBLIOGRAPHY:?\s*(?:.*)?\n?", "", org_text)
    org_text = re.sub(r"(?im)^\s*#\+LATEX_HEADER:\s*\\ExecuteBibliographyOptions\{.*\}\s*\n?", "", org_text)
    org_text = re.sub(r"(?im)^\s*#\+LATEX_HEADER:\s*\\PassOptionsToPackage\{.*\}\{biblatex\}\s*\n?", "", org_text)
    lines = [ln for ln in org_text.splitlines() if ln.strip()]
    insert_at = 0
    for i, line in enumerate(lines):
        if line.startswith("#+LATEX_CLASS") or line.startswith("#+LATEX_CLASS_OPTIONS"):
            insert_at = i + 1
    lines.insert(insert_at, cite_line)
    lines.append(f"#+BIBLIOGRAPHY: {bib_name}")
    return "\n".join(lines).strip() + "\n"

def ensure_document_class(org_text: str, doc_type: str) -> str:
    class_name = "fgv-dissertacao" if normalize_document_type(doc_type) == "dissertacao" else "fgv-paper"
    org_text = re.sub(r"(?im)^\s*#\+LATEX_CLASS:.*\n?", "", org_text)
    return f"#+LATEX_CLASS: {class_name}\n" + org_text.lstrip("\n")

def normalize_bibliography_block(org_text: str) -> str:
    """Normaliza a bibliografia Org sem drawer ou heading manual.

    Para o template de atividade, a bibliografia deve ser inserida apenas por
    #+PRINT_BIBLIOGRAPHY:. O próprio exportador/biblatex cria o título visual
    de referências. Portanto, esta rotina remove blocos antigos como
    '* Referências', ':PROPERTIES:', ':UNNUMBERED: t' e recria somente a
    diretiva final.
    """
    text = org_text.replace("\r\n", "\n")

    # Remove uma seção final de Referências/Bibliography já existente, incluindo
    # drawer :PROPERTIES: e eventual #+PRINT_BIBLIOGRAPHY. Como a bibliografia
    # deve ficar no fim do arquivo, a remoção até EOF é intencional.
    text = re.sub(
        r"(?ims)\n*^\*+\s+(Refer[êe]ncias|References|Bibliography)\s*$.*\Z",
        "",
        text,
    )

    # Remove diretivas/drawers soltos que possam ter ficado órfãos.
    text = re.sub(r"(?im)^\s*#\+PRINT_BIBLIOGRAPHY:?\s*(?:.*)?\n?", "", text)
    text = re.sub(r"(?im)^\s*#\+BIBLIOGRAPHY:.*\n?", "", text)
    text = re.sub(r"(?im)^\s*:PROPERTIES:\s*$\n?", "", text)
    text = re.sub(r"(?im)^\s*:UNNUMBERED:\s*t\s*$\n?", "", text)
    text = re.sub(r"(?im)^\s*:END:\s*$\n?", "", text)

    # Remove forma colada em fim de parágrafo que já apareceu em PDFs gerados.
    text = re.sub(r"(?i)\s*:UNNUMBERED:\s*t\s*#\+PRINT_BIBLIOGRAPHY:?\s*", "\n\n", text)

    text = text.rstrip()
    return text + "\n\n#+PRINT_BIBLIOGRAPHY:\n"

def ensure_cover_command(org_text: str) -> str:
    if "\\usepapercover" not in org_text or "#+LATEX: \\makemytitle" in org_text:
        return org_text
    marker = "#+begin_abstract"
    if marker in org_text:
        return org_text.replace(marker, "#+LATEX: \\makemytitle\n\n" + marker, 1)
    return org_text + "\n#+LATEX: \\makemytitle\n"

def cleanup_generated_org(org_text: str) -> str:
    # Remove diretivas de bibliografia órfãs que a IA possa ter escrito como texto.
    # A forma correta será recriada ao final por normalize_bibliography_block().
    org_text = re.sub(r"(?im)^\s*#\+PRINT_BIBLIOGRAPHY:?\s*$\n?", "", org_text)
    org_text = re.sub(r"(?im)^\s*[,;:]+\s*$\n?", "", org_text)
    org_text = re.sub(r"(?im)^\s*P\d+\\?\}\s*$\n?", "", org_text)
    org_text = re.sub(r"(?im)^\s*<empty citation>\s*$\n?", "", org_text)
    org_text = re.sub(r"\n{3,}", "\n\n", org_text)
    org_text = re.sub(r"[ \t]+\n", "\n", org_text)
    org_text = re.sub(r"(?ms)^\*\s+EP[ÍI]GRAFE\n(?:(?!^\* ).)*?(?:Autor da ep[ií]grafe|\[\.\.\.\]|“\[\.\.\.\]”).*?(?=^\* |\Z)", "", org_text)
    org_text = re.sub(r"(?ms)^\*\s+AP[ÊE]NDICE\s+A\s+[—-]\s+T[ÍI]TULO DO AP[ÊE]NDICE.*?(?=^\* |\Z)", "", org_text)
    org_text = re.sub(r"(?ms)^\*\s+ANEXO\s+A\s+[—-]\s+T[ÍI]TULO DO ANEXO.*?(?=^\* |\Z)", "", org_text)
    return org_text.strip() + "\n"

def strip_org_cite_directives(doc: str) -> str:
    banned = ("#+cite_export:", "#+bibliography:", "#+print_bibliography:")
    kept = []
    for line in doc.splitlines():
        if line.strip().lower().startswith(banned):
            continue
        kept.append(line)
    return "\n".join(kept).strip() + "\n"


def org_uses_citation_pipeline(doc: str) -> bool:
    """Detecta se o documento precisa de BibLaTeX/Biber.

    Versões anteriores só reconheciam [cite:...]. Isso falhava quando a IA
    gerava citações textuais como [cite/t:...] ou quando o documento era
    convertido para citações LaTeX diretas (\parencite, \textcite etc.).
    O efeito era o PDF exibir chaves BibTeX cruas e não imprimir a
    bibliografia.
    """
    if not doc or not doc.strip():
        return False
    patterns = (
        r"(?im)^\s*#\+cite_export:",
        r"(?im)^\s*#\+bibliography:",
        r"(?im)^\s*#\+print_bibliography:",
        r"\[cite(?:/[A-Za-z]+)?(?::|/)[^\]]+\]",
        r"\\(?:paren|text|auto|foot|smart|cite|citeauthor|citeyear)cite\s*\{",
        r"\\printbibliography\b",
    )
    return any(re.search(pattern, doc) for pattern in patterns)

def split_bib_entries(text: str) -> list[str]:
    entries: list[str] = []
    i = 0
    n = len(text)
    while i < n:
        at = text.find("@", i)
        if at == -1:
            break
        brace = text.find("{", at)
        if brace == -1:
            break
        depth = 0
        j = brace
        while j < n:
            ch = text[j]
            if ch == "{":
                depth += 1
            elif ch == "}":
                depth -= 1
                if depth == 0:
                    entry = text[at:j + 1].strip()
                    if entry:
                        entries.append(entry)
                    i = j + 1
                    break
            j += 1
        else:
            entry = text[at:].strip()
            if entry:
                entries.append(entry)
            break
    return entries

def bib_entry_key(entry: str) -> str | None:
    m = re.match(r"\s*@[^{]+\{\s*([^,]+)\s*,", entry, flags=re.DOTALL)
    return m.group(1).strip() if m else None

def parse_bib_entries(path: Path | None) -> tuple[list[str], list[str]]:
    if path is None or not path.exists():
        return [], []
    text = path.read_text(encoding="utf-8", errors="ignore")
    entries = split_bib_entries(text)
    keys = [k for e in entries if (k := bib_entry_key(e))]
    return entries, keys


def normalize_title_loose(text: str) -> str:
    text = (text or "").strip().lower()
    # Simplifica marcas BibTeX/LaTeX comuns antes de comparar títulos.
    text = re.sub(r"\\[a-zA-Z]+\s*\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\[a-zA-Z]+", "", text)
    text = text.replace("{", "").replace("}", "")
    repl = {
        "á": "a", "à": "a", "â": "a", "ã": "a", "é": "e", "ê": "e", "í": "i",
        "ó": "o", "ô": "o", "õ": "o", "ú": "u", "ç": "c",
    }
    for k, v in repl.items():
        text = text.replace(k, v)
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text

def normalize_doi(value: str | None) -> str:
    return (value or "").strip().lower()

def bibtex_to_plain_text(value: str) -> str:
    text = str(value or "")
    replacements = {
        r"{\'a}": "á", r"\'a": "á", r"{\`a}": "à", r"\`a": "à", r"{\^a}": "â", r"\^a": "â", r"{\~a}": "ã", r"\~a": "ã",
        r"{\'e}": "é", r"\'e": "é", r"{\^e}": "ê", r"\^e": "ê",
        r"{\'i}": "í", r"\'i": "í",
        r"{\'o}": "ó", r"\'o": "ó", r"{\^o}": "ô", r"\^o": "ô", r"{\~o}": "õ", r"\~o": "õ",
        r"{\'u}": "ú", r"\'u": "ú",
        r"{\c{c}}": "ç", r"\c{c}": "ç", r"{\c c}": "ç", r"\c c": "ç",
        r"{\~n}": "ñ", r"\~n": "ñ",
        r"{\"u}": "ü", r"\"u": "ü",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    # Remove comandos TeX simples preservando seus argumentos textuais quando possível.
    text = re.sub(r"\\[a-zA-Z]+\s*\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\[a-zA-Z]+", "", text)
    text = text.replace("{", "").replace("}", "")
    text = text.replace("~", " ")
    return re.sub(r"\s+", " ", text).strip()


def extract_bib_field(entry: str, field: str) -> str | None:
    """Extrai campo BibTeX com suporte a chaves aninhadas.

    A versão anterior usava regex não-gulosa e truncava títulos como
    `Pol{\'i}tica...`, o que impedia mapear nomes de PDF para chaves reais.
    """
    m = re.search(rf"(?is)\b{re.escape(field)}\s*=\s*", entry)
    if not m:
        return None
    i = m.end()
    n = len(entry)
    while i < n and entry[i].isspace():
        i += 1
    if i >= n:
        return None
    if entry[i] == "{":
        depth = 0
        start = i + 1
        j = i
        while j < n:
            ch = entry[j]
            if ch == "{":
                depth += 1
            elif ch == "}":
                depth -= 1
                if depth == 0:
                    return re.sub(r"\s+", " ", entry[start:j]).strip()
            j += 1
        return re.sub(r"\s+", " ", entry[start:]).strip()
    if entry[i] == '"':
        start = i + 1
        j = start
        escaped = False
        while j < n:
            ch = entry[j]
            if ch == '"' and not escaped:
                return re.sub(r"\s+", " ", entry[start:j]).strip()
            escaped = (ch == "\\" and not escaped)
            if ch != "\\":
                escaped = False
            j += 1
        return re.sub(r"\s+", " ", entry[start:]).strip()
    # Valor sem aspas/chaves, até vírgula ou fim de entrada.
    j = i
    while j < n and entry[j] not in ",\n}":
        j += 1
    return re.sub(r"\s+", " ", entry[i:j]).strip() or None

def parse_bib_entry_meta(entry: str) -> dict[str, Any]:
    title = extract_bib_field(entry, "title")
    author = extract_bib_field(entry, "author")
    return {
        "key": bib_entry_key(entry),
        "title": bibtex_to_plain_text(title or "") if title else None,
        "doi": extract_bib_field(entry, "doi"),
        "url": extract_bib_field(entry, "url"),
        "year": extract_bib_field(entry, "year"),
        "author": bibtex_to_plain_text(author or "") if author else None,
    }


def _first_author_identity(author_field: str | None) -> str:
    text = bibtex_to_plain_text(author_field or "")
    if not text:
        return ""
    first = re.split(r"\s+and\s+|;", text, maxsplit=1, flags=re.IGNORECASE)[0].strip()
    if not first:
        return ""
    # BibTeX pode vir como "Sobrenome, Nome" ou "Nome Sobrenome".
    if "," in first:
        surname = first.split(",", 1)[0].strip()
    else:
        parts = first.split()
        surname = parts[-1] if parts else first
    return normalize_title_loose(surname)


def _bib_identity_key(entry: str) -> str:
    """Identidade bibliográfica robusta para deduplicação.

    Prioridade:
    1. DOI normalizado;
    2. título normalizado + ano + primeiro autor;
    3. título normalizado, quando suficientemente específico.

    Isso evita que o enriquecimento por múltiplas bases gere Arretche 2018a,
    2018b, 2018c etc. para a mesma obra.
    """
    doi = normalize_doi(extract_bib_field(entry, "doi"))
    if doi:
        doi = re.sub(r"^https?://(?:dx\.)?doi\.org/", "", doi, flags=re.IGNORECASE)
        doi = doi.strip().rstrip(".")
        return "doi:" + doi
    title = normalize_title_loose(bibtex_to_plain_text(extract_bib_field(entry, "title") or ""))
    year = str(extract_bib_field(entry, "year") or "").strip().lower()
    author = _first_author_identity(extract_bib_field(entry, "author") or extract_bib_field(entry, "editor"))
    if title and len(title) >= 20:
        if year or author:
            return "tay:" + "|".join([title, year, author])
        return "title:" + title
    key = bib_entry_key(entry) or ""
    return "key:" + normalize_title_loose(key)


def _bib_quality_score(entry: str) -> tuple[int, int, int]:
    """Pontua entrada BibLaTeX para escolher a melhor entre duplicatas."""
    score = 0
    if extract_bib_field(entry, "doi"):
        score += 50
    if extract_bib_field(entry, "author") or extract_bib_field(entry, "editor"):
        score += 20
    if extract_bib_field(entry, "year"):
        score += 10
    if extract_bib_field(entry, "journaltitle") or extract_bib_field(entry, "journal"):
        score += 10
    if extract_bib_field(entry, "publisher") or extract_bib_field(entry, "booktitle"):
        score += 8
    if extract_bib_field(entry, "pages"):
        score += 5
    note = (extract_bib_field(entry, "note") or "").lower()
    if "metadados" in note or "inferid" in note or "revisar manualmente" in note:
        score -= 25
    author = (extract_bib_field(entry, "author") or "").lower()
    if "fornecido pelo professor" in author or "material fornecido" in author:
        score -= 40
    return (score, len(entry), -len(bib_entry_key(entry) or ""))


def _replace_bib_entry_key(entry: str, new_key: str) -> str:
    old = bib_entry_key(entry)
    if not old or old == new_key:
        return entry
    return re.sub(r"(\s*@[^{]+\{\s*)[^,]+(\s*,)", lambda m: f"{m.group(1)}{new_key}{m.group(2)}", entry, count=1, flags=re.DOTALL)


def deduplicate_bib_entries(entries: list[str], *, preferred_keys: list[str] | None = None) -> tuple[list[str], list[str], dict[str, str], dict[str, Any]]:
    """Deduplica entradas BibLaTeX e devolve mapa old_key -> canonical_key.

    A deduplicação é feita antes da geração do documento para que o prompt da
    IA só veja chaves canônicas e não produza citações como 2018a/2018b para a
    mesma obra.
    """
    preferred = {k: i for i, k in enumerate(preferred_keys or [])}
    groups: dict[str, list[str]] = {}
    for entry in entries:
        key = bib_entry_key(entry)
        if not key:
            continue
        groups.setdefault(_bib_identity_key(entry), []).append(entry)

    alias: dict[str, str] = {}
    deduped: list[str] = []
    duplicate_groups: list[dict[str, Any]] = []
    used_keys: set[str] = set()

    for identity, group in groups.items():
        if len(group) == 1:
            entry = group[0]
            key = bib_entry_key(entry) or "ref"
            if key in used_keys:
                key = unique_key(key, used_keys)
                entry = _replace_bib_entry_key(entry, key)
            else:
                used_keys.add(key)
            deduped.append(entry)
            alias[key] = key
            continue

        def rank(entry: str) -> tuple[int, int, int, int]:
            key = bib_entry_key(entry) or ""
            pref = -preferred.get(key, 10_000)
            q = _bib_quality_score(entry)
            return (pref, *q)

        canonical_entry = max(group, key=rank)
        canonical_key = bib_entry_key(canonical_entry) or "ref"
        if canonical_key in used_keys:
            new_key = unique_key(canonical_key, used_keys)
            canonical_entry = _replace_bib_entry_key(canonical_entry, new_key)
            canonical_key = new_key
        else:
            used_keys.add(canonical_key)
        deduped.append(canonical_entry)
        old_keys = [bib_entry_key(e) for e in group if bib_entry_key(e)]
        for old in old_keys:
            alias[old] = canonical_key
        duplicate_groups.append({
            "identity": identity,
            "canonical_key": canonical_key,
            "merged_keys": old_keys,
            "count": len(old_keys),
            "title": parse_bib_entry_meta(canonical_entry).get("title"),
            "doi": parse_bib_entry_meta(canonical_entry).get("doi"),
        })

    # Preserva ordem de entrada original de forma aproximada: deduped foi
    # montado na ordem dos grupos vistos no dicionário, que segue a ordem de
    # inserção em Python 3.7+.
    keys = [k for e in deduped if (k := bib_entry_key(e))]
    report = {
        "before_count": len([e for e in entries if bib_entry_key(e)]),
        "after_count": len(deduped),
        "duplicate_groups_count": len(duplicate_groups),
        "duplicate_groups": duplicate_groups,
        "alias_map": alias,
    }
    return deduped, keys, alias, report


def apply_bib_key_aliases(cfg: dict[str, Any], research_paths: ResearchPaths, alias: dict[str, str]) -> None:
    """Atualiza mapas internos para chaves canônicas após deduplicação."""
    if not alias:
        return
    for entry in research_paths.selected_entries:
        key = entry.get("bib_key")
        if key in alias:
            entry["bib_key"] = alias[key]
    for cfg_key in ("__local_revised_bib_keys__", "__citation_target_keys__"):
        raw = cfg.get(cfg_key)
        if isinstance(raw, list):
            cfg[cfg_key] = list(dict.fromkeys(alias.get(k, k) for k in raw if k))
    for map_key in ("__local_revised_bib_key_by_path__", "__local_revised_bib_key_by_name__"):
        mp = cfg.get(map_key)
        if isinstance(mp, dict):
            cfg[map_key] = {k: alias.get(v, v) for k, v in mp.items()}


def replace_org_citation_keys(org_text: str, alias: dict[str, str]) -> str:
    """Reescreve [cite:@old] para [cite:@canonical] quando necessário."""
    if not alias:
        return org_text
    def repl(match: re.Match[str]) -> str:
        key = match.group(1)
        return "@" + alias.get(key, key)
    return re.sub(r"@([A-Za-z0-9_:\-]+)", repl, org_text)

def match_bib_key_for_selected(entry: dict[str, Any], bib_entries: list[str]) -> str | None:
    doi = normalize_doi(entry.get("doi"))
    title_norm = normalize_title_loose(str(entry.get("title") or ""))
    for bib in bib_entries:
        meta = parse_bib_entry_meta(bib)
        if not meta.get("key"):
            continue
        if doi and normalize_doi(meta.get("doi")) == doi:
            return str(meta["key"])
    for bib in bib_entries:
        meta = parse_bib_entry_meta(bib)
        if not meta.get("key"):
            continue
        bib_title_norm = normalize_title_loose(str(meta.get("title") or ""))
        if title_norm and bib_title_norm and (title_norm == bib_title_norm or title_norm in bib_title_norm or bib_title_norm in title_norm):
            return str(meta["key"])
    return None

def assign_bib_keys_to_selected_docs(base_docs: list[SourceDoc], research_paths: ResearchPaths, bib_entries: list[str]) -> None:
    by_path: dict[str, str] = {}
    by_title: dict[str, str] = {}
    for entry in research_paths.selected_entries:
        key = match_bib_key_for_selected(entry, bib_entries)
        if not key:
            continue
        title = str(entry.get("title") or "")
        if title:
            by_title[normalize_title_loose(title)] = key
        pdf_raw = entry.get("downloaded_pdf_path")
        if pdf_raw:
            p = Path(str(pdf_raw)).expanduser()
            try:
                p = p.resolve()
                by_path[str(p)] = key
                by_path[p.name] = key
            except Exception:
                by_path[Path(str(pdf_raw)).name] = key

    for doc in base_docs:
        if not doc.kind.startswith("texto_selecionado"):
            continue
        if doc.bib_key:
            continue
        key = None
        p = Path(doc.path)
        if p.exists():
            key = by_path.get(str(p.resolve())) or by_path.get(p.name)
        if key is None:
            key = _bib_key_from_filename(p, bib_entries)
        if key is None:
            title_norm = normalize_title_loose(doc.label)
            key = by_title.get(title_norm)
        if key is None:
            label_norm = normalize_title_loose(doc.label)
            for bib in bib_entries:
                meta = parse_bib_entry_meta(bib)
                bib_key = meta.get("key")
                bib_title_norm = normalize_title_loose(str(meta.get("title") or ""))
                if bib_key and label_norm and bib_title_norm and (label_norm in bib_title_norm or bib_title_norm in label_norm):
                    key = str(bib_key)
                    break
        if key:
            doc.bib_key = key

def extract_cited_keys_from_org(org_text: str) -> list[str]:
    keys = re.findall(r"@([A-Za-z0-9_:\-]+)", org_text)
    seen = []
    used = set()
    for k in keys:
        if k not in used:
            used.add(k)
            seen.append(k)
    return seen

def count_words(text: str) -> int:
    return len(re.findall(r"\b\w+\b", text or "", flags=re.UNICODE))

def count_org_words_per_top_section(org_text: str) -> dict[str, int]:
    sections: dict[str, int] = {}
    matches = list(re.finditer(r"(?m)^\*\s+(.+)$", org_text))
    if not matches:
        return {"documento_total": count_words(org_text)}
    for idx, m in enumerate(matches):
        title = m.group(1).strip()
        start = m.end()
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(org_text)
        body = org_text[start:end]
        sections[title] = count_words(body)
    sections["documento_total"] = count_words(org_text)
    return sections


def normalize_heading_loose(text: str) -> str:
    return normalize_title_loose(text or "")


def classify_top_section_bucket(title: str) -> str:
    t = normalize_heading_loose(title)
    if any(x in t for x in ("introducao",)):
        return "introducao"
    if any(x in t for x in ("referencial", "revisao", "fundamentacao", "marco teorico")):
        return "referencial_teorico"
    if any(x in t for x in ("metodologia", "metodo", "procedimentos", "estrategia de pesquisa")):
        return "metodologia"
    if any(x in t for x in ("resultados", "discussao", "analise", "achados", "implicacoes")):
        return "resultados_discussao"
    if any(x in t for x in ("consideracoes finais", "conclusao", "conclusoes", "sintese final", "agenda futura")):
        return "conclusao"
    return "outros"


def aggregate_section_buckets(word_counts: dict[str, int]) -> dict[str, int]:
    buckets = {
        "introducao": 0,
        "referencial_teorico": 0,
        "metodologia": 0,
        "resultados_discussao": 0,
        "conclusao": 0,
        "outros": 0,
        "documento_total": int(word_counts.get("documento_total", 0) or 0),
    }
    for title, words in word_counts.items():
        if title == "documento_total":
            continue
        bucket = classify_top_section_bucket(title)
        buckets[bucket] += int(words or 0)
    return buckets


def dissertation_generation_targets(cfg: dict[str, Any]) -> dict[str, int]:
    documento = cfg.get("documento", {})
    doc_type_local = normalize_document_type(documento.get("tipo_documento"))
    if doc_type_local != "dissertacao":
        return {
            "documento_total_min": int(documento.get("limite_palavras_total") or 6000),
            "introducao_min": int(documento.get("limite_palavras_introducao") or 800),
            "referencial_teorico_min": 1800,
            "metodologia_min": 1000,
            "resultados_discussao_min": 1800,
            "conclusao_min": int(documento.get("limite_palavras_conclusao") or 800),
        }
    return {
        "documento_total_min": int(documento.get("min_palavras_total") or documento.get("limite_palavras_total") or 14000),
        "documento_total_alvo": int(documento.get("alvo_palavras_total") or 18000),
        "introducao_min": int(documento.get("min_palavras_introducao") or documento.get("limite_palavras_introducao") or 1400),
        "referencial_teorico_min": int(documento.get("min_palavras_referencial") or documento.get("limite_palavras_revisao") or 4200),
        "metodologia_min": int(documento.get("min_palavras_metodologia") or 1800),
        "resultados_discussao_min": int(documento.get("min_palavras_resultados") or 5200),
        "conclusao_min": int(documento.get("min_palavras_conclusao") or documento.get("limite_palavras_conclusao") or 1400),
    }


def should_expand_dissertation(cfg: dict[str, Any], org_text: str, selected_keys: list[str]) -> tuple[bool, dict[str, Any]]:
    documento = cfg.get("documento", {})
    doc_type_local = normalize_document_type(documento.get("tipo_documento"))
    if doc_type_local != "dissertacao":
        return False, {}
    counts = count_org_words_per_top_section(org_text)
    buckets = aggregate_section_buckets(counts)
    targets = dissertation_generation_targets(cfg)
    cited = extract_cited_keys_from_org(org_text)
    cited_selected = [k for k in cited if k in set(selected_keys)]
    reasons: list[str] = []
    if buckets.get("documento_total", 0) < targets.get("documento_total_min", 0):
        reasons.append(f"documento_total<{targets.get('documento_total_min')}")
    for bucket_key, target_key in (
        ("introducao", "introducao_min"),
        ("referencial_teorico", "referencial_teorico_min"),
        ("metodologia", "metodologia_min"),
        ("resultados_discussao", "resultados_discussao_min"),
        ("conclusao", "conclusao_min"),
    ):
        if buckets.get(bucket_key, 0) < targets.get(target_key, 0):
            reasons.append(f"{bucket_key}<{targets.get(target_key)}")
    return bool(reasons), {
        "word_counts": counts,
        "bucket_counts": buckets,
        "targets": targets,
        "selected_cited_total": len(cited_selected),
        "selected_keys_total": len(selected_keys),
        "reasons": reasons,
    }


def build_fulltext_cache_pdf_paths(research_paths: ResearchPaths) -> list[Path]:
    """Return every PDF actually present in the research fulltext cache.

    The dissertation citation-coverage target is now the download/cache folder
    (prefixo_fulltext_cache/), not an arbitrary minimum number of selected items.
    """
    seen: set[str] = set()
    paths: list[Path] = []

    def add(path: Path | None) -> None:
        if not path:
            return
        try:
            p = Path(path).expanduser()
        except Exception:
            return
        if not p.exists() or p.suffix.lower() not in READABLE_SUFFIXES:
            return
        try:
            key = str(p.resolve())
        except Exception:
            key = str(p)
        if key not in seen:
            seen.add(key)
            paths.append(p)

    if research_paths.fulltext_cache_dir and research_paths.fulltext_cache_dir.exists():
        for pdf in sorted(p for p in research_paths.fulltext_cache_dir.iterdir() if p.is_file() and p.suffix.lower() in READABLE_SUFFIXES):
            add(pdf)
    for pdf in research_paths.fulltext_paths:
        add(pdf)
    for pdf in research_paths.selected_fulltext_paths:
        add(pdf)
    return paths


def _bib_key_from_filename(path: Path, bib_entries: list[str]) -> str | None:
    stem_norm = normalize_title_loose(path.stem)
    if not stem_norm:
        return None
    # 1) chave exatamente igual ou contida no nome do arquivo
    for entry in bib_entries:
        key = bib_entry_key(entry)
        if key and normalize_title_loose(key) == stem_norm:
            return key
    for entry in bib_entries:
        key = bib_entry_key(entry)
        if not key:
            continue
        key_norm = normalize_title_loose(key)
        if key_norm and (key_norm in stem_norm or stem_norm in key_norm):
            return key
    # 2) título BibTeX comparado ao nome do arquivo
    best_key = None
    best_score = 0.0
    stem_tokens = {t for t in stem_norm.split() if len(t) >= 4}
    for entry in bib_entries:
        meta = parse_bib_entry_meta(entry)
        key = meta.get("key")
        title_norm = normalize_title_loose(str(meta.get("title") or ""))
        if not key or not title_norm:
            continue
        if title_norm in stem_norm or stem_norm in title_norm:
            return str(key)
        title_tokens = {t for t in title_norm.split() if len(t) >= 4}
        if not title_tokens or not stem_tokens:
            continue
        inter = stem_tokens & title_tokens
        score = len(inter) / max(1, min(len(stem_tokens), len(title_tokens)))
        if score > best_score:
            best_score = score
            best_key = str(key)
    # score moderado porque nomes de arquivos podem estar truncados/sem acentos.
    if best_key and best_score >= 0.45:
        return best_key
    return None

def match_bib_key_for_fulltext_pdf(path: Path, research_paths: ResearchPaths, bib_entries: list[str]) -> str | None:
    try:
        resolved = str(path.resolve())
    except Exception:
        resolved = str(path)
    name = path.name
    stem_norm = normalize_title_loose(path.stem)

    for entry in research_paths.selected_entries:
        raw = entry.get("downloaded_pdf_path") or entry.get("pdf_path") or entry.get("local_pdf_path")
        if raw:
            try:
                p = Path(str(raw)).expanduser()
                if str(p.resolve()) == resolved or p.name == name:
                    key = match_bib_key_for_selected(entry, bib_entries)
                    if key:
                        return key
            except Exception:
                if Path(str(raw)).name == name:
                    key = match_bib_key_for_selected(entry, bib_entries)
                    if key:
                        return key

    key = _bib_key_from_filename(path, bib_entries)
    if key:
        return key

    for entry in research_paths.selected_entries:
        title = str(entry.get("title") or "")
        title_norm = normalize_title_loose(title)
        if title_norm and stem_norm and (title_norm in stem_norm or stem_norm in title_norm):
            key = match_bib_key_for_selected(entry, bib_entries)
            if key:
                return key

    for bib in bib_entries:
        meta = parse_bib_entry_meta(bib)
        key = meta.get("key")
        title_norm = normalize_title_loose(str(meta.get("title") or ""))
        if key and title_norm and stem_norm and (title_norm in stem_norm or stem_norm in title_norm):
            return str(key)
    return None


def build_fulltext_cache_citation_keys(research_paths: ResearchPaths, bib_entries: list[str]) -> tuple[list[str], list[dict[str, Any]]]:
    keys: list[str] = []
    unresolved: list[dict[str, Any]] = []
    seen: set[str] = set()
    for pdf in build_fulltext_cache_pdf_paths(research_paths):
        key = match_bib_key_for_fulltext_pdf(pdf, research_paths, bib_entries)
        if key and key not in seen:
            seen.add(key)
            keys.append(key)
        elif not key:
            unresolved.append({"path": str(pdf), "filename": pdf.name})
    return keys, unresolved


def build_selected_corpus_catalog(research_paths: ResearchPaths, bib_entries: list[str], limit: int = 40) -> list[dict[str, Any]]:
    catalog: list[dict[str, Any]] = []
    seen_keys: set[str] = set()
    for pdf in build_fulltext_cache_pdf_paths(research_paths):
        key = match_bib_key_for_fulltext_pdf(pdf, research_paths, bib_entries)
        if key:
            seen_keys.add(key)
        meta = None
        if key:
            for bib in bib_entries:
                if bib_entry_key(bib) == key:
                    meta = parse_bib_entry_meta(bib)
                    break
        catalog.append({
            "title": (meta or {}).get("title") or pdf.stem,
            "year": (meta or {}).get("year"),
            "venue": None,
            "doi": (meta or {}).get("doi"),
            "bib_key": key,
            "keywords": [],
            "paper_id": None,
            "has_pdf": True,
            "pdf_path": str(pdf),
            "source": "fulltext_cache",
            "abstract_excerpt": "",
        })
    for entry in research_paths.selected_entries:
        key = match_bib_key_for_selected(entry, bib_entries)
        if key and key in seen_keys:
            continue
        catalog.append({
            "title": entry.get("title"),
            "year": entry.get("year"),
            "venue": entry.get("venue"),
            "doi": entry.get("doi"),
            "bib_key": key,
            "keywords": entry.get("keywords") or [],
            "paper_id": entry.get("paper_id"),
            "has_pdf": bool(entry.get("downloaded_pdf_path") or entry.get("pdf_url")),
            "source": "selected_all",
            "abstract_excerpt": shorten_text(str(entry.get("abstract") or entry.get("tldr") or ""), 900),
        })
    return catalog[:limit]

def build_reference_usage_map(org_text: str, base_docs: list[SourceDoc], extra_docs: list[SourceDoc], bib_keys: list[str], required_keys: list[str] | None = None) -> dict[str, Any]:
    cited_keys = extract_cited_keys_from_org(org_text)
    selected_keys = sorted({d.bib_key for d in base_docs if d.kind.startswith("texto_selecionado") and d.bib_key})
    extra_keys = sorted({d.bib_key for d in extra_docs if d.bib_key})
    required_keys = sorted({k for k in (required_keys or selected_keys) if k})
    selected_cited = [k for k in cited_keys if k in selected_keys]
    extra_cited = [k for k in cited_keys if k in extra_keys]
    required_cited = [k for k in cited_keys if k in required_keys]
    required_not_cited = [k for k in required_keys if k not in cited_keys]
    cited_known = set(selected_keys) | set(extra_keys) | set(required_keys)
    other_cited = [k for k in cited_keys if k not in cited_known]
    bib_not_cited = [k for k in bib_keys if k not in cited_keys]

    return {
        "cited_keys": cited_keys,
        "selected_keys": selected_keys,
        "extra_keys": extra_keys,
        "required_fulltext_cache_keys": required_keys,
        "selected_cited_keys": selected_cited,
        "extra_cited_keys": extra_cited,
        "required_fulltext_cache_cited_keys": required_cited,
        "required_fulltext_cache_not_cited_keys": required_not_cited,
        "other_cited_keys": other_cited,
        "bib_keys_not_cited": bib_not_cited,
        "counts": {
            "cited_total": len(cited_keys),
            "selected_total": len(selected_keys),
            "selected_cited_total": len(selected_cited),
            "extra_total": len(extra_keys),
            "extra_cited_total": len(extra_cited),
            "required_fulltext_cache_total": len(required_keys),
            "required_fulltext_cache_cited_total": len(required_cited),
            "required_fulltext_cache_not_cited_total": len(required_not_cited),
            "bib_not_cited_total": len(bib_not_cited),
        },
    }

def render_reference_usage_markdown(usage: dict[str, Any]) -> str:
    lines = ["# Mapa de uso das referências", ""]
    counts = usage.get("counts", {})
    lines += [
        f"- Citações únicas no documento: {counts.get('cited_total', 0)}",
        f"- Chaves de artigos selecionados: {counts.get('selected_total', 0)}",
        f"- Chaves de selecionados efetivamente citadas: {counts.get('selected_cited_total', 0)}",
        f"- Chaves obrigatórias do fulltext_cache: {counts.get('required_fulltext_cache_total', 0)}",
        f"- Chaves do fulltext_cache efetivamente citadas: {counts.get('required_fulltext_cache_cited_total', 0)}",
        f"- Chaves do fulltext_cache não citadas: {counts.get('required_fulltext_cache_not_cited_total', 0)}",
        f"- Chaves de artigos extras: {counts.get('extra_total', 0)}",
        f"- Chaves de extras efetivamente citadas: {counts.get('extra_cited_total', 0)}",
        f"- Entradas do .bib não citadas: {counts.get('bib_not_cited_total', 0)}",
        "",
        "## Selecionados citados",
    ]
    sel = usage.get("selected_cited_keys", [])
    lines += [f"- `{k}`" for k in sel] if sel else ["- Nenhum"]
    req = usage.get("required_fulltext_cache_not_cited_keys", [])
    lines += ["", "## Fulltext cache obrigatório não citado"]
    lines += [f"- `{k}`" for k in req] if req else ["- Nenhum"]
    lines += ["", "## Extras citados"]
    ex = usage.get("extra_cited_keys", [])
    lines += [f"- `{k}`" for k in ex] if ex else ["- Nenhum"]
    lines += ["", "## Outras chaves citadas"]
    other = usage.get("other_cited_keys", [])
    lines += [f"- `{k}`" for k in other] if other else ["- Nenhuma"]
    lines += ["", "## Entradas do .bib não citadas"]
    notc = usage.get("bib_keys_not_cited", [])
    lines += [f"- `{k}`" for k in notc] if notc else ["- Nenhuma"]
    lines += [""]
    return "\n".join(lines)

def make_bib_key(authors: list[str], year: str | None, title: str) -> str:
    surname = slugify((authors[0].split()[-1] if authors else "anon"))
    year_part = re.search(r"(19|20)\d{2}", year or "")
    year_str = year_part.group(0) if year_part else "sd"
    title_words = [w for w in slugify(title).split("_") if w][:2]
    return f"{surname}{year_str}_{'_'.join(title_words) if title_words else 'trabalho'}"

def unique_key(candidate: str, used: set[str]) -> str:
    key = candidate
    idx = 2
    while key in used:
        key = f"{candidate}_{idx}"
        idx += 1
    used.add(key)
    return key

def render_biblatex_entry(key: str, meta: InferredBibMetadata) -> str:
    """Renderiza uma entrada BibLaTeX robusta a partir de metadados inferidos.

    O .bib não é "APA" ou "ABNT" em si: ele precisa ter campos corretos
    para que o estilo escolhido em #+CITE_EXPORT / biblatex renderize a
    bibliografia final no padrão desejado.
    """
    allowed_types = {"article", "book", "incollection", "inbook", "report", "thesis", "misc"}
    entry_type = (meta.entry_type or "misc").strip().lower()
    if entry_type not in allowed_types:
        entry_type = "misc"
    fields: list[tuple[str, str | None]] = []
    if meta.authors:
        fields.append(("author", " and ".join(a for a in meta.authors if str(a).strip())))
    if meta.editors:
        fields.append(("editor", " and ".join(e for e in meta.editors if str(e).strip())))
    fields.extend([
        ("title", meta.title),
        ("booktitle", meta.booktitle),
        ("journaltitle", meta.journaltitle),
        ("publisher", meta.publisher),
        ("location", meta.location),
        ("year", meta.year),
        ("edition", meta.edition),
        ("volume", meta.volume),
        ("number", meta.number),
        ("pages", meta.pages),
        ("doi", meta.doi),
        ("isbn", meta.isbn),
        ("url", meta.url),
        ("note", meta.note),
    ])
    body_parts: list[str] = []
    for field_name, value in fields:
        clean = str(value or "").strip()
        if not clean:
            continue
        body_parts.append(f"{field_name} = {{{bibtex_escape(clean)}}}")
    body = ",\n  ".join(body_parts)
    return f"@{entry_type}{{{key},\n  {body}\n}}"


# ---------------------------------------------------------------------------
# Enriquecimento bibliográfico de documentos locais por bases externas
# ---------------------------------------------------------------------------
def local_metadata_enrichment_enabled(cfg: dict[str, Any]) -> bool:
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    return bool(local.get("enriquecer_metadados_buscadores", local.get("enriquecer_metadados_online", False)))


def local_metadata_sources(cfg: dict[str, Any]) -> list[str]:
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    raw = local.get("fontes_metadados") or local.get("metadata_sources") or ["crossref", "openalex", "semantic_scholar"]
    sources = [s.strip().lower().replace("-", "_") for s in _ensure_list_of_strings(raw)]
    allowed = {"crossref", "openalex", "semantic_scholar", "semanticscholar", "scopus"}
    normalized: list[str] = []
    for s in sources:
        if s == "semanticscholar":
            s = "semantic_scholar"
        if s in allowed and s not in normalized:
            normalized.append(s)
    return normalized or ["crossref", "openalex", "semantic_scholar"]


def include_inferred_metadata_notes(cfg: dict[str, Any]) -> bool:
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    return bool(local.get("incluir_notas_metadados_inferidos", False))


def bibliographic_metadata_incomplete(meta: InferredBibMetadata | None, cfg: dict[str, Any] | None = None) -> bool:
    if meta is None:
        return True
    title = normalize_title_loose(meta.title or "")
    author_blob = normalize_title_loose(" ".join(meta.authors or []) + " " + " ".join(meta.editors or []))
    weak_authors = not (meta.authors or meta.editors) or any(x in author_blob for x in (
        "material fornecido", "fornecido professor", "bibliografia disciplina", "autor desconhecido"
    ))
    weak_title = not title or len(title) < 8 or is_placeholder_value(meta.title)
    weak_year = not bool(re.search(r"(18|19|20)\d{2}", str(meta.year or "")))
    weak_source = not any([meta.journaltitle, meta.booktitle, meta.publisher, meta.doi, meta.url, meta.isbn])
    weak_type = (meta.entry_type or "").strip().lower() in {"", "misc"}
    return bool(weak_title or weak_authors or weak_year or (weak_type and weak_source))


def _http_get_json(url: str, *, headers: dict[str, str] | None = None, timeout: int = 20) -> dict[str, Any] | list[Any] | None:
    req = urllib.request.Request(url, headers=headers or {})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:  # nosec - user-controlled academic metadata lookups
            charset = resp.headers.get_content_charset() or "utf-8"
            raw = resp.read().decode(charset, errors="replace")
            return json.loads(raw)
    except Exception as exc:
        debug_print(f"Falha em consulta de metadados: {url} -> {exc}")
        return None


def _clean_doi(value: str | None) -> str:
    doi = normalize_doi(value)
    doi = re.sub(r"^https?://(?:dx\.)?doi\.org/", "", doi)
    doi = doi.strip().strip(".").strip()
    return doi


def _extract_doi_from_text(text: str) -> str | None:
    m = re.search(r"10\.\d{4,9}/[-._;()/:A-Za-z0-9]+", text or "")
    if not m:
        return None
    return _clean_doi(m.group(0).rstrip(".,);]"))


def _doi_lookup_keys_for_path(path: str | Path) -> list[str]:
    """Gera chaves de procura para DOI manual por caminho/nome/stem.

    Permite que o manifesto CSV/TOML use tanto o nome exato do arquivo quanto
    caminhos relativos ou apenas o stem sem extensão.
    """
    p = Path(str(path))
    keys: list[str] = []
    raw_values = [str(path), p.name, p.stem]
    try:
        raw_values.append(str(p.resolve()))
    except Exception:
        pass
    for raw in raw_values:
        raw = raw.replace("\\", "/").strip()
        for value in (raw, raw.lower(), normalize_title_loose(raw)):
            if value and value not in keys:
                keys.append(value)
    return keys


def _add_doi_manifest_item(result: dict[str, str], filename: str, doi: str | None) -> None:
    clean_doi = _clean_doi(doi)
    if not filename or not clean_doi:
        return
    for key in _doi_lookup_keys_for_path(filename):
        result[key] = clean_doi


def load_local_doi_manifest(cfg: dict[str, Any]) -> dict[str, str]:
    """Lê DOI manual de [documentos_locais].doi_manifest_path e doi_map.

    Formatos aceitos:
      CSV: colunas arquivo/filename/path e doi
      TOML:
        [[documentos_locais.doi_map]]
        arquivo = "falleti_2005.pdf"
        doi = "10.1017/S0003055405051695"
    """
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    result: dict[str, str] = {}

    raw_manifest = local.get("doi_manifest_path") or local.get("doi_csv_path") or local.get("manifesto_doi_path")
    if raw_manifest:
        manifest_path = resolve_configured_path(raw_manifest, cfg)
        if manifest_path and manifest_path.exists() and manifest_path.is_file():
            suffix = manifest_path.suffix.lower()
            try:
                if suffix == ".csv":
                    with manifest_path.open("r", encoding="utf-8-sig", newline="") as f:
                        reader = csv.DictReader(f)
                        for row in reader:
                            filename = str(row.get("arquivo") or row.get("filename") or row.get("file") or row.get("path") or row.get("caminho") or "").strip()
                            doi = str(row.get("doi") or row.get("DOI") or "").strip()
                            _add_doi_manifest_item(result, filename, doi)
                elif suffix == ".json":
                    data = json.loads(manifest_path.read_text(encoding="utf-8"))
                    items = data if isinstance(data, list) else data.get("doi_map", []) if isinstance(data, dict) else []
                    for item in items:
                        if isinstance(item, dict):
                            filename = str(item.get("arquivo") or item.get("filename") or item.get("file") or item.get("path") or item.get("caminho") or "").strip()
                            doi = str(item.get("doi") or item.get("DOI") or "").strip()
                            _add_doi_manifest_item(result, filename, doi)
                else:
                    # TSV ou texto simples: arquivo,doi por linha.
                    for line in manifest_path.read_text(encoding="utf-8", errors="ignore").splitlines():
                        if not line.strip() or line.lower().startswith("arquivo"):
                            continue
                        parts = [p.strip() for p in re.split(r"[,;\t]", line) if p.strip()]
                        if len(parts) >= 2:
                            _add_doi_manifest_item(result, parts[0], parts[1])
            except Exception as exc:
                debug_print(f"Falha ao ler manifesto DOI {manifest_path}: {exc}")

    raw_map = local.get("doi_map") or local.get("mapa_doi") or []
    if isinstance(raw_map, list):
        for item in raw_map:
            if isinstance(item, dict):
                filename = str(item.get("arquivo") or item.get("filename") or item.get("file") or item.get("path") or item.get("caminho") or "").strip()
                doi = str(item.get("doi") or item.get("DOI") or "").strip()
                _add_doi_manifest_item(result, filename, doi)
    elif isinstance(raw_map, dict):
        for filename, doi in raw_map.items():
            _add_doi_manifest_item(result, str(filename), str(doi))

    return result


def doi_from_manifest_for_doc(doc: SourceDoc, doi_manifest: dict[str, str]) -> str | None:
    if not doi_manifest:
        return None
    for key in _doi_lookup_keys_for_path(doc.path):
        if key in doi_manifest:
            return _clean_doi(doi_manifest[key])
    return None


def known_doi_for_doc(doc: SourceDoc, cfg: dict[str, Any], doi_manifest: dict[str, str] | None = None) -> tuple[str | None, str | None]:
    """Retorna DOI conhecido e origem: manual_csv_toml, extraido_texto ou None."""
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    prefer_manual = bool(local.get("preferir_doi_manual", True))
    extract_auto = bool(local.get("extrair_doi_dos_pdfs", local.get("extrair_doi_dos_documentos", True)))
    manifest = doi_manifest or {}

    manual = doi_from_manifest_for_doc(doc, manifest)
    if prefer_manual and manual:
        return manual, "manual_csv_toml"
    if extract_auto:
        extracted = _extract_doi_from_text(doc.extracted_text)
        if extracted:
            return extracted, "extraido_texto"
    if manual:
        return manual, "manual_csv_toml"
    return None, None


def _first_year(*values: Any) -> str | None:
    for value in values:
        m = re.search(r"(18|19|20)\d{2}", str(value or ""))
        if m:
            return m.group(0)
    return None


def _author_list_from_crossref(authors: list[dict[str, Any]] | None) -> list[str]:
    out: list[str] = []
    for a in authors or []:
        given = str(a.get("given") or "").strip()
        family = str(a.get("family") or "").strip()
        name = " ".join(p for p in (given, family) if p).strip()
        if name:
            out.append(name)
    return out


def _author_list_from_openalex(authorships: list[dict[str, Any]] | None) -> list[str]:
    out: list[str] = []
    for item in authorships or []:
        author = item.get("author") if isinstance(item, dict) else None
        name = str((author or {}).get("display_name") or "").strip()
        if name:
            out.append(name)
    return out


def _author_list_from_s2(authors: list[dict[str, Any]] | None) -> list[str]:
    out: list[str] = []
    for a in authors or []:
        name = str(a.get("name") or "").strip()
        if name:
            out.append(name)
    return out


def _candidate_to_meta(candidate: dict[str, Any]) -> InferredBibMetadata:
    return InferredBibMetadata(
        entry_type=str(candidate.get("entry_type") or "article"),
        title=str(candidate.get("title") or "").strip(),
        authors=[str(a).strip() for a in candidate.get("authors") or [] if str(a).strip()],
        editors=[str(a).strip() for a in candidate.get("editors") or [] if str(a).strip()],
        year=str(candidate.get("year") or "").strip() or None,
        booktitle=candidate.get("booktitle"),
        journaltitle=candidate.get("journaltitle"),
        publisher=candidate.get("publisher"),
        location=candidate.get("location"),
        pages=candidate.get("pages"),
        volume=candidate.get("volume"),
        number=candidate.get("number"),
        edition=candidate.get("edition"),
        doi=_clean_doi(candidate.get("doi")) or None,
        isbn=candidate.get("isbn"),
        url=candidate.get("url"),
        note=None,
    )


def _crossref_item_to_candidate(item: dict[str, Any]) -> dict[str, Any]:
    title = " ".join(item.get("title") or []).strip()
    container = " ".join(item.get("container-title") or []).strip()
    issued = item.get("issued") or item.get("published-print") or item.get("published-online") or {}
    date_parts = issued.get("date-parts") or []
    year = None
    if date_parts and date_parts[0]:
        year = str(date_parts[0][0])
    pages = item.get("page")
    typ = str(item.get("type") or "journal-article")
    entry_type = "article" if "journal" in typ else ("book" if typ == "book" else ("incollection" if "book-chapter" in typ or "chapter" in typ else "misc"))
    return {
        "source": "crossref",
        "entry_type": entry_type,
        "title": title,
        "authors": _author_list_from_crossref(item.get("author")),
        "editors": _author_list_from_crossref(item.get("editor")),
        "year": year,
        "journaltitle": container if entry_type == "article" else None,
        "booktitle": container if entry_type in {"incollection", "inbook"} else None,
        "publisher": item.get("publisher"),
        "pages": pages,
        "volume": item.get("volume"),
        "number": item.get("issue"),
        "doi": item.get("DOI"),
        "isbn": (item.get("ISBN") or [None])[0] if isinstance(item.get("ISBN"), list) else item.get("ISBN"),
        "url": item.get("URL"),
    }


def query_crossref_metadata(meta: InferredBibMetadata, doc: SourceDoc, cfg: dict[str, Any]) -> list[dict[str, Any]]:
    mailto = os.getenv("CROSSREF_MAILTO") or os.getenv("OPENALEX_MAILTO") or ""
    headers = {"User-Agent": f"academic-pipeline/0.2.6 ({mailto or 'local-use'})"}
    candidates: list[dict[str, Any]] = []
    doi = _clean_doi(meta.doi) or _extract_doi_from_text(doc.extracted_text)
    if doi:
        url = "https://api.crossref.org/works/" + urllib.parse.quote(doi, safe="")
        data = _http_get_json(url, headers=headers)
        if isinstance(data, dict) and isinstance(data.get("message"), dict):
            candidates.append(_crossref_item_to_candidate(data["message"]))
            return candidates
    query = str(meta.title or "").strip() or Path(doc.path).stem.replace("_", " ").replace("-", " ")
    if not query:
        return candidates
    params = {"query.title": query, "rows": "5"}
    if mailto:
        params["mailto"] = mailto
    url = "https://api.crossref.org/works?" + urllib.parse.urlencode(params)
    data = _http_get_json(url, headers=headers)
    items = (((data or {}).get("message") or {}).get("items") or []) if isinstance(data, dict) else []
    for item in items[:5]:
        if isinstance(item, dict):
            candidates.append(_crossref_item_to_candidate(item))
    return candidates


def _openalex_item_to_candidate(item: dict[str, Any]) -> dict[str, Any]:
    primary_location = item.get("primary_location") or {}
    source = primary_location.get("source") or {}
    biblio = item.get("biblio") or {}
    doi = item.get("doi")
    if doi:
        doi = _clean_doi(str(doi))
    typ = str(item.get("type") or "article")
    entry_type = "article" if typ in {"article", "journal-article"} else ("book" if typ == "book" else ("incollection" if typ in {"book-chapter", "chapter"} else "misc"))
    return {
        "source": "openalex",
        "entry_type": entry_type,
        "title": item.get("display_name"),
        "authors": _author_list_from_openalex(item.get("authorships")),
        "year": str(item.get("publication_year") or "") or None,
        "journaltitle": source.get("display_name") if entry_type == "article" else None,
        "booktitle": source.get("display_name") if entry_type in {"incollection", "inbook"} else None,
        "publisher": source.get("host_organization_name") or item.get("publisher"),
        "pages": biblio.get("first_page") + "--" + biblio.get("last_page") if biblio.get("first_page") and biblio.get("last_page") else None,
        "volume": biblio.get("volume"),
        "number": biblio.get("issue"),
        "doi": doi,
        "url": item.get("landing_page_url") or item.get("id"),
    }


def query_openalex_metadata(meta: InferredBibMetadata, doc: SourceDoc, cfg: dict[str, Any]) -> list[dict[str, Any]]:
    mailto = os.getenv("OPENALEX_MAILTO") or os.getenv("CROSSREF_MAILTO") or ""
    headers = {"User-Agent": f"academic-pipeline/0.2.6 ({mailto or 'local-use'})"}
    candidates: list[dict[str, Any]] = []
    doi = _clean_doi(meta.doi) or _extract_doi_from_text(doc.extracted_text)
    if doi:
        url = "https://api.openalex.org/works/doi:" + urllib.parse.quote(doi, safe="")
        if mailto:
            url += "?" + urllib.parse.urlencode({"mailto": mailto})
        data = _http_get_json(url, headers=headers)
        if isinstance(data, dict) and data.get("display_name"):
            candidates.append(_openalex_item_to_candidate(data))
            return candidates
    query = str(meta.title or "").strip() or Path(doc.path).stem.replace("_", " ").replace("-", " ")
    if not query:
        return candidates
    params = {"search": query, "per-page": "5"}
    if mailto:
        params["mailto"] = mailto
    url = "https://api.openalex.org/works?" + urllib.parse.urlencode(params)
    data = _http_get_json(url, headers=headers)
    results = (data.get("results") or []) if isinstance(data, dict) else []
    for item in results[:5]:
        if isinstance(item, dict):
            candidates.append(_openalex_item_to_candidate(item))
    return candidates


def _s2_item_to_candidate(item: dict[str, Any]) -> dict[str, Any]:
    publication_venue = item.get("publicationVenue") or {}
    venue = item.get("venue") or publication_venue.get("name")
    return {
        "source": "semantic_scholar",
        "entry_type": "article",
        "title": item.get("title"),
        "authors": _author_list_from_s2(item.get("authors")),
        "year": str(item.get("year") or "") or None,
        "journaltitle": venue,
        "doi": item.get("externalIds", {}).get("DOI") if isinstance(item.get("externalIds"), dict) else None,
        "url": item.get("url"),
    }


def query_semantic_scholar_metadata(meta: InferredBibMetadata, doc: SourceDoc, cfg: dict[str, Any]) -> list[dict[str, Any]]:
    key = os.getenv("SEMANTIC_SCHOLAR_API_KEY") or os.getenv("S2_API_KEY") or ""
    headers = {"User-Agent": "academic-pipeline/0.2.6"}
    if key:
        headers["x-api-key"] = key
    candidates: list[dict[str, Any]] = []
    doi = _clean_doi(meta.doi) or _extract_doi_from_text(doc.extracted_text)
    fields = "title,authors,year,venue,externalIds,url,publicationVenue"
    if doi:
        url = "https://api.semanticscholar.org/graph/v1/paper/DOI:" + urllib.parse.quote(doi, safe="") + "?" + urllib.parse.urlencode({"fields": fields})
        data = _http_get_json(url, headers=headers)
        if isinstance(data, dict) and data.get("title"):
            candidates.append(_s2_item_to_candidate(data))
            return candidates
    query = str(meta.title or "").strip() or Path(doc.path).stem.replace("_", " ").replace("-", " ")
    if not query:
        return candidates
    url = "https://api.semanticscholar.org/graph/v1/paper/search?" + urllib.parse.urlencode({"query": query, "limit": "5", "fields": fields})
    data = _http_get_json(url, headers=headers)
    results = (data.get("data") or []) if isinstance(data, dict) else []
    for item in results[:5]:
        if isinstance(item, dict):
            candidates.append(_s2_item_to_candidate(item))
    return candidates


def _scopus_item_to_candidate(item: dict[str, Any]) -> dict[str, Any]:
    authors = []
    if item.get("dc:creator"):
        authors.append(str(item.get("dc:creator")))
    doi = item.get("prism:doi") or item.get("doi")
    typ = str(item.get("subtypeDescription") or item.get("subtype") or "article").lower()
    entry_type = "article" if "article" in typ or typ in {"ar"} else "misc"
    return {
        "source": "scopus",
        "entry_type": entry_type,
        "title": item.get("dc:title"),
        "authors": authors,
        "year": _first_year(item.get("prism:coverDate"), item.get("prism:coverDisplayDate")),
        "journaltitle": item.get("prism:publicationName"),
        "volume": item.get("prism:volume"),
        "number": item.get("prism:issueIdentifier"),
        "doi": doi,
        "url": item.get("prism:url"),
    }


def query_scopus_metadata(meta: InferredBibMetadata, doc: SourceDoc, cfg: dict[str, Any]) -> list[dict[str, Any]]:
    key = os.getenv("SCOPUS_API_KEY") or os.getenv("ELSEVIER_API_KEY") or ""
    if not key:
        return []
    headers = {"X-ELS-APIKey": key, "Accept": "application/json", "User-Agent": "academic-pipeline/0.2.6"}
    candidates: list[dict[str, Any]] = []
    doi = _clean_doi(meta.doi) or _extract_doi_from_text(doc.extracted_text)
    if doi:
        query = f"DOI({doi})"
    else:
        title = str(meta.title or "").strip() or Path(doc.path).stem.replace("_", " ").replace("-", " ")
        if not title:
            return []
        query = f'TITLE("{title[:180]}")'
    params = {"query": query, "count": "5"}
    url = "https://api.elsevier.com/content/search/scopus?" + urllib.parse.urlencode(params)
    data = _http_get_json(url, headers=headers)
    entries = (((data or {}).get("search-results") or {}).get("entry") or []) if isinstance(data, dict) else []
    for item in entries[:5]:
        if isinstance(item, dict) and not item.get("error"):
            candidates.append(_scopus_item_to_candidate(item))
    return candidates


def candidate_match_score(candidate: dict[str, Any], meta: InferredBibMetadata, doc: SourceDoc) -> float:
    cand_doi = _clean_doi(candidate.get("doi"))
    known_doi = _clean_doi(meta.doi) or _extract_doi_from_text(doc.extracted_text)
    if cand_doi and known_doi and cand_doi == known_doi:
        return 1.0
    cand_title = normalize_title_loose(str(candidate.get("title") or ""))
    local_title = normalize_title_loose(str(meta.title or "")) or normalize_title_loose(Path(doc.path).stem)
    title_score = difflib.SequenceMatcher(None, cand_title, local_title).ratio() if cand_title and local_title else 0.0
    cand_year = _first_year(candidate.get("year"))
    local_year = _first_year(meta.year, doc.extracted_text[:5000])
    year_bonus = 0.06 if cand_year and local_year and cand_year == local_year else 0.0
    cand_authors = normalize_title_loose(" ".join(candidate.get("authors") or []))
    local_authors = normalize_title_loose(" ".join(meta.authors or meta.editors or []))
    author_bonus = 0.0
    if cand_authors and local_authors:
        cand_tokens = {t for t in cand_authors.split() if len(t) >= 4}
        local_tokens = {t for t in local_authors.split() if len(t) >= 4}
        if cand_tokens and local_tokens and (cand_tokens & local_tokens):
            author_bonus = 0.06
    return min(0.99, title_score + year_bonus + author_bonus)


def enrich_metadata_from_sources(meta: InferredBibMetadata, doc: SourceDoc, cfg: dict[str, Any]) -> tuple[InferredBibMetadata | None, dict[str, Any]]:
    if not local_metadata_enrichment_enabled(cfg):
        return None, {"enabled": False}
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    min_score = float(local.get("min_score_match_metadados") or local.get("min_metadata_match_score") or 0.82)
    diagnostics: dict[str, Any] = {"enabled": True, "min_score": min_score, "sources": [], "candidates": []}
    all_candidates: list[dict[str, Any]] = []
    for source in local_metadata_sources(cfg):
        diagnostics["sources"].append(source)
        try:
            if source == "crossref":
                candidates = query_crossref_metadata(meta, doc, cfg)
            elif source == "openalex":
                candidates = query_openalex_metadata(meta, doc, cfg)
            elif source == "semantic_scholar":
                candidates = query_semantic_scholar_metadata(meta, doc, cfg)
            elif source == "scopus":
                candidates = query_scopus_metadata(meta, doc, cfg)
            else:
                candidates = []
        except Exception as exc:
            diagnostics["candidates"].append({"source": source, "error": str(exc)})
            candidates = []
        for cand in candidates:
            score = candidate_match_score(cand, meta, doc)
            cand = {**cand, "score": score}
            diagnostics["candidates"].append({k: v for k, v in cand.items() if k not in {"abstract"}})
            all_candidates.append(cand)
    if not all_candidates:
        diagnostics["selected"] = None
        return None, diagnostics
    best = sorted(all_candidates, key=lambda c: float(c.get("score") or 0), reverse=True)[0]
    diagnostics["selected"] = {k: v for k, v in best.items() if k not in {"abstract"}}
    if float(best.get("score") or 0) < min_score:
        return None, diagnostics
    return _candidate_to_meta(best), diagnostics


def apply_discreet_metadata_fallback(meta: InferredBibMetadata, doc: SourceDoc, cfg: dict[str, Any]) -> InferredBibMetadata:
    """Aplica fallback sem expor ruído técnico no paper nem no .bib."""
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    if not meta.title or is_placeholder_value(meta.title):
        meta.title = Path(doc.path).stem.replace("_", " ").replace("-", " ").strip().title() or Path(doc.path).name
    if not meta.authors and not meta.editors:
        fallback_author = str(local.get("autor_padrao") or "").strip()
        if fallback_author and normalize_title_loose(fallback_author) not in {"material fornecido pelo professor", "fornecido pelo professor"}:
            meta.authors = [fallback_author]
    if not meta.year:
        meta.year = str(local.get("ano_padrao") or "s.d.").strip() or "s.d."
    if not include_inferred_metadata_notes(cfg):
        meta.note = None
    return meta

def collect_readable_files(raw_items: list[Any], base_dir: Path | None = None) -> list[Path]:
    out: list[Path] = []
    seen: set[str] = set()
    for raw in raw_items or []:
        if not str(raw).strip():
            continue
        p = safe_resolve_user_path(raw, base_dir=base_dir)
        if p is None:
            continue
        candidates: list[Path] = []
        if p.is_dir():
            candidates = [c for c in sorted(p.rglob("*")) if c.is_file() and c.suffix.lower() in READABLE_SUFFIXES]
        elif p.is_file() and p.suffix.lower() in READABLE_SUFFIXES:
            candidates = [p]
        for c in candidates:
            key = str(c)
            if key not in seen:
                out.append(c)
                seen.add(key)
    return out

def infer_bib_metadata_for_doc(client: OpenAI, model: str, doc: SourceDoc, citation_style: str | None = None) -> InferredBibMetadata:
    style_hint = (citation_style or DEFAULT_STYLE).strip().lower() or DEFAULT_STYLE
    prompt = textwrap.dedent(
        f"""
        Extraia metadados bibliográficos do documento abaixo e retorne JSON estruturado.

        O objetivo é gerar uma entrada BibLaTeX revisada para posterior renderização no estilo {style_hint.upper()}.
        O .bib deve ser tecnicamente correto; o estilo visual será aplicado depois pelo biblatex.

        Escolha entry_type de forma criteriosa:
        - article: artigo de periódico;
        - book: livro inteiro;
        - incollection: capítulo de livro/coletânea;
        - report: relatório técnico/institucional;
        - thesis: tese/dissertação;
        - misc: somente quando os metadados forem insuficientes.

        Retorne, quando disponíveis:
        - entry_type
        - title
        - authors
        - editors
        - year
        - booktitle
        - journaltitle
        - publisher
        - location
        - pages
        - volume
        - number
        - edition
        - doi
        - isbn
        - url
        - note

        Regras de confiabilidade:
        - não invente autores, ano, DOI, páginas, editora ou periódico;
        - se não houver certeza, deixe o campo nulo;
        - para capítulo de livro, prefira entry_type = incollection e preencha booktitle/editor/publisher/pages se o texto permitir;
        - para metadados inferidos com baixa certeza, deixe note nulo; não escreva observações técnicas para aparecerem no paper;
        - não inclua comentários fora do JSON.

        Documento: {doc.label}
        Caminho: {doc.path}

        Texto extraído:
        {shorten_text(doc.extracted_text, 22000)}
        """
    ).strip()
    resp = client.responses.parse(
        model=model,
        input=[{"role": "user", "content": prompt}],
        text_format=InferredBibMetadata,
    )
    if resp.output_parsed is None:
        raise RuntimeError("IA não retornou metadados bibliográficos estruturados.")
    return resp.output_parsed

def build_bib_entries_for_extra_docs(client: OpenAI, model: str, docs: list[SourceDoc], existing_keys: list[str]) -> tuple[list[SourceDoc], list[str], list[str]]:
    used = set(existing_keys)
    added_entries: list[str] = []
    added_keys: list[str] = []

    for doc in docs:
        try:
            meta = infer_bib_metadata_for_doc(client, model, doc)
        except Exception:
            meta = InferredBibMetadata(
                entry_type="misc",
                title=Path(doc.path).stem.replace("_", " ").replace("-", " "),
                url=doc.metadata.get("url") if isinstance(doc.metadata, dict) else None,
                note="Artigo extra adicionado fora da triagem da pesquisa; revisar metadados manualmente.",
            )

        key = unique_key(make_bib_key(meta.authors, meta.year, meta.title), used)
        entry = render_biblatex_entry(key, meta)
        doc.bib_key = key
        doc.metadata = {**doc.metadata, **meta.model_dump()}
        added_entries.append(entry)
        added_keys.append(key)

    return docs, added_entries, added_keys

def find_fallback_template(doc_type: str = DEFAULT_DOC_TYPE) -> Path | None:
    doc_type = normalize_document_type(doc_type)
    preferred = DEFAULT_DISSERTATION_TEMPLATE if doc_type == "dissertacao" else DEFAULT_PAPER_TEMPLATE
    candidates = [
        Path.cwd() / "templates" / preferred,
        Path.cwd() / preferred,
        Path.cwd() / "templates" / DEFAULT_FALLBACK_TEMPLATE,
        Path.cwd() / DEFAULT_FALLBACK_TEMPLATE,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate.resolve()
    return None

def build_default_document_template(doc_type) -> str:
    class_name = "fgv-dissertacao" if normalize_document_type(doc_type) == "dissertacao" else "fgv-paper"
    return textwrap.dedent(f"""
        #+TITLE: TÍTULO A SER GERADO PELA IA
        #+AUTHOR: Gustavo M. Mendes de Tarso
        #+LANGUAGE: pt_BR
        #+OPTIONS: toc:nil num:t ^:nil
        #+LATEX_CLASS: {class_name}
        #+LATEX_HEADER: \institution{{Faculdade Getúlio Vargas}}
        #+LATEX_HEADER: \coursename{{}}
        #+LATEX_HEADER: \disciplinename{{}}
        #+LATEX_HEADER: \professorname{{}}
        #+LATEX_HEADER: \cityname{{Brasília}}
        #+LATEX_HEADER: \papertype{{Texto gerado automaticamente pela IA após a conclusão do documento}}
        #+LATEX_HEADER: \covernote{{Nota filosófica a ser gerada pela IA após a conclusão do documento}}

        #+begin_abstract
        Resumo a ser gerado pela IA.
        #+end_abstract

        * Introdução
        * Desenvolvimento
        * Considerações finais
        #+PRINT_BIBLIOGRAPHY:
        """).strip() + "\n"


# ---------------------------------------------------------------------------
# Documentos locais / ZIP como corpus de entrada
# ---------------------------------------------------------------------------
def is_local_documents_mode(cfg: dict[str, Any]) -> bool:
    pipeline = cfg.get("pipeline", {}) if isinstance(cfg.get("pipeline"), dict) else {}
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    mode = str(pipeline.get("modo_entrada") or local.get("modo_entrada") or "").strip().lower()
    return mode in {"documentos_locais", "local", "zip", "pasta_local"} or bool(local.get("ativos", False))


def get_bundle_root_from_config(cfg: dict[str, Any]) -> Path:
    cfg_dir = get_config_base_dir(cfg)
    candidates = [cfg_dir, *cfg_dir.parents]
    for cand in candidates:
        if (cand / "scripts").exists() and (cand / "templates").exists():
            return cand
        if cand.name.startswith("bundle_projeto"):
            return cand
    return cfg_dir


def local_supported_suffixes(cfg: dict[str, Any]) -> set[str]:
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    raw = local.get("tipos") or ["pdf", "docx", "txt", "md", "org"]
    values = _ensure_list_of_strings(raw)
    suffixes: set[str] = set()
    for value in values:
        v = value.strip().lower()
        if not v:
            continue
        if not v.startswith("."):
            v = "." + v
        if v in READABLE_SUFFIXES:
            suffixes.add(v)
    return suffixes or {".pdf", ".docx", ".txt", ".md", ".org"}


def safe_extract_zip(zip_path: Path, dest_dir: Path) -> list[Path]:
    dest_dir.mkdir(parents=True, exist_ok=True)
    extracted: list[Path] = []
    root = dest_dir.resolve()
    with zipfile.ZipFile(zip_path, "r") as zf:
        for info in zf.infolist():
            name = info.filename
            if not name or name.endswith("/"):
                continue
            target = (dest_dir / name).resolve()
            if not str(target).startswith(str(root) + os.sep) and target != root:
                raise RuntimeError(f"Entrada insegura no ZIP recusada: {name}")
            target.parent.mkdir(parents=True, exist_ok=True)
            with zf.open(info, "r") as src, target.open("wb") as dst:
                shutil.copyfileobj(src, dst)
            extracted.append(target)
    return extracted


def unique_copy_to_dir(src: Path, dest_dir: Path) -> Path:
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / src.name
    if not dest.exists():
        shutil.copy2(src, dest)
        return dest
    stem, suffix = src.stem, src.suffix
    idx = 2
    while True:
        candidate = dest_dir / f"{stem}_{idx}{suffix}"
        if not candidate.exists():
            shutil.copy2(src, candidate)
            return candidate
        idx += 1


def discover_local_input_files(cfg: dict[str, Any], work_root: Path) -> tuple[list[Path], dict[str, Any]]:
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    suffixes = local_supported_suffixes(cfg)
    recursive = bool(local.get("recursive", True))
    extracted_dir = work_root / "extracted"
    source_info: dict[str, Any] = {"input_zip": None, "input_dir": None, "extracted_dir": str(extracted_dir), "supported_suffixes": sorted(suffixes)}
    candidates: list[Path] = []
    input_zip = resolve_configured_path(local.get("input_zip"), cfg) if local.get("input_zip") else None
    input_dir = resolve_configured_path(local.get("input_dir"), cfg) if local.get("input_dir") else None
    if input_zip:
        if not input_zip.exists() or not input_zip.is_file():
            raise RuntimeError(f"ZIP de documentos locais não encontrado: {input_zip}")
        source_info["input_zip"] = str(input_zip)
        unzip_dir = extracted_dir / slugify(input_zip.stem)
        if unzip_dir.exists() and bool(local.get("limpar_extracao_anterior", True)):
            shutil.rmtree(unzip_dir)
        candidates.extend(p for p in safe_extract_zip(input_zip, unzip_dir) if p.is_file())
    elif input_dir:
        if not input_dir.exists() or not input_dir.is_dir():
            raise RuntimeError(f"Diretório de documentos locais não encontrado: {input_dir}")
        source_info["input_dir"] = str(input_dir)
        pattern_iter = input_dir.rglob("*") if recursive else input_dir.glob("*")
        candidates.extend(p for p in pattern_iter if p.is_file())
    else:
        raise RuntimeError("Modo documentos_locais ativo, mas nenhum input_zip ou input_dir foi informado em [documentos_locais].")
    files = sorted({p.resolve() for p in candidates if p.suffix.lower() in suffixes})
    if not files:
        raise RuntimeError("Nenhum documento local suportado foi encontrado no ZIP/diretório informado.")
    return files, source_info


def bibtex_escape(value: str) -> str:
    """Escapa valores gerados automaticamente para campos BibTeX.

    Usado principalmente no modo documentos_locais quando nÃ£o hÃ¡ .bib externo.
    Nomes de arquivos locais frequentemente trazem `_`, `%`, `&` etc.; se esses
    caracteres entrarem crus no .bbl, o LaTeX quebra no final da compilaÃ§Ã£o.
    """
    text = str(value or "")
    text = text.replace("\\", "/")
    text = text.replace("\n", " ").strip()
    replacements = {
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
        "{": r"\{",
        "}": r"\}",
    }
    return "".join(replacements.get(ch, ch) for ch in text)


def latex_escape_text(value: str) -> str:
    """Escapa texto comum que será inserido dentro de comandos LaTeX crus.

    Não use esta função em blocos LaTeX completos, apenas em valores textuais
    como títulos, campos de ficha técnica e rótulos visíveis.
    """
    text = str(value or "")
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    return "".join(replacements.get(ch, ch) for ch in text)


def pretty_title_from_prefix(prefix: str) -> str:
    text = str(prefix or "").strip()
    text = re.sub(r"[_\-]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text.title() if text else "Documento"


def configured_external_bib_path(cfg: dict[str, Any]) -> Path | None:
    """Localiza um .bib externo informado ou detectado no modo documentos_locais.

    Se [documentos_locais].auto_detect_bib nÃ£o for falso, procura tambÃ©m por
    arquivos .bib ao lado do ZIP, dentro da pasta de entrada e no diretÃ³rio do TOML.
    Isso evita cair em bibliografia sintÃ©tica quando a atividade jÃ¡ tem .bib prÃ³prio.
    """
    candidates: list[Any] = []
    for section_name in ("documentos_locais", "documento", "bibliografia"):
        section = cfg.get(section_name, {}) if isinstance(cfg.get(section_name), dict) else {}
        for key in ("bib_path", "referencias_bib", "referencias_bib_path", "arquivo_bib", "arquivo_bibliografia"):
            if section.get(key):
                candidates.append(section.get(key))
    for raw in candidates:
        path = resolve_configured_path(raw, cfg)
        if path and path.exists() and path.is_file() and path.suffix.lower() == ".bib":
            return path

    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    if not bool(local.get("auto_detect_bib", True)):
        return None

    search_roots: list[Path] = []
    for key in ("input_zip", "input_dir"):
        raw = local.get(key)
        if not raw:
            continue
        path = resolve_configured_path(raw, cfg)
        if not path:
            continue
        if path.is_file():
            search_roots.append(path.parent)
        elif path.is_dir():
            search_roots.append(path)
    cfg_dir = cfg.get("__config_dir__")
    if cfg_dir:
        try:
            search_roots.append(Path(str(cfg_dir)))
        except Exception:
            pass

    seen_roots: set[str] = set()
    bibs: list[Path] = []
    for root in search_roots:
        try:
            root = root.expanduser().resolve()
        except Exception:
            continue
        if not root.exists() or str(root) in seen_roots:
            continue
        seen_roots.add(str(root))
        bibs.extend(sorted(root.glob("*.bib")))
        bibs.extend(sorted(root.rglob("*.bib")))

    unique: list[Path] = []
    seen_files: set[str] = set()
    for b in bibs:
        try:
            rb = b.resolve()
        except Exception:
            rb = b
        if str(rb) not in seen_files and rb.exists() and rb.is_file():
            seen_files.add(str(rb))
            unique.append(rb)
    if not unique:
        return None

    def priority(path: Path) -> tuple[int, str]:
        name = normalize_title_loose(path.stem)
        score = 0
        if any(token in name for token in ("referencia", "referencias", "references", "bibliografia", "bibliography")):
            score -= 10
        return (score, str(path))

    return sorted(unique, key=priority)[0]


def build_local_bib_entry(key: str, path: Path, cfg: dict[str, Any]) -> str:
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    author = str(local.get("autor_padrao") or "Material fornecido pelo professor")
    year = str(local.get("ano_padrao") or "s.d.")
    title = path.stem.replace("_", " ").replace("-", " ").strip().title() or path.name
    return textwrap.dedent(f"""
    @misc{{{key},
      title = {{{bibtex_escape(title)}}},
      author = {{{bibtex_escape(author)}}},
      year = {{{bibtex_escape(year)}}},
      note = {{{bibtex_escape('Documento local: ' + path.name)}}}
    }}
    """).strip()




def build_revised_bib_entries_for_local_docs(
    client: OpenAI,
    model: str,
    cfg: dict[str, Any],
    docs: list[SourceDoc],
    existing_keys: list[str] | None = None,
) -> tuple[list[SourceDoc], list[str], list[str], list[dict[str, Any]]]:
    """Gera .bib revisado por IA para o corpus local.

    v0.2.4: quando [documentos_locais].enriquecer_metadados_buscadores=true,
    tenta enriquecer metadados incompletos em Crossref, OpenAlex,
    Semantic Scholar e Scopus antes de cair em fallback. Essa etapa é usada
    apenas para corrigir bibliografia dos documentos locais, sem executar PRISMA.
    """
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    bibliografia = cfg.get("bibliografia", {}) if isinstance(cfg.get("bibliografia"), dict) else {}
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    style = str(
        local.get("estilo_bib_revisado")
        or documento.get("estilo_citacao")
        or bibliografia.get("estilo_citacao")
        or DEFAULT_STYLE
    ).strip().lower()
    used = set(existing_keys or [])
    entries: list[str] = []
    keys: list[str] = []
    diagnostics: list[dict[str, Any]] = []
    doi_manifest = load_local_doi_manifest(cfg)

    for doc in docs:
        source = "ia"
        inference_error = None
        known_doi, known_doi_source = known_doi_for_doc(doc, cfg, doi_manifest)
        try:
            meta = infer_bib_metadata_for_doc(client, model, doc, citation_style=style)
        except Exception as exc:
            inference_error = str(exc)
            meta = InferredBibMetadata(
                entry_type="misc",
                title=Path(doc.path).stem.replace("_", " ").replace("-", " ").strip().title() or Path(doc.path).name,
                authors=[],
                year=str(local.get("ano_padrao") or "s.d.").strip() or "s.d.",
                note=None,
            )
            source = "fallback_local"

        if known_doi:
            meta.doi = known_doi

        enrichment_diag: dict[str, Any] | None = None
        force_doi_lookup = bool(known_doi and local.get("buscar_metadados_por_doi", True))
        if local_metadata_enrichment_enabled(cfg) and (force_doi_lookup or bibliographic_metadata_incomplete(meta, cfg)):
            enriched_meta, enrichment_diag = enrich_metadata_from_sources(meta, doc, cfg)
            if enriched_meta is not None:
                # Se o DOI foi fornecido manualmente, preserva-o mesmo que a base retorne variação.
                if known_doi and not enriched_meta.doi:
                    enriched_meta.doi = known_doi
                meta = enriched_meta
                source = f"enriquecido_buscadores:{(enrichment_diag.get('selected') or {}).get('source', 'desconhecido')}"
                if known_doi_source:
                    source += f"+doi_{known_doi_source}"
            else:
                source = source + "+sem_match_buscadores"

        meta = apply_discreet_metadata_fallback(meta, doc, cfg)
        if not include_inferred_metadata_notes(cfg):
            meta.note = None

        key = unique_key(make_bib_key(meta.authors or meta.editors, meta.year, meta.title), used)
        doc.bib_key = key
        doc.metadata = {**doc.metadata, **meta.model_dump(), "bib_inference_source": source}
        entries.append(render_biblatex_entry(key, meta))
        keys.append(key)
        diagnostics.append({
            "path": doc.path,
            "label": doc.label,
            "bib_key": key,
            "source": source,
            "known_doi": known_doi,
            "known_doi_source": known_doi_source,
            "inference_error": inference_error,
            "metadata_incomplete_after_enrichment": bibliographic_metadata_incomplete(meta, cfg),
            "metadata_enrichment": enrichment_diag,
            "metadata": meta.model_dump(),
        })
    return docs, entries, keys, diagnostics


def maybe_generate_revised_local_bib(
    client: OpenAI,
    model: str,
    cfg: dict[str, Any],
    research_paths: ResearchPaths,
) -> dict[str, Any] | None:
    """Opcionalmente substitui o .bib sintético por .bib revisado por IA.

    Ative no TOML:
      [documentos_locais]
      gerar_bib_revisado_ia = true

    Se houver .bib externo detectado/informado, a rotina só sobrescreve quando
    gerar_bib_revisado_ia=true; caso contrário, preserva o .bib externo.
    """
    if not is_local_documents_mode(cfg):
        return None
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    enabled = bool(local.get("gerar_bib_revisado_ia", local.get("gerar_bib_revisado_por_ia", False)))
    if not enabled:
        return None

    files = build_fulltext_cache_pdf_paths(research_paths)
    if not files:
        files = list(research_paths.selected_fulltext_paths or [])
    files = [p for p in files if p.exists() and p.is_file() and p.suffix.lower() in READABLE_SUFFIXES]
    if not files:
        return {"generated": False, "reason": "Nenhum documento local legível encontrado para inferência bibliográfica."}

    max_docs = int(local.get("max_docs_bib_revisado_ia") or len(files))
    files = files[:max_docs]
    docs: list[SourceDoc] = []
    for path in files:
        try:
            docs.append(SourceDoc(
                path=str(path),
                kind="texto_local_bib_revisado_ia",
                label=path.name,
                extracted_text=read_text_file(path, int(local.get("max_caracteres_bib_revisado") or 35000)),
                metadata={"source_path": str(path)},
            ))
        except Exception as exc:
            debug_print(f"Falha ao ler documento local para bib revisado {path}: {exc}")

    if not docs:
        return {"generated": False, "reason": "Documentos encontrados, mas nenhum texto pôde ser extraído."}

    existing_entries: list[str] = []
    existing_keys: list[str] = []
    if research_paths.bib_path and research_paths.bib_path.exists() and bool(local.get("preservar_chaves_bib_existente", False)):
        existing_entries, existing_keys = parse_bib_entries(research_paths.bib_path)

    revised_docs, revised_entries, revised_keys, diagnostics = build_revised_bib_entries_for_local_docs(
        client=client,
        model=model,
        cfg=cfg,
        docs=docs,
        existing_keys=existing_keys,
    )

    combined_entries = (existing_entries if existing_entries else []) + revised_entries
    local_dedup_enabled = bool(local.get("deduplicar_bib", local.get("deduplicar_referencias", True)))
    local_dedup_report: dict[str, Any] | None = None
    local_alias_map: dict[str, str] = {}
    if combined_entries and local_dedup_enabled:
        combined_entries, combined_keys, local_alias_map, local_dedup_report = deduplicate_bib_entries(combined_entries, preferred_keys=existing_keys + revised_keys)
        revised_keys = list(dict.fromkeys(local_alias_map.get(k, k) for k in revised_keys))
        for doc in revised_docs:
            if doc.bib_key in local_alias_map:
                doc.bib_key = local_alias_map[doc.bib_key]
    else:
        combined_keys = [k for e in combined_entries if (k := bib_entry_key(e))]

    if not research_paths.bib_path:
        research_paths.bib_path = research_paths.root_dir / f"{research_paths.root_dir.name}.bib"
    assert research_paths.bib_path is not None
    bib_path = research_paths.bib_path
    write_text(bib_path, "\n\n".join(combined_entries).strip() + "\n")
    revised_sidecar = bib_path.with_name(bib_path.stem + "_revisado_ia.bib")
    write_text(revised_sidecar, "\n\n".join(combined_entries).strip() + "\n")
    metadata_diag_path = bib_path.with_name(bib_path.stem + "_metadados_diagnostico.json")
    if bool(local.get("salvar_diagnostico_metadados", True)):
        write_text(metadata_diag_path, json.dumps(diagnostics, ensure_ascii=False, indent=2, default=str))

    path_key_map: dict[str, str] = {}
    name_key_map: dict[str, str] = {}
    meta_by_key: dict[str, dict[str, Any]] = {}
    for doc in revised_docs:
        if not doc.bib_key:
            continue
        p = Path(doc.path)
        try:
            path_key_map[str(p.resolve())] = doc.bib_key
        except Exception:
            path_key_map[str(p)] = doc.bib_key
        name_key_map[p.name] = doc.bib_key
        meta_by_key[doc.bib_key] = doc.metadata

    for entry in research_paths.selected_entries:
        raw = entry.get("downloaded_pdf_path") or entry.get("local_file_path")
        matched_key = None
        if raw:
            p = Path(str(raw))
            try:
                matched_key = path_key_map.get(str(p.resolve()))
            except Exception:
                matched_key = path_key_map.get(str(p))
            matched_key = matched_key or name_key_map.get(p.name)
        if matched_key:
            entry["bib_key"] = matched_key
            meta = meta_by_key.get(matched_key, {})
            if meta.get("title"):
                entry["title"] = meta.get("title")
            if meta.get("year"):
                entry["year"] = meta.get("year")
            if meta.get("authors"):
                entry["authors"] = meta.get("authors")
            if meta.get("doi"):
                entry["doi"] = meta.get("doi")
            if meta.get("url"):
                entry["url"] = meta.get("url")

    if research_paths.debug_path and research_paths.debug_path.exists():
        try:
            payload = json.loads(research_paths.debug_path.read_text(encoding="utf-8"))
        except Exception:
            payload = {}
        payload["selected_all"] = research_paths.selected_entries
        payload["local_revised_bib"] = {
            "generated": True,
            "bib_path": str(bib_path),
            "sidecar_path": str(revised_sidecar),
            "keys": revised_keys,
            "diagnostics": diagnostics,
            "metadata_diagnostics_path": str(metadata_diag_path),
        }
        write_text(research_paths.debug_path, json.dumps(payload, ensure_ascii=False, indent=2, default=str))

    cfg["__local_revised_bib_path__"] = str(bib_path)
    cfg["__local_revised_bib_keys__"] = revised_keys
    cfg["__local_revised_bib_sidecar_path__"] = str(revised_sidecar)
    cfg["__local_revised_bib_key_by_path__"] = path_key_map
    cfg["__local_revised_bib_key_by_name__"] = name_key_map
    return {
        "generated": True,
        "bib_path": str(bib_path),
        "sidecar_path": str(revised_sidecar),
        "keys": revised_keys,
        "documents": len(revised_docs),
        "diagnostics": diagnostics,
        "deduplication": local_dedup_report,
        "metadata_diagnostics_path": str(metadata_diag_path),
    }



def apply_local_revised_bib_keys_to_docs(cfg: dict[str, Any], docs: list[SourceDoc]) -> None:
    """Aplica chaves do .bib revisado aos SourceDoc do corpus local."""
    path_map = cfg.get("__local_revised_bib_key_by_path__") or {}
    name_map = cfg.get("__local_revised_bib_key_by_name__") or {}
    if not isinstance(path_map, dict):
        path_map = {}
    if not isinstance(name_map, dict):
        name_map = {}
    for doc in docs:
        if doc.bib_key:
            continue
        p = Path(doc.path)
        key = None
        try:
            key = path_map.get(str(p.resolve()))
        except Exception:
            key = path_map.get(str(p))
        key = key or name_map.get(p.name)
        if key:
            doc.bib_key = str(key)


class MindmapOutput(BaseModel):
    plantuml: str


def mindmap_config(cfg: dict[str, Any]) -> dict[str, Any]:
    section = cfg.get("mapa_mental", {}) if isinstance(cfg.get("mapa_mental"), dict) else {}
    if not section and isinstance(cfg.get("mindmap"), dict):
        section = cfg.get("mindmap", {})
    return section if isinstance(section, dict) else {}


def should_generate_mindmap(cfg: dict[str, Any]) -> bool:
    mm = mindmap_config(cfg)
    return bool(
        mm.get("gerar")
        or mm.get("ativo")
        or mm.get("habilitado")
        or mm.get("enabled")
    )


def sanitize_plantuml_mindmap(text: str) -> str:
    """Normaliza a saída da IA para código PlantUML mindmap válido.

    Esta função retorna apenas o conteúdo do arquivo .puml. O código não deve
    ser inserido como texto exportável no .org; o .org deve receber apenas a
    imagem renderizada, quando existir.
    """
    raw = strip_code_fences(text or "")
    m = re.search(r"(?is)@startmindmap.*?@endmindmap", raw)
    if m:
        raw = m.group(0)
    else:
        lines = [ln.rstrip() for ln in raw.splitlines() if ln.strip()]
        if not lines or not any(ln.lstrip().startswith("*") for ln in lines):
            lines = ["* Mapa mental", "** Síntese do corpus", "** Conceitos centrais", "** Relações analíticas"]
        raw = "@startmindmap\n" + "\n".join(lines) + "\n@endmindmap"
    raw = raw.replace("\r\n", "\n").strip()
    raw = re.sub(r"(?im)^\s*```.*$", "", raw)
    raw = re.sub(r"\n{3,}", "\n\n", raw)
    if not raw.lower().startswith("@startmindmap"):
        raw = "@startmindmap\n" + raw
    if not raw.lower().endswith("@endmindmap"):
        raw = raw.rstrip() + "\n@endmindmap"
    return raw.strip() + "\n"



def _toml_bool(value: Any, default: bool = False) -> bool:
    """Converte valores comuns de TOML/string para booleano."""
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    if text in {"1", "true", "t", "sim", "s", "yes", "y", "on"}:
        return True
    if text in {"0", "false", "f", "nao", "não", "n", "no", "off"}:
        return False
    return default


def _normalize_plantuml_color(value: Any) -> str:
    """Normaliza cores aceitas pelo PlantUML.

    Aceita valores como '#DCEBFF', 'DCEBFF' ou nomes de cores PlantUML
    como 'LightBlue'. Retorna string vazia se o valor for inválido.
    """
    color = str(value or "").strip()
    if not color:
        return ""
    if color.startswith("#"):
        return color
    if re.fullmatch(r"[0-9A-Fa-f]{6}|[0-9A-Fa-f]{3}", color):
        return "#" + color
    if re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", color):
        return color
    return ""


def apply_mindmap_level_colors(plantuml_code: str, cfg: dict[str, Any]) -> str:
    """Aplica cores por nível aos nós de um mindmap PlantUML.

    A coloração é controlada pelo TOML em [mapa_mental]. Exemplos:

      colorir_niveis = true
      cores_niveis = ["#DCEBFF", "#DCFCE7", "#FEF3C7", "#FEE2E2"]

    ou:

      [mapa_mental.cores_por_nivel]
      1 = "#DCEBFF"
      2 = "#DCFCE7"
      3 = "#FEF3C7"

    A função não duplica cores em linhas que já estejam coloridas.
    """
    mm = mindmap_config(cfg or {})

    raw_list = mm.get("cores_niveis") or mm.get("level_colors") or mm.get("colors_by_level_list") or []
    raw_map = mm.get("cores_por_nivel") or mm.get("colors_by_level") or {}

    has_color_config = bool(raw_list) or bool(raw_map)
    enabled = _toml_bool(
        mm.get("colorir_niveis", mm.get("usar_cores_niveis", mm.get("color_by_level", has_color_config))),
        default=has_color_config,
    )
    if not enabled:
        return plantuml_code.strip() + "\n"

    default_colors = [
        "#DCEBFF",  # nível 1 / raiz
        "#DCFCE7",  # nível 2
        "#FEF3C7",  # nível 3
        "#FEE2E2",  # nível 4
        "#F3E8FF",  # nível 5+
    ]

    colors: list[str] = []
    if isinstance(raw_list, list):
        colors = [_normalize_plantuml_color(c) for c in raw_list]
        colors = [c for c in colors if c]
    elif raw_list:
        # Permite string separada por vírgula, útil em TOMLs mais simples.
        colors = [_normalize_plantuml_color(c.strip()) for c in str(raw_list).split(",")]
        colors = [c for c in colors if c]
    if not colors:
        colors = default_colors

    color_by_level: dict[int, str] = {}
    if isinstance(raw_map, dict):
        for key, value in raw_map.items():
            try:
                level = int(str(key).strip())
            except Exception:
                continue
            color = _normalize_plantuml_color(value)
            if level >= 1 and color:
                color_by_level[level] = color

    out_lines: list[str] = []
    for line in plantuml_code.splitlines():
        stripped = line.lstrip()

        # Só colore nós de mindmap. Linhas como @startmindmap, skinparam,
        # left side/right side e @endmindmap permanecem intactas.
        if not stripped.startswith("*"):
            out_lines.append(line)
            continue

        indent = line[: len(line) - len(stripped)]
        match = re.match(r"^(\*+)(.*)$", stripped)
        if not match:
            out_lines.append(line)
            continue

        stars = match.group(1)
        rest = match.group(2).lstrip()
        if not rest:
            out_lines.append(line)
            continue

        # Preserva linhas que já têm cor ou marcador PlantUML explícito.
        if rest.startswith("[#") or (rest.startswith("[") and "]" in rest[:40]):
            out_lines.append(line)
            continue

        level = len(stars)
        color = color_by_level.get(level) or colors[min(level - 1, len(colors) - 1)]
        color = _normalize_plantuml_color(color)
        if not color:
            out_lines.append(line)
            continue

        out_lines.append(f"{indent}{stars}[{color}] {rest}")

    return "\n".join(out_lines).strip() + "\n"

def apply_plantuml_render_options(plantuml_code: str, cfg: dict[str, Any]) -> str:
    """Aplica opções de renderização de alta definição ao .puml."""
    mm = mindmap_config(cfg)
    try:
        dpi = int(mm.get("dpi") or mm.get("resolucao_dpi") or 300)
    except Exception:
        dpi = 300
    dpi = max(96, min(dpi, 600))
    code = sanitize_plantuml_mindmap(plantuml_code)
    code = apply_mindmap_level_colors(code, cfg)
    lines = code.splitlines()
    lower = "\n".join(lines).lower()
    inserts: list[str] = []
    if "skinparam dpi" not in lower:
        inserts.append(f"skinparam dpi {dpi}")
    if "skinparam backgroundcolor" not in lower:
        inserts.append("skinparam backgroundColor white")
    if "skinparam defaultfontname" not in lower:
        font_name = str(mm.get("fonte") or mm.get("font_name") or "Arial").strip() or "Arial"
        inserts.append(f'skinparam defaultFontName "{font_name}"')
    if not inserts:
        return code
    out: list[str] = []
    inserted = False
    for line in lines:
        out.append(line)
        if not inserted and line.strip().lower().startswith("@startmindmap"):
            out.extend(inserts)
            inserted = True
    if not inserted:
        out = ["@startmindmap", *inserts, *[ln for ln in lines if not ln.strip().lower().startswith("@startmindmap")]]
    return "\n".join(out).strip() + "\n"


def build_mindmap_prompt(
    cfg: dict[str, Any],
    context: DocumentContext,
    org_text: str,
    base_docs: list[SourceDoc],
    guidance_docs: list[SourceDoc],
    selected_corpus_catalog: list[dict[str, Any]] | None = None,
) -> str:
    mm = mindmap_config(cfg)
    max_niveis = int(mm.get("max_niveis") or 4)
    max_nos = int(mm.get("max_nos") or 45)
    titulo = str(mm.get("titulo") or "Mapa mental dos textos analisados").strip()
    foco = str(mm.get("foco") or "síntese analítica da atividade").strip()
    return textwrap.dedent(f"""
    Gere um mapa mental em PlantUML para ser renderizado em Org-mode/LaTeX.

    Regras obrigatórias:
    1. Retorne SOMENTE código PlantUML, sem explicações e sem cercas Markdown.
    2. Use a sintaxe @startmindmap ... @endmindmap.
    3. O nó central deve ser: {titulo}
    4. O mapa deve representar o foco: {foco}.
    5. Use no máximo {max_niveis} níveis hierárquicos.
    6. Use no máximo {max_nos} nós no total.
    7. Para atividade com vários textos/capítulos, crie ramos que mostrem:
       - cada texto/capítulo relevante;
       - conceitos centrais;
       - argumentos principais;
       - convergências e divergências;
       - síntese comparativa;
       - implicações para a análise de políticas públicas/administração pública.
    8. Não use citações Org no mapa mental.
    9. Não use caracteres de controle, HTML ou Markdown.
    10. Mantenha rótulos curtos, legíveis e adequados a uma figura acadêmica.

    Contexto do documento:
    {json.dumps(asdict(context), ensure_ascii=False, indent=2)}

    Catálogo do corpus:
    {shorten_text(json.dumps(selected_corpus_catalog or [], ensure_ascii=False, indent=2), 16000)}

    Textos-base resumidos:
    {json.dumps(summarize_docs(base_docs, excerpt_chars=1600), ensure_ascii=False, indent=2)}

    Orientações relevantes:
    {json.dumps(summarize_docs(guidance_docs[:8], excerpt_chars=1200), ensure_ascii=False, indent=2)}

    Documento já gerado:
    {shorten_text(org_text, 22000)}
    """).strip()


def render_plantuml_file(puml_path: Path, formato: str = "png", limit_size: int = 8192, cfg: dict[str, Any] | None = None) -> tuple[Path | None, str | None]:
    """Renderiza PlantUML em arquivo de imagem antes da exportação Org.

    A figura deve existir fisicamente antes de ser referenciada no .org.
    A rotina tenta, nesta ordem:
    1. comando `plantuml` disponível no PATH;
    2. JAR informado em [mapa_mental].plantuml_jar_path;
    3. JAR informado em [documento].plantuml_jar_path;
    4. JAR informado na variável de ambiente PLANTUML_JAR.
    """
    formato = (formato or "png").strip().lower().lstrip(".")
    if formato not in {"png", "svg"}:
        formato = "png"
    try:
        limit_size = int(limit_size or 8192)
    except Exception:
        limit_size = 8192
    limit_size = max(4096, min(limit_size, 32768))

    rendered = puml_path.with_suffix("." + formato)
    env = os.environ.copy()
    env["PLANTUML_LIMIT_SIZE"] = str(limit_size)

    commands: list[list[str]] = []
    plantuml = shutil.which("plantuml")
    if plantuml:
        commands.append([plantuml, f"-DPLANTUML_LIMIT_SIZE={limit_size}", f"-t{formato}", str(puml_path.name)])

    cfg_local = cfg or {}
    mm = mindmap_config(cfg_local)
    documento_cfg = cfg_local.get("documento", {}) if isinstance(cfg_local.get("documento", {}), dict) else {}
    jar_raw = mm.get("plantuml_jar_path") or documento_cfg.get("plantuml_jar_path") or os.getenv("PLANTUML_JAR")
    jar_path = None
    if jar_raw:
        try:
            jar_path = Path(str(jar_raw)).expanduser().resolve()
        except Exception:
            jar_path = None
    java = shutil.which("java")
    if java and jar_path and jar_path.exists():
        commands.append([java, f"-DPLANTUML_LIMIT_SIZE={limit_size}", "-jar", str(jar_path), f"-t{formato}", str(puml_path.name)])

    if not commands:
        return None, "PlantUML não encontrado. Instale o comando `plantuml` ou informe [mapa_mental].plantuml_jar_path, [documento].plantuml_jar_path ou PLANTUML_JAR. O .puml foi preservado, mas a imagem não foi renderizada."

    errors: list[str] = []
    for cmd in commands:
        proc = subprocess.run(cmd, cwd=str(puml_path.parent), capture_output=True, text=True, env=env)
        if proc.returncode == 0 and rendered.exists():
            return rendered, None
        errors.append(
            f"Comando: {' '.join(cmd)}\nRETURN: {proc.returncode}\nSTDOUT:\n{proc.stdout}\nSTDERR:\n{proc.stderr}"
        )

    if rendered.exists():
        return rendered, None
    return None, "Falha ao renderizar PlantUML. Tentativas:\n" + "\n---\n".join(errors)


def remove_existing_mindmap_sections(org_text: str) -> str:
    """Remove blocos e headings automáticos de mapa mental.

    A versão anterior removia apenas blocos automáticos e preservava headings
    contendo "mapa mental". Isso causava o erro: o heading ficava antes das
    referências e a figura depois. Agora removemos também headings vazios ou
    automáticos de mapa mental para reinseri-los corretamente após a bibliografia.
    """
    text = org_text.replace("\r\n", "\n")

    # 1) Remove bloco automático novo, marcado por sentinelas.
    text = re.sub(
        r"(?ims)^\s*#\s*MINDMAP_AUTO_START\b.*?^\s*#\s*MINDMAP_AUTO_END\b.*?(?:\n|$)",
        "\n\n",
        text,
    )

    # 2) Remove blocos LaTeX antigos do mapa mental.
    text = re.sub(
        r"(?ims)^\s*#\+begin_export\s+latex\s*$"
        r".*?\\includegraphics\[[^\]]*\]\{(?:[^}]*/)?mapa_mental\.(?:png|svg)\}"
        r".*?^\s*#\+end_export\s*$\n?",
        "\n\n",
        text,
    )

    # 3) Remove blocos src PlantUML contendo mindmap.
    text = re.sub(
        r"(?ims)^\s*#\+begin_src\s+plantuml\b(?:(?!^\s*#\+end_src\s*$).)*@startmindmap.*?^\s*#\+end_src\s*$\n?",
        "\n\n",
        text,
    )

    # 4) Remove trechos crus @startmindmap...@endmindmap.
    text = re.sub(r"(?ims)^\s*@startmindmap\b.*?@endmindmap\s*$\n?", "\n\n", text)
    text = re.sub(r"(?im)^\s*Arquivo PlantUML gerado:.*$\n?", "", text)

    # 5) Remove comentários antigos contendo PlantUML.
    text = re.sub(
        r"(?ims)^\s*#\+begin_comment\s*$.*?Arquivo PlantUML:.*?^\s*#\+end_comment\s*$\n?",
        "\n\n",
        text,
    )

    # 6) Remove headings automáticos/vazios de mapa mental antes da bibliografia.
    # Ex.: "* Mapa mental dos textos analisados" ou "* 6 Mapa Mental..."
    text = re.sub(
        r"(?ims)^\*+\s+(?:\d+\s+)?Mapa\s+mental[^\n]*\n\s*(?=^\s*#\+PRINT_BIBLIOGRAPHY:)",
        "\n",
        text,
    )

    # 7) Remove seção de mapa mental completa quando ela contém apenas artefatos automáticos.
    text = re.sub(
        r"(?ims)^\*+\s+(?:\d+\s+)?Mapa\s+mental[^\n]*\n"
        r"(?:(?!^\*+\s+|^\s*#\+PRINT_BIBLIOGRAPHY:).)*"
        r"(?=^\*+\s+|^\s*#\+PRINT_BIBLIOGRAPHY:|\Z)",
        "\n\n",
        text,
    )

    return re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"


def insert_mindmap_section(
    org_text: str,
    *,
    title: str,
    puml_rel_path: str,
    image_rel_path: str | None = None,
    plantuml_code: str | None = None,
    incluir_codigo_fonte: bool = False,
    posicao: str = "apos_referencias",
) -> str:
    """Insere o heading e a figura do mapa mental juntos após as referências.

    Para atividades e papers, o título do mapa deve ser heading Org, não título
    LaTeX solto. Assim ele aparece como seção numerada, por exemplo:
    "6 MAPA MENTAL DOS TEXTOS ANALISADOS".

    A ordem final fica:
      texto principal
      referências
      nova página
      heading do mapa mental
      figura do mapa mental
    """
    title = title.strip() or "Mapa mental dos textos analisados"
    org_text = remove_existing_mindmap_sections(org_text)

    if not image_rel_path:
        return org_text

    safe_image = image_rel_path.replace("\\", "/")

    latex_lines = [
        "# MINDMAP_AUTO_START: mapa_mental",
        "#+LATEX: \\clearpage",
        f"* {title}",
        "#+begin_export latex",
        r"\vspace{0.8em}",
        r"\begingroup",
        r"\centering",
        r"\includegraphics[width=\textwidth,height=0.78\textheight,keepaspectratio]{" + safe_image + r"}",
        r"\par",
        r"\endgroup",
        r"\clearpage",
        "#+end_export",
    ]

    if incluir_codigo_fonte and plantuml_code:
        latex_lines.extend([
            "#+begin_comment",
            f"Arquivo PlantUML: {puml_rel_path}",
            plantuml_code.strip(),
            "#+end_comment",
        ])

    latex_lines.extend([
        "# MINDMAP_AUTO_END: mapa_mental",
        "",
    ])

    mindmap_block = "\n".join(latex_lines) + "\n"
    pos = (posicao or "apos_referencias").strip().lower()

    # Padrão: depois da bibliografia.
    m = re.search(r"(?im)^\s*#\+PRINT_BIBLIOGRAPHY:?\s*$", org_text)
    if m:
        line_end = org_text.find("\n", m.end())
        if line_end == -1:
            line_end = m.end()
        return org_text[:line_end].rstrip() + "\n\n" + mindmap_block + org_text[line_end:].lstrip("\n")

    # Se a bibliografia não existir, recria a diretiva antes do mapa.
    return org_text.rstrip() + "\n\n#+PRINT_BIBLIOGRAPHY:\n\n" + mindmap_block

def maybe_generate_activity_mindmap(
    client: OpenAI,
    model: str,
    cfg: dict[str, Any],
    context: DocumentContext,
    org_text: str,
    base_docs: list[SourceDoc],
    guidance_docs: list[SourceDoc],
    selected_corpus_catalog: list[dict[str, Any]] | None,
    documento_output_dir: Path,
    documento_prefix: str,
) -> tuple[str, dict[str, Any] | None, str | None]:
    """Gera mapa mental PlantUML opcional para atividade/resumo/fichamento.

    Ative no TOML:
      [mapa_mental]
      gerar = true
      linguagem = "plantuml"
      formato = "png"
      renderizar = true
    """
    if not should_generate_mindmap(cfg):
        return org_text, None, None
    mm = mindmap_config(cfg)
    linguagem = str(mm.get("linguagem") or "plantuml").strip().lower()
    if linguagem not in {"plantuml", "puml"}:
        return org_text, {"generated": False, "reason": f"Linguagem de mapa mental não suportada: {linguagem}"}, None

    title = str(mm.get("titulo") or "Mapa mental dos textos analisados").strip()
    arquivo = slugify(str(mm.get("arquivo") or "mapa_mental"))
    formato = str(mm.get("formato") or "png").strip().lower().lstrip(".")
    renderizar = bool(mm.get("renderizar", True))
    inserir_no_org = bool(mm.get("inserir_no_org", True))
    incluir_codigo = bool(mm.get("incluir_codigo_fonte", False))
    falhar = bool(mm.get("falhar_se_nao_renderizar", False))
    posicao = str(mm.get("posicao") or "apos_referencias")

    images_dir = documento_output_dir / str(mm.get("diretorio_imagens") or "images")
    images_dir.mkdir(parents=True, exist_ok=True)
    puml_path = images_dir / f"{arquivo}.puml"

    prompt = build_mindmap_prompt(cfg, context, org_text, base_docs, guidance_docs, selected_corpus_catalog)
    resp = client.responses.parse(
        model=model,
        input=[{"role": "user", "content": prompt}],
        text_format=MindmapOutput,
    )
    if resp.output_parsed is None:
        raise RuntimeError("A IA não retornou o mapa mental em PlantUML.")
    plantuml_code = apply_plantuml_render_options(resp.output_parsed.plantuml, cfg)
    write_text(puml_path, plantuml_code)

    rendered_path: Path | None = None
    render_error: str | None = None
    if renderizar:
        rendered_path, render_error = render_plantuml_file(puml_path, formato=formato, limit_size=int(mm.get("plantuml_limit_size") or 8192), cfg=cfg)
        if render_error and falhar:
            raise RuntimeError(render_error)

    image_rel = None
    if rendered_path and rendered_path.exists():
        image_rel = os.path.relpath(rendered_path, documento_output_dir).replace(os.sep, "/")
    puml_rel = os.path.relpath(puml_path, documento_output_dir).replace(os.sep, "/")

    if inserir_no_org:
        org_text = insert_mindmap_section(
            org_text,
            title=title,
            puml_rel_path=puml_rel,
            image_rel_path=image_rel,
            plantuml_code=plantuml_code,
            incluir_codigo_fonte=incluir_codigo,
            posicao=posicao,
        )

    info = {
        "generated": True,
        "language": "plantuml",
        "title": title,
        "puml_path": str(puml_path),
        "puml_rel_path": puml_rel,
        "image_path": str(rendered_path) if rendered_path else None,
        "image_rel_path": image_rel,
        "rendered": bool(rendered_path and rendered_path.exists()),
        "render_error": render_error,
        "inserted_in_org": inserir_no_org,
        "position": posicao,
    }
    return org_text, info, prompt


def prepare_local_corpus(cfg: dict[str, Any]) -> ResearchPaths:
    """Transforma ZIP/pasta local em estrutura compatível com uma pesquisa pronta."""
    local = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    saida = cfg.setdefault("saida", {})
    pipeline = cfg.setdefault("pipeline", {})
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    prefixo = str(local.get("prefixo") or documento.get("prefixo") or saida.get("prefixo") or "documentos_locais").strip() or "documentos_locais"
    prefixo = slugify(prefixo)
    output_base = resolve_configured_path(local.get("output_dir"), cfg) if local.get("output_dir") else None
    if output_base is None:
        output_base = get_bundle_root_from_config(cfg) / "output" / "corpus_local"
    root_dir = output_base / prefixo if bool(local.get("criar_subdiretorio", True)) else output_base
    root_dir.mkdir(parents=True, exist_ok=True)
    cache_dir = root_dir / f"{prefixo}_fulltext_cache"
    if cache_dir.exists() and bool(local.get("limpar_cache_anterior", True)):
        shutil.rmtree(cache_dir)
    cache_dir.mkdir(parents=True, exist_ok=True)
    input_files, source_info = discover_local_input_files(cfg, root_dir)
    selected_entries: list[dict[str, Any]] = []
    bib_entries: list[str] = []
    external_bib_path = configured_external_bib_path(cfg)
    external_bib_entries: list[str] = []
    external_bib_keys: list[str] = []
    if external_bib_path:
        external_bib_entries, external_bib_keys = parse_bib_entries(external_bib_path)
        bib_entries = list(external_bib_entries)
        cfg["__local_external_bib_path__"] = str(external_bib_path)
        cfg["__local_external_bib_keys__"] = external_bib_keys
    copied_files: list[Path] = []
    used_keys: set[str] = set()
    for idx, src in enumerate(input_files, start=1):
        copied = unique_copy_to_dir(src, cache_dir) if bool(local.get("copiar_para_fulltext_cache", True)) else src
        copied_files.append(copied)
        base_key = slugify(copied.stem)[:70] or f"texto_local_{idx:02d}"
        matched_external_key = _bib_key_from_filename(copied, external_bib_entries) if external_bib_entries else None
        internal_key = matched_external_key or base_key
        n = 2
        while internal_key in used_keys:
            internal_key = f"{base_key}_{n}"
            n += 1
        used_keys.add(internal_key)
        if not external_bib_entries:
            bib_entries.append(build_local_bib_entry(internal_key, copied, cfg))
        selected_entries.append({
            "paper_id": f"local:{internal_key}",
            "title": copied.stem.replace("_", " ").replace("-", " ").strip().title() or copied.name,
            "year": None,
            "authors": [str(local.get("autor_padrao") or "Material fornecido pelo professor")],
            "source": "documentos_locais",
            "doi": None,
            "url": None,
            "downloaded_pdf_path": str(copied),
            "local_file_path": str(copied),
            # Se houver .bib externo e o arquivo não puder ser mapeado com segurança,
            # não gravamos uma chave sintética aqui. As chaves válidas serão as do .bib externo.
            "bib_key": matched_external_key if external_bib_entries else internal_key,
            "reason_selected": "Documento fornecido localmente pelo usuário/professor.",
            "filename": copied.name,
            "suffix": copied.suffix.lower(),
        })
    org_path = root_dir / f"{prefixo}.org"
    bib_path = root_dir / f"{prefixo}.bib"
    debug_path = root_dir / f"{prefixo}_debug.json"
    contexto_path = root_dir / f"{prefixo}_contexto_local.json"
    corpus_title = str(
        local.get("titulo_corpus")
        or documento.get("titulo_trabalho")
        or documento.get("titulo")
        or f"Corpus local — {pretty_title_from_prefix(prefixo)}"
    ).strip()
    lines = [
        f"#+TITLE: {corpus_title}",
        f"#+AUTHOR: {cfg.get('atividade', {}).get('aluno') or DEFAULT_AUTHOR}",
        "#+LANGUAGE: pt_BR",
        "#+OPTIONS: toc:t num:t",
        "#+CITE_EXPORT: biblatex apa",
        f"#+BIBLIOGRAPHY: {prefixo}.bib",
        "",
        "* Corpus local",
        "Este arquivo foi gerado automaticamente a partir de documentos locais/ZIP.",
        "",
        "* Documentos incluídos",
    ]
    for entry in selected_entries:
        lines.append(f"** {entry['title']}")
        lines.append(f"- Arquivo: {entry['filename']}")
        if entry.get("bib_key"):
            lines.append(f"- Chave BibTeX sugerida: [cite:@{entry['bib_key']}]")
        elif external_bib_keys:
            lines.append("- Chave BibTeX sugerida: ver bibliografia externa informada no TOML.")
        else:
            lines.append("- Chave BibTeX: não mapeada automaticamente.")
        lines.append("")
    lines.append("* Referências")
    lines.append("#+PRINT_BIBLIOGRAPHY:")
    write_text(org_path, "\n".join(lines).strip() + "\n")
    write_text(bib_path, "\n\n".join(bib_entries).strip() + "\n")
    debug_payload = {
        "local_corpus": True,
        "generated_at": datetime.now().isoformat(),
        "source_info": source_info,
        "proposal": {
            "tema": cfg.get("pesquisa", {}).get("tema") or cfg.get("documento", {}).get("tema") or cfg.get("atividade", {}).get("tema"),
            "recorte": cfg.get("pesquisa", {}).get("recorte") or cfg.get("documento", {}).get("recorte"),
            "objetivo": cfg.get("pesquisa", {}).get("objetivo") or cfg.get("documento", {}).get("objetivo"),
            "pergunta_pesquisa": cfg.get("pesquisa", {}).get("pergunta_pesquisa") or cfg.get("documento", {}).get("pergunta_pesquisa"),
        },
        "selected_all": selected_entries,
        "selected_count": len(selected_entries),
        "fulltext_cache_dir": str(cache_dir),
        "contexto_local_path": str(contexto_path),
        "external_bib_path": str(external_bib_path) if external_bib_path else None,
        "external_bib_keys": external_bib_keys,
    }
    write_text(debug_path, json.dumps(debug_payload, ensure_ascii=False, indent=2, default=str))
    write_text(contexto_path, json.dumps({
        "generated_at": datetime.now().isoformat(),
        "prefixo": prefixo,
        "documentos": selected_entries,
        "external_bib_path": str(external_bib_path) if external_bib_path else None,
        "external_bib_keys": external_bib_keys,
        "instrucoes": {
            "uso_obrigatorio": "Todos os documentos do fulltext_cache local devem ser interpretados pela IA e usados substantivamente no documento final.",
            "nao_inventar_fontes": True,
        },
    }, ensure_ascii=False, indent=2, default=str))
    pipeline["pesquisa_dir_existente"] = str(root_dir)
    pipeline["executar_pesquisa"] = False
    saida["prefixo"] = prefixo
    if not saida.get("output_dir"):
        saida["output_dir"] = str(output_base)
    cfg.setdefault("documento", {}).setdefault("usar_bib_da_pesquisa", True)
    cfg.setdefault("documento", {}).setdefault("usar_artigos_selecionados_pesquisa", True)
    cfg["__local_corpus_prepared__"] = True
    rp = detect_research_paths(cfg)
    rp.fulltext_cache_dir = cache_dir
    rp.selected_entries = selected_entries
    rp.selected_fulltext_paths = copied_files
    rp.fulltext_paths = copied_files
    return rp

def detect_research_output_dir(cfg: dict[str, Any]) -> Path:
    pipeline = cfg.get("pipeline", {})
    pesquisa_dir_existente = resolve_configured_path(pipeline.get("pesquisa_dir_existente"), cfg)
    if pesquisa_dir_existente is not None:
        return pesquisa_dir_existente

    saida = cfg.get("saida", {})
    base_dir = resolve_configured_path(saida.get("output_dir") or ".", cfg)
    if base_dir is None:
        base_dir = get_config_base_dir(cfg)
    prefixo = (saida.get("prefixo") or "atividade").strip()
    create_subdir = bool(saida.get("criar_subdiretorio", True))
    return (base_dir / prefixo) if create_subdir else base_dir

def detect_research_paths(cfg: dict[str, Any]) -> ResearchPaths:
    root_dir = detect_research_output_dir(cfg)
    prefixo = (cfg.get("saida", {}).get("prefixo") or "atividade").strip()
    rp = ResearchPaths(root_dir=root_dir)
    candidates = {
        "org": root_dir / f"{prefixo}.org",
        "bib": root_dir / f"{prefixo}.bib",
        "debug": root_dir / f"{prefixo}_debug.json",
        "pdf": root_dir / f"{prefixo}.pdf",
        "prisma_svg": root_dir / f"{prefixo}_prisma.svg",
        "prisma_pdf": root_dir / f"{prefixo}_prisma.pdf",
    }
    rp.org_path = candidates["org"] if candidates["org"].exists() else None
    rp.bib_path = candidates["bib"] if candidates["bib"].exists() else None
    rp.debug_path = candidates["debug"] if candidates["debug"].exists() else None
    rp.pdf_path = candidates["pdf"] if candidates["pdf"].exists() else None
    rp.prisma_svg_path = candidates["prisma_svg"] if candidates["prisma_svg"].exists() else None
    rp.prisma_pdf_path = candidates["prisma_pdf"] if candidates["prisma_pdf"].exists() else None

    config_output = cfg.get("controle", {}).get("config_output")
    if config_output:
        cp = resolve_configured_path(config_output, cfg)
        rp.config_path = cp if cp is not None and cp.exists() else None
    else:
        for pattern in ("*config*.toml",):
            matches = sorted(root_dir.glob(pattern))
            if matches:
                rp.config_path = matches[0]
                break

    cache_dir = root_dir / f"{prefixo}_fulltext_cache"
    rp.fulltext_cache_dir = cache_dir if cache_dir.exists() and cache_dir.is_dir() else None

    seen_selected: set[str] = set()
    if rp.debug_path and rp.debug_path.exists():
        try:
            debug_json = json.loads(rp.debug_path.read_text(encoding="utf-8"))
            for entry in (debug_json.get("selected_all") or []):
                if not isinstance(entry, dict):
                    continue
                rp.selected_entries.append(entry)
                pdf_raw = entry.get("downloaded_pdf_path")
                if pdf_raw:
                    pdf_path = Path(str(pdf_raw)).expanduser()
                    if not pdf_path.is_absolute():
                        pdf_path = (root_dir / pdf_path).resolve()
                    else:
                        pdf_path = pdf_path.resolve()
                    if pdf_path.exists() and pdf_path.is_file():
                        key = str(pdf_path)
                        if key not in seen_selected:
                            rp.selected_fulltext_paths.append(pdf_path)
                            seen_selected.add(key)
        except Exception as exc:
            debug_print(f"Falha ao ler debug JSON da pesquisa: {exc}")

    # fallback: se houver cache local, inclui somente o que casar com os selecionados
    if rp.fulltext_cache_dir and rp.selected_entries:
        by_name = {p.name: p for p in rp.fulltext_cache_dir.iterdir() if p.is_file() and p.suffix.lower() in READABLE_SUFFIXES}
        for entry in rp.selected_entries:
            pdf_raw = entry.get("downloaded_pdf_path")
            if not pdf_raw:
                continue
            name = Path(str(pdf_raw)).name
            cand = by_name.get(name)
            if cand and str(cand) not in seen_selected:
                rp.selected_fulltext_paths.append(cand)
                seen_selected.add(str(cand))

    # fallback residual: qualquer PDF do cache, se nada explícito foi encontrado
    if rp.fulltext_cache_dir and not rp.selected_fulltext_paths:
        for pdf in sorted(p for p in rp.fulltext_cache_dir.iterdir() if p.is_file() and p.suffix.lower() in READABLE_SUFFIXES):
            rp.selected_fulltext_paths.append(pdf)

    # backward compatibility: mantém também lista ampla de PDFs adicionais úteis
    for pdf in sorted(root_dir.glob("*.pdf")):
        if pdf in {rp.pdf_path, rp.prisma_pdf_path}:
            continue
        rp.fulltext_paths.append(pdf)
    if rp.fulltext_cache_dir:
        for pdf in sorted(p for p in rp.fulltext_cache_dir.iterdir() if p.is_file() and p.suffix.lower() in READABLE_SUFFIXES):
            if pdf not in rp.fulltext_paths:
                rp.fulltext_paths.append(pdf)
    return rp

def _ensure_list_of_strings(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(v).strip() for v in value if str(v).strip()]
    text = str(value).strip()
    return [text] if text else []

def collect_orientation_values(section: dict[str, Any] | None) -> list[str]:
    values: list[str] = []
    if not isinstance(section, dict):
        return values
    seen: set[str] = set()

    def add(raw: Any) -> None:
        for item in _ensure_list_of_strings(raw):
            if item not in seen:
                values.append(item)
                seen.add(item)

    add(section.get("orientacoes_paths"))
    add(section.get("orientacao_inline"))
    return values


def orientation_file_sort_key(path: Path) -> tuple[int, str]:
    """Ordena orientações extraídas de pasta/ZIP por provável relevância.

    Arquivos com "final" no nome vêm primeiro, pois é comum o usuário enviar
    versão preliminar e versão final no mesmo pacote. A ordenação é estável o
    suficiente para manter previsibilidade sem descartar nenhum arquivo.
    """
    name = normalize_title_loose(path.name)
    score = 0
    if "final" in name or "versao final" in name or "versao_final" in name:
        score -= 30
    if any(token in name for token in ("roteiro", "orientacao", "orientacoes", "instrucao", "instrucoes", "prompt")):
        score -= 10
    if any(token in name for token in ("rascunho", "draft", "antigo", "old")):
        score += 20
    suffix_order = {".docx": 0, ".pdf": 1, ".org": 2, ".md": 3, ".txt": 4}
    score += suffix_order.get(path.suffix.lower(), 9)
    return score, str(path).lower()


def expand_orientation_source_paths(
    raw: str,
    *,
    base_dir: Path | None = None,
    extract_dir: Path | None = None,
    max_files: int = MAX_ORIENTATION_FILES_FROM_ARCHIVE,
) -> list[Path] | None:
    """Expande uma orientação informada no TOML em arquivos legíveis.

    Aceita:
    - arquivo legível direto (.pdf, .docx, .txt, .md, .org etc.);
    - diretório com arquivos legíveis;
    - arquivo .zip contendo arquivos legíveis.

    Retorna None quando `raw` deve ser tratado como texto inline, e lista vazia
    quando o caminho existe mas não contém arquivos suportados.
    """
    raw = str(raw or "").strip()
    if not raw:
        return []
    candidate = safe_resolve_user_path(raw, base_dir=base_dir)
    if candidate is None or not candidate.exists():
        return None

    try:
        candidate = candidate.resolve()
    except Exception:
        pass

    if candidate.is_file() and candidate.suffix.lower() in READABLE_SUFFIXES:
        return [candidate]

    if candidate.is_dir():
        files = [p for p in candidate.rglob("*") if p.is_file() and p.suffix.lower() in READABLE_SUFFIXES]
        return sorted(files, key=orientation_file_sort_key)[:max_files]

    if candidate.is_file() and candidate.suffix.lower() in ORIENTATION_ARCHIVE_SUFFIXES:
        root = extract_dir or ((base_dir or Path.cwd()) / "_orientacoes_extraidas")
        target_dir = root / slugify(candidate.stem)
        if target_dir.exists():
            shutil.rmtree(target_dir)
        extracted = safe_extract_zip(candidate, target_dir)
        files = [p for p in extracted if p.is_file() and p.suffix.lower() in READABLE_SUFFIXES]
        files = sorted(files, key=orientation_file_sort_key)[:max_files]
        if len(files) >= max_files:
            debug_print(f"ZIP de orientação limitado aos primeiros {max_files} arquivos suportados: {candidate}")
        return files

    return []

def resolve_orientation_contents(
    values: list[str],
    *,
    max_chars: int = 40000,
    base_dir: Path | None = None,
    extract_dir: Path | None = None,
) -> list[tuple[str, str]]:
    chunks: list[tuple[str, str]] = []
    for idx, raw in enumerate(values, start=1):
        label = f"orientacao_{idx}"
        expanded_paths = expand_orientation_source_paths(
            str(raw),
            base_dir=base_dir,
            extract_dir=extract_dir,
        )
        if expanded_paths is not None:
            if not expanded_paths:
                debug_print(f"Orientação ignorada por não conter arquivo suportado: {raw}")
                continue
            for path_idx, candidate in enumerate(expanded_paths, start=1):
                try:
                    text = read_text_file(candidate, max_chars=max_chars)
                    chunks.append((str(candidate), text))
                except Exception as exc:
                    debug_print(f"Falha ao ler orientação {candidate}: {exc}")
            continue
        chunks.append((f"inline:{label}", shorten_text(str(raw), max_chars)))
    return chunks

def write_combined_orientation_file(values: list[str], output_path: Path, *, title: str, base_dir: Path | None = None) -> Path | None:
    chunks = resolve_orientation_contents(
        values,
        max_chars=50000,
        base_dir=base_dir,
        extract_dir=output_path.parent / "_orientacoes_extraidas",
    )
    if not chunks:
        return None
    parts: list[str] = [f"# {title}"]
    for idx, (source, text) in enumerate(chunks, start=1):
        parts.append(f"\n## Bloco {idx} — {source}\n")
        parts.append(text.strip())
    write_text(output_path, "\n\n".join(parts).strip() + "\n")
    return output_path

def filter_research_config(cfg: dict[str, Any], work_dir: Path) -> dict[str, Any]:
    research_cfg = {k: json.loads(json.dumps(v)) for k, v in cfg.items() if k in RESEARCH_SECTIONS}

    saida = research_cfg.setdefault("saida", {})
    triagem = research_cfg.setdefault("triagem", {})
    config_base_dir = get_config_base_dir(cfg)

    saida_values = collect_orientation_values(cfg.get("saida", {}))
    triagem_values = collect_orientation_values(cfg.get("triagem", {}))

    saida.pop("orientacoes_paths", None)
    saida.pop("orientacao_inline", None)
    triagem.pop("orientacoes_paths", None)
    triagem.pop("orientacao_inline", None)

    if saida_values:
        combined = write_combined_orientation_file(
            saida_values,
            work_dir / "_orientacoes_saida_combinadas.txt",
            title="Orientações gerais da pesquisa",
            base_dir=config_base_dir,
        )
        if combined is not None:
            saida[RESEARCH_BRIDGE_OUT_ORIENTACAO] = str(combined)

    if triagem_values:
        combined = write_combined_orientation_file(
            triagem_values,
            work_dir / "_orientacoes_triagem_combinadas.txt",
            title="Orientações específicas de triagem",
            base_dir=config_base_dir,
        )
        if combined is not None:
            triagem[RESEARCH_BRIDGE_TRIAGEM_PROMPT] = str(combined)
            triagem[RESEARCH_BRIDGE_TRIAGEM_EXTRAS] = ""

    return research_cfg

def dumps_toml_value(value: Any) -> str:
    if isinstance(value, bool):
        return "true" if value else "false"
    if value is None:
        return '""'
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(value)
    if isinstance(value, list):
        return "[" + ", ".join(dumps_toml_value(v) for v in value) + "]"
    if isinstance(value, dict):
        raise TypeError("Dict não deve ser serializado diretamente aqui")
    text = str(value).replace("\\", "\\\\").replace('"', '\\"')
    return f'"{text}"'

def dict_to_toml(cfg: dict[str, Any]) -> str:
    lines: list[str] = []
    for section, values in cfg.items():
        if not isinstance(values, dict):
            continue
        lines.append(f"[{section}]")
        for key, value in values.items():
            lines.append(f"{key} = {dumps_toml_value(value)}")
        lines.append("")
    return "\n".join(lines).strip() + "\n"

def run_research_stage(config_path: Path, script_path: Path, cwd: Path | None = None) -> None:
    ensure_command(sys.executable)
    if not script_path.exists():
        raise RuntimeError(f"Script de pesquisa não encontrado: {script_path}")
    cmd = [sys.executable, str(script_path), "--config", str(config_path)]
    debug_print("Executando etapa de pesquisa:", cmd)
    proc = subprocess.run(cmd, cwd=str(cwd or script_path.parent), text=True)
    if proc.returncode != 0:
        raise RuntimeError(f"Etapa de pesquisa falhou com código {proc.returncode}.")

def extract_selected_titles_from_org(org_text: str) -> list[str]:
    titles: list[str] = []
    for m in re.finditer(r"\*\s+Texto selecionado\s+\d+.*?\n.*?\n([^\n]+)", org_text, flags=re.IGNORECASE | re.DOTALL):
        cand = m.group(1).strip()
        if cand and cand not in titles:
            titles.append(cand)
    return titles

def build_document_context(cfg: dict[str, Any], research_paths: ResearchPaths) -> tuple[DocumentContext, dict[str, Any]]:
    pesquisa = cfg.get("pesquisa", {})
    atividade = cfg.get("atividade", {})
    debug_json: dict[str, Any] = {}
    if research_paths.debug_path and research_paths.debug_path.exists():
        try:
            debug_json = json.loads(research_paths.debug_path.read_text(encoding="utf-8"))
        except Exception:
            debug_json = {}

    pergunta = None
    hipotese = None
    titulo_sugerido = pesquisa.get("trabalho") or None
    if isinstance(debug_json.get("proposal"), dict):
        proposal = debug_json["proposal"]
        pergunta = proposal.get("pergunta_pesquisa") or proposal.get("pergunta") or pergunta
        hipotese = proposal.get("hipotese") or hipotese
    if research_paths.org_path and research_paths.org_path.exists() and not titulo_sugerido:
        org_text = research_paths.org_path.read_text(encoding="utf-8", errors="ignore")
        m = re.search(r"(?im)^#\+TITLE:\s*(.+)$", org_text)
        if m:
            titulo_sugerido = m.group(1).strip()

    context = DocumentContext(
        tema=pesquisa.get("tema", ""),
        recorte=pesquisa.get("recorte", ""),
        objetivo=pesquisa.get("objetivo", ""),
        pergunta_pesquisa=pergunta,
        hipotese=hipotese,
        palavras_chave=list(pesquisa.get("palavras_chave") or []),
        titulo_sugerido=titulo_sugerido,
        tipo_estudo=pesquisa.get("tipo_estudo") or None,
        idiomas=list(pesquisa.get("idiomas") or ["português", "inglês"]),
        modo_origem=atividade.get("modo", "revisao_sistematica"),
        titulo_trabalho_base=pesquisa.get("trabalho") or None,
    )
    return context, debug_json


def is_empty_context_field(value: Any) -> bool:
    """Retorna True para campos de contexto que devem ser inferidos pela IA.

    Além de string vazia, trata como vazios os placeholders gerados pelo
    próprio gerador de TOML, como "Preencher...", "Tema da pesquisa" etc.
    """
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if not text:
        return True
    low = text.lower()
    placeholders = {
        "tema da pesquisa",
        "recorte analítico",
        "objetivo geral",
        "preencher o tema do paper",
        "preencher o recorte conforme o roteiro",
        "paper",
        "dissertação",
        "dissertacao",
        "atividade",
    }
    if low in placeholders:
        return True
    return low.startswith("preencher ") or low.startswith("a preencher") or low.startswith("título a ser gerado") or low.startswith("titulo a ser gerado")


def configured_document_title_value(cfg: dict[str, Any]) -> str:
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    atividade = cfg.get("atividade", {}) if isinstance(cfg.get("atividade"), dict) else {}
    pesquisa = cfg.get("pesquisa", {}) if isinstance(cfg.get("pesquisa"), dict) else {}
    for section, keys in (
        (documento, ("titulo_trabalho", "titulo", "title")),
        (atividade, ("titulo_trabalho", "titulo", "title")),
        (pesquisa, ("titulo_trabalho", "titulo", "title")),
    ):
        for key in keys:
            value = str(section.get(key) or "").strip()
            if value:
                return value
    return ""


def should_infer_missing_context_fields(cfg: dict[str, Any]) -> bool:
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    pesquisa = cfg.get("pesquisa", {}) if isinstance(cfg.get("pesquisa"), dict) else {}
    raw = documento.get("inferir_campos_vazios_ia", pesquisa.get("inferir_campos_vazios_ia", True))
    return _toml_bool(raw, default=True)


def infer_missing_document_context(
    client: OpenAI,
    model: str,
    cfg: dict[str, Any],
    context: DocumentContext,
    base_docs: list[SourceDoc],
    guidance_docs: list[SourceDoc],
    selected_corpus_catalog: list[dict[str, Any]] | None = None,
) -> tuple[DocumentContext, dict[str, Any]]:
    """Preenche título, tema, recorte e objetivo quando vierem vazios no TOML.

    A inferência usa somente o corpus local/pesquisa já carregado e as
    orientações lidas de arquivos, pastas ou ZIP. Quando o campo já está
    preenchido com valor substantivo, ele é preservado.
    """
    if not should_infer_missing_context_fields(cfg):
        return context, {"used": False, "source": "disabled"}

    pesquisa_cfg = cfg.get("pesquisa", {}) if isinstance(cfg.get("pesquisa"), dict) else {}
    documento_cfg = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    atividade_cfg = cfg.get("atividade", {}) if isinstance(cfg.get("atividade"), dict) else {}

    raw_title = configured_document_title_value(cfg)
    missing = {
        "titulo_sugerido": is_empty_context_field(raw_title) and is_empty_context_field(context.titulo_sugerido),
        "tema": is_empty_context_field(pesquisa_cfg.get("tema", context.tema)),
        "recorte": is_empty_context_field(pesquisa_cfg.get("recorte", context.recorte)),
        "objetivo": is_empty_context_field(pesquisa_cfg.get("objetivo", context.objetivo)),
    }
    if not any(missing.values()):
        return context, {"used": False, "source": "all_fields_already_filled", "missing": missing}

    prompt = textwrap.dedent(
        f"""
        Você deve inferir campos acadêmicos faltantes para um documento em português.

        Contexto atual vindo do TOML/pesquisa:
        {json.dumps(asdict(context), ensure_ascii=False, indent=2)}

        Campos que precisam ser inferidos porque vieram vazios ou genéricos:
        {json.dumps(missing, ensure_ascii=False, indent=2)}

        Tipo de documento: {normalize_document_type(documento_cfg.get('tipo_documento'))}
        Disciplina/atividade:
        {json.dumps(atividade_cfg, ensure_ascii=False, indent=2)}

        Regras obrigatórias:
        1. Não invente tema alheio ao corpus local, à pesquisa carregada ou às orientações.
        2. Use as orientações como comando hierarquicamente superior ao corpus.
        3. Se houver roteiro, edital, enunciado, prompt ou arquivo com "final" no nome, trate-o como orientação principal.
        4. Preserve qualquer campo já preenchido com conteúdo substantivo.
        5. Gere título acadêmico específico, não genérico.
        6. Gere tema, recorte e objetivo compatíveis entre si.
        7. Se possível, gere também pergunta de pesquisa e palavras-chave.
        8. Retorne JSON estruturado, sem comentários fora do JSON.

        Orientações lidas de arquivos/pastas/ZIP:
        {json.dumps(summarize_docs(guidance_docs, excerpt_chars=2600), ensure_ascii=False, indent=2)}

        Catálogo do corpus/pesquisa:
        {shorten_text(json.dumps(selected_corpus_catalog or [], ensure_ascii=False, indent=2), 16000)}

        Trechos dos documentos-base:
        {json.dumps(summarize_docs(base_docs[:30], excerpt_chars=2200), ensure_ascii=False, indent=2)}
        """
    ).strip()

    resp = client.responses.parse(
        model=model,
        input=[{"role": "user", "content": prompt}],
        text_format=InferredDocumentContextOutput,
    )
    parsed = resp.output_parsed
    if parsed is None:
        return context, {"used": False, "source": "ia_no_structured_output", "missing": missing}

    new_context = DocumentContext(
        tema=(parsed.tema if missing["tema"] and parsed.tema else context.tema),
        recorte=(parsed.recorte if missing["recorte"] and parsed.recorte else context.recorte),
        objetivo=(parsed.objetivo if missing["objetivo"] and parsed.objetivo else context.objetivo),
        pergunta_pesquisa=parsed.pergunta_pesquisa or context.pergunta_pesquisa,
        hipotese=parsed.hipotese or context.hipotese,
        palavras_chave=parsed.palavras_chave or context.palavras_chave,
        titulo_sugerido=(parsed.titulo_sugerido if missing["titulo_sugerido"] and parsed.titulo_sugerido else context.titulo_sugerido),
        tipo_estudo=context.tipo_estudo,
        idiomas=context.idiomas,
        modo_origem=context.modo_origem,
        titulo_trabalho_base=context.titulo_trabalho_base,
    )

    # Atualiza a configuração em memória para que prompts, capa e auditoria usem
    # os valores inferidos na mesma execução.
    pesquisa_cfg["tema"] = new_context.tema
    pesquisa_cfg["recorte"] = new_context.recorte
    pesquisa_cfg["objetivo"] = new_context.objetivo
    if new_context.pergunta_pesquisa:
        pesquisa_cfg["pergunta_pesquisa"] = new_context.pergunta_pesquisa
    if new_context.palavras_chave:
        pesquisa_cfg["palavras_chave"] = new_context.palavras_chave
    if missing["titulo_sugerido"] and new_context.titulo_sugerido:
        documento_cfg.setdefault("titulo_trabalho", new_context.titulo_sugerido)
        atividade_cfg.setdefault("titulo_trabalho", new_context.titulo_sugerido)

    return new_context, {
        "used": True,
        "source": "ia_inferred_missing_fields",
        "missing": missing,
        "rationale": parsed.rationale,
        "prompt": prompt,
        "parsed": parsed.model_dump(),
    }

def maybe_rewrite_document_context(
    client: OpenAI,
    model: str,
    cfg: dict[str, Any],
    context: DocumentContext,
    guidance_docs: list[SourceDoc],
) -> tuple[DocumentContext, dict[str, Any]]:
    documento = cfg.get("documento", {})
    usar_contexto_consolidado = bool(documento.get("usar_contexto_consolidado_da_pesquisa", True))
    reformular = bool(documento.get("reformular_tema_recorte_objetivo", False))
    modo_escrita = str(documento.get("modo_escrita") or "novo").strip().lower()

    if usar_contexto_consolidado and not reformular:
        return context, {"used": False, "source": "pesquisa_consolidada"}

    override_tema = str(documento.get("tema") or "").strip()
    override_recorte = str(documento.get("recorte") or "").strip()
    override_objetivo = str(documento.get("objetivo") or "").strip()
    if not reformular and (override_tema or override_recorte or override_objetivo):
        new_context = DocumentContext(
            tema=override_tema or context.tema,
            recorte=override_recorte or context.recorte,
            objetivo=override_objetivo or context.objetivo,
            pergunta_pesquisa=context.pergunta_pesquisa,
            hipotese=context.hipotese,
            palavras_chave=context.palavras_chave,
            titulo_sugerido=context.titulo_sugerido,
            tipo_estudo=context.tipo_estudo,
            idiomas=context.idiomas,
            modo_origem=context.modo_origem,
            titulo_trabalho_base=context.titulo_trabalho_base,
        )
        return new_context, {"used": True, "source": "paper_override_manual"}

    prompt = textwrap.dedent(
        f"""
        Reformule, para a etapa de redação do documento, o tema, o recorte e o objetivo abaixo.

        Regras:
        - preserve o núcleo analítico da pesquisa já consolidada;
        - NÃO mude o assunto central;
        - apenas refine a formulação para a escrita do documento;
        - leve em conta o modo de escrita: {modo_escrita};
        - se o modo for "reescrever", priorize coesão e clareza do argumento;
        - se o modo for "expandir", mantenha o núcleo e permita maior abrangência analítica moderada;
        - retorne JSON estruturado.

        Contexto consolidado da pesquisa:
        {json.dumps(asdict(context), ensure_ascii=False, indent=2)}

        Orientações e artefatos disponíveis:
        {json.dumps(summarize_docs(guidance_docs), ensure_ascii=False, indent=2)}
        """
    ).strip()
    resp = client.responses.parse(
        model=model,
        input=[{"role": "user", "content": prompt}],
        text_format=RewrittenDocumentContextOutput,
    )
    parsed = resp.output_parsed
    if parsed is None:
        return context, {"used": False, "source": "fallback_contexto_consolidado"}

    new_context = DocumentContext(
        tema=parsed.tema or context.tema,
        recorte=parsed.recorte or context.recorte,
        objetivo=parsed.objetivo or context.objetivo,
        pergunta_pesquisa=parsed.pergunta_pesquisa or context.pergunta_pesquisa,
        hipotese=parsed.hipotese or context.hipotese,
        palavras_chave=context.palavras_chave,
        titulo_sugerido=parsed.titulo_sugerido or context.titulo_sugerido,
        tipo_estudo=context.tipo_estudo,
        idiomas=context.idiomas,
        modo_origem=context.modo_origem,
        titulo_trabalho_base=context.titulo_trabalho_base,
    )
    return new_context, {
        "used": True,
        "source": "ia_rewrite",
        "rationale": parsed.rationale,
        "modo_escrita": modo_escrita,
    }

def collect_guidance_docs(cfg: dict[str, Any], research_paths: ResearchPaths) -> list[SourceDoc]:
    docs: list[SourceDoc] = []
    config_base_dir = get_config_base_dir(cfg)

    def add(path: Path, kind: str, label: str | None = None) -> None:
        if not path.exists() or not path.is_file():
            return
        try:
            if path.suffix.lower() in READABLE_SUFFIXES:
                text = read_text_file(path, 20000)
            else:
                text = shorten_text(path.read_text(encoding="utf-8", errors="ignore"), 20000)
            docs.append(SourceDoc(
                path=str(path),
                kind=kind,
                label=label or path.name,
                extracted_text=text,
            ))
        except Exception as exc:
            debug_print(f"Falha ao ler orientação {path}: {exc}")

    def add_values(values: list[str], kind_prefix: str) -> None:
        extract_dir = research_paths.root_dir / "_orientacoes_extraidas" / kind_prefix
        for idx, raw in enumerate(values, start=1):
            raw = str(raw).strip()
            if not raw:
                continue
            expanded_paths = expand_orientation_source_paths(
                raw,
                base_dir=config_base_dir,
                extract_dir=extract_dir,
            )
            if expanded_paths is not None:
                if not expanded_paths:
                    debug_print(f"Orientação ignorada por não conter arquivo suportado: {raw}")
                    continue
                for path_idx, maybe_path in enumerate(expanded_paths, start=1):
                    add(maybe_path, kind_prefix, f"{kind_prefix}_{idx}_{path_idx}")
            else:
                docs.append(SourceDoc(
                    path=f"inline:{kind_prefix}_{idx}",
                    kind=kind_prefix,
                    label=f"{kind_prefix}_{idx}",
                    extracted_text=shorten_text(raw, 20000),
                ))

    # orientações gerais da pesquisa
    saida_values = collect_orientation_values(cfg.get("saida", {}))
    add_values(saida_values, "orientacao_pesquisa")

    # artefatos da pesquisa como orientação
    if research_paths.org_path:
        add(research_paths.org_path, "pesquisa_org", "pesquisa_org")
    if research_paths.debug_path:
        add(research_paths.debug_path, "pesquisa_debug", "pesquisa_debug")

    # orientações específicas do documento acadêmico
    documento_values = collect_orientation_values(cfg.get("documento", {}))
    add_values(documento_values, "orientacao_documento")

    prompts_section = cfg.get("prompts", {}) if isinstance(cfg.get("prompts"), dict) else {}
    prompt_values: list[str] = []
    for key, value in prompts_section.items():
        k = str(key).lower()
        if k.endswith("_path") or k.endswith("_paths") or "prompt" in k or k.endswith("_inline"):
            prompt_values.extend(_ensure_list_of_strings(value))
    add_values(prompt_values, "prompt_externo")

    # reescrita/expansão do documento acadêmico: usa o .org anterior como orientação
    documento = cfg.get("documento", {})
    if bool(documento.get("reescrever_a_partir_do_org_atual", False)):
        org_existente_raw = str(documento.get("documento_org_existente") or "").strip()
        org_existente: Path | None = None
        if org_existente_raw:
            p = safe_resolve_user_path(org_existente_raw, base_dir=config_base_dir)
            if p is None:
                p = Path(os.path.expanduser(org_existente_raw)).resolve()
            if p.exists() and p.is_file():
                org_existente = p
        else:
            doc_type_local = normalize_document_type(documento.get("tipo_documento"))
            prefix = (documento.get("prefixo") or ((cfg.get("saida", {}).get("prefixo") or "atividade") + f"_{DEFAULT_DOC_PREFIX_BY_TYPE.get(doc_type_local, 'documento')}" )).strip()
            output_dir_raw = documento.get("output_dir")
            if output_dir_raw:
                base = Path(output_dir_raw).expanduser().resolve()
                create_subdir = bool(documento.get("criar_subdiretorio", True))
                documento_dir = base / prefix if create_subdir else base
            else:
                documento_dir = build_document_output_dir(cfg, research_paths.root_dir)
            candidate = documento_dir / f"{prefix}.org"
            if candidate.exists():
                org_existente = candidate

        if org_existente is not None:
            add(org_existente, "paper_org_anterior", "paper_org_anterior")

    return docs

def collect_base_docs(cfg: dict[str, Any], research_paths: ResearchPaths, max_files: int | None = None) -> list[SourceDoc]:
    docs: list[SourceDoc] = []
    documento = cfg.get("documento", {})

    def add(path: Path, kind: str, label: str | None = None) -> None:
        if not path.exists() or path.suffix.lower() not in READABLE_SUFFIXES:
            return
        try:
            docs.append(SourceDoc(
                path=str(path),
                kind=kind,
                label=label or path.name,
                extracted_text=read_text_file(path, 25000),
            ))
        except Exception as exc:
            debug_print(f"Falha ao ler texto-base {path}: {exc}")

    usar_selecionados = bool(documento.get("usar_artigos_selecionados_pesquisa", True))
    permitir_correlata_extra = bool(documento.get("permitir_busca_correlata_extra", False))
    citar_todos_fulltext_cache = True if normalize_document_type(documento.get("tipo_documento")) == "dissertacao" else bool(documento.get("citar_todos_fulltext_cache", True))
    if max_files is None:
        doc_type_local = normalize_document_type(documento.get("tipo_documento"))
        if citar_todos_fulltext_cache and doc_type_local == "dissertacao":
            default_max = max(len(build_fulltext_cache_pdf_paths(research_paths)), len(research_paths.selected_fulltext_paths), 9999)
        else:
            default_max = 18 if doc_type_local == "dissertacao" else 12
        max_files = int(documento.get("max_selected_base_docs", default_max) or default_max)

    if usar_selecionados:
        source_paths = build_fulltext_cache_pdf_paths(research_paths) if citar_todos_fulltext_cache else research_paths.selected_fulltext_paths
        for path in source_paths[:max_files]:
            add(path, "texto_selecionado_fulltext_cache" if citar_todos_fulltext_cache else "texto_selecionado_formal")

        if not docs and research_paths.selected_entries:
            for entry in research_paths.selected_entries[:max_files]:
                title = str(entry.get("title") or "Sem título")
                abstract = str(entry.get("abstract") or entry.get("tldr") or "").strip()
                if not abstract:
                    continue
                docs.append(SourceDoc(
                    path=f"selected_all:{entry.get('paper_id') or slugify(title)}",
                    kind="texto_selecionado_resumo",
                    label=title,
                    extracted_text=shorten_text(abstract, 25000),
                    metadata={
                        "paper_id": entry.get("paper_id"),
                        "doi": entry.get("doi"),
                        "url": entry.get("url"),
                        "pdf_url": entry.get("pdf_url"),
                        "downloaded_pdf_path": entry.get("downloaded_pdf_path"),
                    },
                ))

    if permitir_correlata_extra:
        selected_set = {str(p.resolve()) for p in research_paths.selected_fulltext_paths}
        extra_count = 0
        for path in research_paths.fulltext_paths:
            if str(path.resolve()) in selected_set:
                continue
            add(path, "texto_correlato_extra_cache")
            extra_count += 1
            if len(docs) >= max_files + 8:
                break
        if extra_count:
            debug_print(f"Textos correlatos extras de cache adicionados: {extra_count}")

    if not docs and research_paths.org_path:
        add(research_paths.org_path, "texto_base_revisao", "pesquisa_org")
    if research_paths.pdf_path:
        add(research_paths.pdf_path, "relatorio_pesquisa_pdf", "pesquisa_pdf")
    return docs

def summarize_docs(docs: list[SourceDoc], excerpt_chars: int = 2500) -> list[dict[str, Any]]:
    out = []
    for d in docs:
        excerpt = shorten_text(d.extracted_text, excerpt_chars)
        out.append({
            "kind": d.kind,
            "label": d.label,
            "path": d.path,
            "excerpt": excerpt,
            "excerpt_word_count": count_words(excerpt),
            "fulltext_word_count": count_words(d.extracted_text),
            "bib_key": d.bib_key,
        })
    return out

def collect_extra_article_docs(cfg: dict[str, Any], max_files: int = 20) -> list[SourceDoc]:
    documento = cfg.get("documento", {})
    raw_items = documento.get("artigos_extras_paths", []) or []
    files = collect_readable_files(raw_items, get_config_base_dir(cfg))[:max_files]
    docs: list[SourceDoc] = []
    for path in files:
        try:
            docs.append(SourceDoc(
                path=str(path),
                kind="artigo_extra",
                label=path.name,
                extracted_text=read_text_file(path, 25000),
                metadata={"source_path": str(path)},
            ))
        except Exception as exc:
            debug_print(f"Falha ao ler artigo extra {path}: {exc}")
    return docs

def infer_final_front_matter(client: OpenAI, model: str, context: DocumentContext, org_text: str) -> FinalFrontMatterOutput:
    prompt = textwrap.dedent(
        f"""
        Gere os elementos finais de capa para um documento acadêmico em português.

        Contexto:
        {json.dumps(asdict(context), ensure_ascii=False, indent=2)}

        Conteúdo preliminar do documento:
        {shorten_text(org_text, 15000)}

        Regras:
        - title: título acadêmico final, claro e elegante.
        - documento_type: descrição curta do tipo do documento para a capa.
        - cover_note: nota curta, formal e compatível com a capa.
        """
    ).strip()
    resp = client.responses.parse(
        model=model,
        input=[{"role": "user", "content": prompt}],
        text_format=FinalFrontMatterOutput,
    )
    if resp.output_parsed is None:
        raise RuntimeError("A IA não retornou os elementos finais de capa.")
    return resp.output_parsed

def apply_final_front_matter(org_text: str, *, title: str, author: str, documento_type: str, cover_note: str, institution_name: str, course_name: str = "", discipline_name: str = "", professor_name: str = "", city_name: str = "Brasília") -> str:
    org_text = replace_org_header_line(org_text, "#+TITLE:", title)
    org_text = replace_org_header_line(org_text, "#+AUTHOR:", author)
    org_text = replace_latex_header_macro(org_text, "institution", institution_name)
    org_text = replace_latex_header_macro(org_text, "coursename", course_name)
    org_text = replace_latex_header_macro(org_text, "disciplinename", discipline_name)
    org_text = replace_latex_header_macro(org_text, "professorname", professor_name)
    org_text = replace_latex_header_macro(org_text, "cityname", city_name)
    org_text = replace_latex_header_macro(org_text, "papertype", documento_type)
    org_text = replace_latex_header_macro(org_text, "covernote", cover_note)
    return org_text

def force_activity_template_visible_title(org_text: str, title: str) -> str:
    """Força o título visual do template de atividade sem expor prefixo técnico.

    O template de atividade contém um bloco LaTeX cru com {\\Large\\bfseries ...}.
    Como esse bloco não passa pelo escape automático do Org, o texto precisa ser
    sanitizado para LaTeX.

    A substituição é feita por função, não por string de replacement do re.sub,
    porque comandos LaTeX como \\Large podem ser interpretados pelo mecanismo
    de regex como escape inválido (ex.: \\L).
    """
    clean = latex_escape_text(str(title or "").strip())
    if not clean:
        return org_text
    pattern = r"(?m)^\{\\Large\\bfseries\s+.*?\\par\}$"
    return re.sub(
        pattern,
        lambda _m: r"{\Large\bfseries " + clean + r"\par}",
        org_text,
        count=1,
    )

def _first_nonempty(*values: Any) -> str:
    for value in values:
        if value is None:
            continue
        text = str(value).strip()
        if text and not is_placeholder_value(text):
            return text
    return ""


def template_placeholder_values(cfg: dict[str, Any], final_title: str = "") -> dict[str, str]:
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    atividade = cfg.get("atividade", {}) if isinstance(cfg.get("atividade"), dict) else {}
    pesquisa = cfg.get("pesquisa", {}) if isinstance(cfg.get("pesquisa"), dict) else {}
    title = _first_nonempty(final_title, documento.get("titulo_trabalho"), documento.get("titulo"), atividade.get("titulo_trabalho"), atividade.get("titulo"), pesquisa.get("titulo"), pesquisa.get("tema"), "Atividade acadÃªmica")
    today = datetime.now().strftime("%d/%m/%Y")
    raw_values = {
        "TITULO_TRABALHO": title,
        "TITULO": title,
        "TITLE": title,
        "AUTOR": _first_nonempty(documento.get("autor"), atividade.get("aluno"), atividade.get("autor"), DEFAULT_AUTHOR),
        "AUTHOR": _first_nonempty(documento.get("autor"), atividade.get("aluno"), atividade.get("autor"), DEFAULT_AUTHOR),
        "ALUNO": _first_nonempty(atividade.get("aluno"), documento.get("aluno"), documento.get("autor"), DEFAULT_AUTHOR),
        "ALUNOS": _first_nonempty(atividade.get("alunos"), atividade.get("aluno"), documento.get("aluno"), DEFAULT_AUTHOR),
        "DATA": _first_nonempty(documento.get("data"), atividade.get("data"), today),
        "CURSO": _first_nonempty(documento.get("curso"), atividade.get("curso"), documento.get("course_name")),
        "TURMA": _first_nonempty(documento.get("turma"), atividade.get("turma")),
        "POLO": _first_nonempty(documento.get("polo"), atividade.get("polo"), documento.get("cidade"), "BrasÃ­lia"),
        "PÃLO": _first_nonempty(documento.get("polo"), atividade.get("polo"), documento.get("cidade"), "BrasÃ­lia"),
        "DISCIPLINA": _first_nonempty(documento.get("disciplina"), atividade.get("disciplina")),
        "PROFESSOR": _first_nonempty(documento.get("professor"), atividade.get("professor")),
        "TEMA": _first_nonempty(documento.get("tema"), atividade.get("tema"), pesquisa.get("tema")),
        "CONTEUDO": "",
        "CONTEÃDO": "",
    }
    return {k: latex_escape_text(v) for k, v in raw_values.items()}


def render_template_placeholders(org_text: str, cfg: dict[str, Any], final_title: str = "") -> str:
    values = template_placeholder_values(cfg, final_title=final_title)
    out = org_text
    for key, value in values.items():
        out = out.replace("{{" + key + "}}", value)
        out = out.replace("@" + key + "@", value)
    out = re.sub(r"\{\{[^{}]{1,80}\}\}", "", out)
    out = re.sub(r"@[A-ZÃÃÃÃÃÃÃÃÃÃÃ0-9_ -]{2,80}@", "", out)
    return out


def normalize_snippet_placeholders(text: str) -> str:
    """Normaliza placeholders de snippets Emacs/Yasnippet.

    O template pode vir com marcadores do tipo ``${1:Título}``. Eles são
    úteis no Emacs, mas não devem chegar ao .org final nem ao prompt da IA.
    """
    def repl(match: re.Match[str]) -> str:
        return match.group(1).strip()

    text = re.sub(r"\$\{\d+:([^{}]*)\}", repl, text or "")
    text = re.sub(r"\$\{\d+\}", "", text)
    return text


def apply_paper_template_metadata(org_text: str, cfg: dict[str, Any], final_title: str = "") -> str:
    """Aplica metadados determinísticos ao template fgv-paper.

    O estilo ``fgv-paper`` imprime ``\programname{}`` e ``\coursename{}``
    como linhas separadas. Portanto, para papers de disciplina, ``programname``
    deve permanecer vazio salvo se o TOML o preencher explicitamente. Essa
    rotina evita a duplicidade do nome do mestrado na capa.
    """
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    atividade = cfg.get("atividade", {}) if isinstance(cfg.get("atividade"), dict) else {}

    if normalize_document_type(documento.get("tipo_documento")) != "paper":
        return org_text

    if not bool(documento.get("normalizar_capa_paper", True)):
        return org_text

    title = _first_nonempty(
        final_title,
        documento.get("titulo_trabalho"),
        documento.get("titulo"),
        atividade.get("titulo_trabalho"),
        atividade.get("titulo"),
    )
    author = _first_nonempty(
        documento.get("autor"),
        documento.get("author"),
        atividade.get("aluno"),
        atividade.get("autor"),
        DEFAULT_AUTHOR,
    )
    institution = _first_nonempty(
        documento.get("institution_name"),
        documento.get("instituicao"),
        atividade.get("instituicao"),
        DEFAULT_INSTITUTION,
    )
    program = str(documento.get("program_name") or documento.get("programa") or "").strip()
    course = _first_nonempty(
        documento.get("course_name"),
        documento.get("curso"),
        atividade.get("curso"),
    )
    discipline = _first_nonempty(
        documento.get("discipline_name"),
        documento.get("disciplina"),
        atividade.get("disciplina"),
    )
    professor = _first_nonempty(
        documento.get("professor_name"),
        documento.get("professor"),
        atividade.get("professor"),
    )
    city = _first_nonempty(
        documento.get("city_name"),
        documento.get("cidade"),
        atividade.get("polo"),
        "Brasília",
    )
    paper_type = _first_nonempty(
        documento.get("papertype"),
        documento.get("documento_type"),
        documento.get("tipo_capa"),
        "Paper acadêmico",
    )
    cover_note = _first_nonempty(
        documento.get("covernote"),
        documento.get("cover_note"),
        documento.get("nota_capa"),
        "Trabalho acadêmico elaborado para a disciplina.",
    )

    if title:
        org_text = replace_org_header_line(org_text, "#+TITLE:", title)
    if author:
        org_text = replace_org_header_line(org_text, "#+AUTHOR:", author)

    org_text = replace_or_insert_latex_header_macro(org_text, "institution", latex_escape_text(institution))
    org_text = replace_or_insert_latex_header_macro(org_text, "programname", latex_escape_text(program))
    org_text = replace_or_insert_latex_header_macro(org_text, "coursename", latex_escape_text(course))
    org_text = replace_or_insert_latex_header_macro(org_text, "disciplinename", latex_escape_text(discipline))
    org_text = replace_or_insert_latex_header_macro(org_text, "professorname", latex_escape_text(professor))
    org_text = replace_or_insert_latex_header_macro(org_text, "cityname", latex_escape_text(city))
    org_text = replace_or_insert_latex_header_macro(org_text, "papertype", latex_escape_text(paper_type))
    org_text = replace_or_insert_latex_header_macro(org_text, "covernote", latex_escape_text(cover_note))

    if "\\usepapercover" not in org_text:
        org_text = insert_latex_header_line(org_text, "#+LATEX_HEADER: \\usepapercover")
    return org_text


TECHNICAL_LEAK_TERMS = [
    "metadados incompletos",
    "metadados inferidos",
    "metadados bibliográficos inferidos",
    "metadados bibliograficos inferidos",
    "cache local",
    "fulltext_cache",
    "limitação de acesso",
    "limitacao de acesso",
    "extração textual",
    "extracao textual",
    "ocr",
    "documentos processados",
    "documentos replicados",
    "material fornecido pelo professor",
    "fornecido pelo professor",
    "limitação técnica",
    "limitacao tecnica",
    "pipeline",
]


def technical_leak_terms_found(text: str) -> list[str]:
    low = normalize_title_loose(text or "")
    found: list[str] = []
    for term in TECHNICAL_LEAK_TERMS:
        if normalize_title_loose(term) in low:
            found.append(term)
    return found


def remove_technical_leaks_from_org(org_text: str, cfg: dict[str, Any]) -> str:
    """Remove parágrafos que vazam termos técnicos do processamento.

    O prompt instrui a IA a não mencionar cache/metadados/OCR, mas a blindagem
    precisa ser determinística. Esta função remove apenas parágrafos textuais,
    preservando cabeçalhos, diretivas Org e comandos LaTeX.
    """
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    if not bool(documento.get("remover_mencoes_tecnicas_pipeline", True)):
        return org_text

    pieces = re.split(r"(\n\s*\n)", org_text.replace("\r\n", "\n"))
    kept: list[str] = []
    for piece in pieces:
        stripped = piece.lstrip()
        is_control = stripped.startswith("#+") or stripped.startswith("*") or stripped.startswith("\\") or stripped.startswith(":")
        if not is_control and technical_leak_terms_found(piece):
            continue
        kept.append(piece)
    text = "".join(kept)

    methodology_replacement = (
        "O paper adota abordagem bibliográfica e argumentativa, articulando os textos-base "
        "do corpus local ao problema formulado. A análise prioriza os artigos formalmente "
        "selecionados e os mobiliza de modo integrado, sem recorrer a fontes externas como "
        "base principal. A estratégia metodológica consiste em reconstruir relações conceituais "
        "entre instituições políticas, federalismo, representação e capacidades estatais, "
        "preservando a coerência entre pergunta de pesquisa, tese central e bibliografia mobilizada.\n"
    )
    text = re.sub(
        r"(?ms)(^\*+\s+(?:\d+(?:\.\d+)?\s+)?Metodologia\s*\n)(.*?)(?=^\*+\s+|\Z)",
        lambda m: m.group(1) + methodology_replacement + "\n",
        text,
        count=1,
    )
    return re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"


def textcite_sequence(keys_csv: str) -> str:
    keys = [k.strip() for k in keys_csv.split(",") if k.strip()]
    if not keys:
        return ""
    parts = [r"\textcite{" + key + "}" for key in keys]
    if len(parts) == 1:
        return parts[0]
    if len(parts) == 2:
        return parts[0] + " e " + parts[1]
    return ", ".join(parts[:-1]) + " e " + parts[-1]


def polish_narrative_latex_citations(org_text: str, cfg: dict[str, Any]) -> str:
    """Transforma usos narrativos ruins de \parencite em \textcite.

    Ex.: ``Contra o diagnóstico..., \parencite{a,b} mostram`` vira
    ``Contra o diagnóstico..., \textcite{a} e \textcite{b} mostram``.
    """
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    if not bool(documento.get("corrigir_citacoes_narrativas", True)):
        return org_text

    verb_re = r"(mostra|mostram|evidencia|evidenciam|sugere|sugerem|argumenta|argumentam|indica|indicam|defende|defendem|sustenta|sustentam)"
    org_text = re.sub(
        rf",\s*\\parencite\{{([^}}]+)\}}\s+{verb_re}",
        lambda m: ", " + textcite_sequence(m.group(1)) + " " + m.group(2),
        org_text,
        flags=re.IGNORECASE,
    )
    org_text = re.sub(
        r"(À luz de|A luz de|Segundo|Conforme|Como sugerem|Como sugere|Como indicam|Como indica)\s+\\parencite\{([^}]+)\}",
        lambda m: m.group(1) + " " + textcite_sequence(m.group(2)),
        org_text,
        flags=re.IGNORECASE,
    )
    org_text = re.sub(
        r"\b(de|por|em)\s+\\parencite\{([^}]+)\}",
        lambda m: m.group(1) + " " + textcite_sequence(m.group(2)),
        org_text,
        flags=re.IGNORECASE,
    )
    return org_text


def validate_final_org_or_raise(org_text: str, bib_keys: list[str], cfg: dict[str, Any]) -> None:
    """Valida o .org final antes de gravar/compilar."""
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    if not bool(documento.get("validar_org_final", True)):
        return

    errors: list[str] = []

    if bool(documento.get("falhar_se_org_tiver_empty_citation", True)) and "<empty citation>" in org_text:
        errors.append("O ORG final ainda contém <empty citation>.")

    if "[cite:" in org_text or "[cite/" in org_text:
        errors.append("O ORG final ainda contém citações Org Cite não convertidas.")

    if bool(documento.get("falhar_se_org_tiver_mencao_tecnica", True)):
        leaks = technical_leak_terms_found(org_text)
        if leaks:
            errors.append("O ORG final contém menções técnicas proibidas: " + ", ".join(leaks))

    cited_groups = re.findall(
        r"\\(?:paren|text|auto|foot|smart|cite|citeauthor|citeyear)cite\{([^}]+)\}",
        org_text,
    )
    used_keys: set[str] = set()
    for group in cited_groups:
        used_keys.update(k.strip() for k in group.split(",") if k.strip())

    known = set(bib_keys or [])
    missing = sorted(k for k in used_keys if k not in known)
    if missing:
        errors.append("Há chaves citadas que não existem no .bib: " + ", ".join(missing))

    if bool(documento.get("falhar_se_org_tiver_chave_crua", True)) and known:
        masked = re.sub(
            r"\\(?:paren|text|auto|foot|smart|cite|citeauthor|citeyear)cite\{[^}]+\}",
            "",
            org_text,
        )
        raw_leaks = sorted(k for k in known if re.search(rf"\b{re.escape(k)}\b", masked))
        if raw_leaks:
            errors.append("Há chaves BibTeX cruas no texto: " + ", ".join(raw_leaks[:20]))

    program_match = re.search(r"\\programname\{([^}]*)\}", org_text)
    course_match = re.search(r"\\coursename\{([^}]*)\}", org_text)
    if program_match and course_match:
        program = normalize_title_loose(program_match.group(1))
        course = normalize_title_loose(course_match.group(1))
        if program and course and program == course:
            errors.append("\\programname{} e \\coursename{} estão iguais; isso duplicará o curso na capa.")

    if errors:
        raise RuntimeError("Validação final do ORG falhou:\n- " + "\n- ".join(errors))


def build_paper_prompt(
    cfg: dict[str, Any],
    context: DocumentContext,
    template_text: str,
    base_docs: list[SourceDoc],
    guidance_docs: list[SourceDoc],
    bib_keys: list[str],
    bib_entries: list[str],
    style: str,
    selected_corpus_catalog: list[dict[str, Any]] | None = None,
) -> str:
    documento = cfg.get("documento", {})
    doc_type_local = normalize_document_type(documento.get("tipo_documento"))
    modo_escrita = str(documento.get("modo_escrita") or "novo").strip().lower()
    perfil_redacao = str(documento.get("perfil_redacao") or "academico_equilibrado").strip().lower()
    priorizar_citacoes = bool(documento.get("priorizar_citacoes_dos_selecionados", True))
    usar_contexto_consolidado = bool(documento.get("usar_contexto_consolidado_da_pesquisa", True))
    extras_so_complementam = bool(documento.get("extras_so_complementam", True))
    preservar_estrutura = bool(documento.get("preservar_estrutura_do_org_anterior", False))

    selected_keys = sorted({d.bib_key for d in base_docs if d.kind.startswith("texto_selecionado") and d.bib_key})
    extra_keys = sorted({d.bib_key for d in base_docs if d.kind == "artigo_extra" and d.bib_key})
    excerpt_chars = 4200 if doc_type_local == "dissertacao" else 2500
    targets = dissertation_generation_targets(cfg)

    instrucoes_modo = {
        "novo": "Escreva um documento novo, mas usando a pesquisa consolidada como base metodológica e analítica.",
        "reescrever": "Reescreva o documento com maior coesão, clareza e densidade argumentativa, aproveitando o documento anterior se ele estiver entre as orientações.",
        "expandir": "Mantenha a linha argumentativa principal e expanda o documento com mais desenvolvimento analítico, usando os textos extras quando úteis.",
    }.get(modo_escrita, "Escreva um documento acadêmico coeso a partir da pesquisa consolidada.")

    instrucoes_perfil = {
        "academico_equilibrado": "Adote redação acadêmica equilibrada, com boa relação entre síntese, discussão teórica e fluidez.",
        "mais_teorico": "Privilegie densidade conceitual, diálogo teórico e elaboração interpretativa mais forte.",
        "mais_discursivo": "Privilegie fluidez discursiva, encadeamento argumentativo e texto mais ensaístico, sem perder rigor acadêmico.",
        "mais_sintetico": "Privilegie concisão, objetividade e alta compressão informacional, evitando expansões desnecessárias.",
    }.get(perfil_redacao, "Adote redação acadêmica equilibrada e coesa.")

    regra_citacao = "Priorize as citações dos artigos formalmente selecionados na pesquisa." if priorizar_citacoes else "Use as referências de forma equilibrada, sem obrigação de priorizar os selecionados."
    regra_extras = "Os artigos extras só podem complementar a argumentação; não substituem o papel central dos artigos formalmente selecionados." if extras_so_complementam else "Os artigos extras podem ter papel relevante, desde que não descaracterizem o núcleo da pesquisa."
    regra_estrutura = "Se houver um org anterior do documento entre as orientações, preserve sua estrutura principal e melhore o texto dentro dessa arquitetura." if preservar_estrutura else "Se houver um org anterior do documento entre as orientações, use-o apenas como referência, sem obrigação de preservar a mesma estrutura."
    contexto_regra = (
        "Tema, recorte e objetivo já consolidados pela pesquisa devem ser tratados como referência principal."
        if usar_contexto_consolidado else
        "É permitido trabalhar com formulação refinada do tema, recorte e objetivo para a escrita do documento, sem mudar o núcleo temático."
    )

    limites = {
        "limite_palavras_total": documento.get("limite_palavras_total"),
        "limite_palavras_introducao": documento.get("limite_palavras_introducao"),
        "limite_palavras_revisao": documento.get("limite_palavras_revisao"),
        "limite_palavras_conclusao": documento.get("limite_palavras_conclusao"),
    }
    limites_dict = {k: v for k, v in limites.items() if v not in (None, "", 0)}
    limites_txt = json.dumps(limites_dict, ensure_ascii=False, indent=2) if limites_dict else "Nenhum limite explícito informado."

    if is_local_documents_mode(cfg):
        cobertura_fulltext_txt = (
            "13. Este documento deve ser construído a partir dos documentos locais fornecidos no ZIP/pasta e copiados para *_fulltext_cache/. Todos os textos disponíveis nesse cache devem ser lidos, interpretados e mobilizados substantivamente ao longo da análise; não use fontes externas como base principal."
        )
    else:
        cobertura_fulltext_txt = (
            "13. A cobertura bibliográfica obrigatória é definida pelos PDFs/textos efetivamente baixados no diretório *_fulltext_cache/: todas as chaves mapeadas a esses arquivos devem ser interpretadas e citadas ao longo do texto, de forma substantiva e não meramente ornamental."
        )

    dissertation_rules = ""
    dissertation_targets = ""
    if doc_type_local == "dissertacao":
        dissertation_rules = textwrap.dedent(f"""
        18. Este trabalho deve ter fôlego de dissertação de mestrado baseada em revisão de literatura robusta, e não de paper curto ou resumo expandido.
        19. Desenvolva em profundidade os conceitos, os debates teóricos, os contrastes entre autores, as lacunas da literatura e as implicações para o governo público federal brasileiro.
        20. Absorva de modo explícito o material já consolidado na pesquisa, aproveitando a massa crítica de artigos selecionados e seus resumos/trechos.
        21. Amplie especialmente as seções de referencial teórico, metodologia e resultados/discussão; evite compressão excessiva.
        22. Em vez de apenas resumir autores, compare abordagens, explicite convergências e divergências e construa sínteses analíticas próprias.
        23. Sempre que houver literatura suficiente, subdivida a discussão em eixos temáticos densos e bem desenvolvidos, com transições argumentativas claras.
        24. Não reduza a metodologia a um parágrafo curto: explique estratégia de busca, critérios, corpus, limites e racionalidade analítica em densidade compatível com dissertação.
        25. Em resultados e discussão, explore de forma aprofundada aplicações da IA, modelos de governança, ética, accountability, compliance, riscos, capacidades estatais, assimetrias institucionais e implicações para a APF.
        26. Não encurte o documento para caber em poucas páginas. O alvo de desenvolvimento textual é compatível com um corpo principal de dissertação.
        27. Metas mínimas desejáveis de extensão textual por bloco analítico:
        {json.dumps(targets, ensure_ascii=False, indent=2)}
        """).strip()
        dissertation_targets = textwrap.dedent("""
        Estrutura desejável para dissertação:
        - Introdução substantiva, com contextualização, delimitação, problema, objetivos e justificativa bem desenvolvidos.
        - Referencial teórico aprofundado, com múltiplos subtópicos e comparação entre autores e correntes.
        - Metodologia detalhada, absorvendo os elementos relevantes da pesquisa PRISMA já realizada.
        - Resultados e discussão como núcleo do trabalho, com desenvolvimento significativamente mais longo do que introdução e conclusão.
        - Considerações finais com síntese, limitações e agenda futura bem explicitadas.
        """).strip()

    return textwrap.dedent(
        f"""
        Gere um documento acadêmico completo em Org-mode.

        Regras obrigatórias:
        1. Preserve o cabeçalho técnico do template (linhas #+...).
        2. O documento deve ser escrito em português, em tom acadêmico, argumentativo e coeso.
        3. Use apenas citações com as chaves bibliográficas fornecidas.
        4. Não invente chaves bibliográficas.
        5. Use citações nativas do Org Cite, como [cite:@chave] ou [cite/t:@chave].
        6. Não escreva numeração manual nos títulos.
        7. Não crie manualmente uma seção de referências fora do mecanismo de bibliografia do template.
        8. Use o estilo bibliográfico final {style.upper()}.
        9. {instrucoes_modo}
        10. {instrucoes_perfil}
        11. {regra_citacao}
        12. {regra_extras}
        {cobertura_fulltext_txt}
        14. {contexto_regra}
        15. {regra_estrutura}
        16. Se os textos-base forem insuficientes em algum ponto, trate essa limitação de forma honesta.
        17. Respeite, quando possível, os limites de extensão abaixo:
        {limites_txt}
        {dissertation_rules}

        Contexto do documento:
        {json.dumps(asdict(context), ensure_ascii=False, indent=2)}

        Estrutura e profundidade esperadas:
        {dissertation_targets or 'Use a estrutura do template e desenvolva cada seção com densidade acadêmica apropriada.'}

        Chaves bibliográficas de artigos formalmente selecionados:
        {json.dumps(selected_keys, ensure_ascii=False, indent=2)}

        Chaves bibliográficas de artigos extras:
        {json.dumps(extra_keys, ensure_ascii=False, indent=2)}

        Catálogo do corpus selecionado da pesquisa:
        {shorten_text(json.dumps(selected_corpus_catalog or [], ensure_ascii=False, indent=2), 40000)}

        Template Org:
        {shorten_text(template_text, 30000)}

        Chaves bibliográficas disponíveis:
        {json.dumps(bib_keys, ensure_ascii=False, indent=2)}

        Entradas bib disponíveis:
        {shorten_text(json.dumps(bib_entries, ensure_ascii=False, indent=2), 25000)}

        Textos-base:
        {json.dumps(summarize_docs(base_docs, excerpt_chars=excerpt_chars), ensure_ascii=False, indent=2)}

        Orientações e artefatos da pesquisa:
        {json.dumps(summarize_docs(guidance_docs, excerpt_chars=3200), ensure_ascii=False, indent=2)}

        Retorne apenas o conteúdo completo do arquivo .org final.
        """
    ).strip()

def cleanup_placeholder_sections_in_tex_ready_text(text: str) -> str:
    text = re.sub(r"(?im)^\s*[,;:]+\s*$\n?", "", text)
    text = re.sub(r"(?im)^\s*P\d+\\?\}\s*$\n?", "", text)
    text = re.sub(r"(?im)^\s*<empty citation>\s*$\n?", "", text)
    return text


def normalize_org_citations(text: str) -> str:
    """Normaliza citações Org Cite geradas pela IA.

    Corrige principalmente:
    - [cite/t:@chave] e demais variantes;
    - múltiplas chaves sem espaçamento consistente;
    - restos como <empty citation>;
    - citações com chaves concatenadas quando a IA remove espaços.
    """
    text = (text or "").replace("<empty citation>", "")
    text = text.replace("[][]", "")

    cite_pat = re.compile(r"\[cite(?P<mode>/[A-Za-z]+)?:(?P<body>[^\]]+)\]")

    def repl(m: re.Match) -> str:
        mode = m.group("mode") or ""
        body = m.group("body") or ""
        keys = re.findall(r"@([A-Za-z0-9_:\-]+)", body)
        if not keys:
            return m.group(0)
        # Remove duplicatas mantendo ordem.
        keys = list(dict.fromkeys(k.strip() for k in keys if k.strip()))
        if not keys:
            return ""
        return "[cite%s:%s]" % (mode, "; ".join("@" + k for k in keys))

    return cite_pat.sub(repl, text)


def should_use_latex_citations(cfg: dict[str, Any]) -> bool:
    """Usa citações LaTeX diretas para PDF robusto.

    Em alguns ambientes Org/Emacs, o org-cite/biblatex fica instável em batch.
    Converter [cite:...] para \parencite{} e [cite/t:...] para \textcite{}
    evita que o PDF mostre chaves BibTeX cruas e garante que o Biber/BibLaTeX
    controle as referências.
    """
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    raw = documento.get("usar_citacoes_latex_diretas", documento.get("citacoes_latex_diretas", True))
    return bool(raw)


def convert_org_cites_to_latex(text: str) -> str:
    text = normalize_org_citations(text)
    cite_pat = re.compile(r"\[cite(?P<mode>/[A-Za-z]+)?:(?P<body>[^\]]+)\]")

    def repl(m: re.Match) -> str:
        mode = (m.group("mode") or "").lower()
        body = m.group("body") or ""
        keys = re.findall(r"@([A-Za-z0-9_:\-]+)", body)
        keys = list(dict.fromkeys(k.strip() for k in keys if k.strip()))
        if not keys:
            return ""
        key_csv = ",".join(keys)
        if mode in {"/t", "/text", "/textcite"}:
            return r"\textcite{" + key_csv + "}"
        return r"\parencite{" + key_csv + "}"

    return cite_pat.sub(repl, text)


def normalize_latex_bibliography_block(org_text: str) -> str:
    """Garante bibliografia final via LaTeX direto quando não usamos org-cite."""
    text = org_text.replace("\r\n", "\n")
    text = re.sub(r"(?im)^\s*#\+CITE_EXPORT:.*\n?", "", text)
    text = re.sub(r"(?im)^\s*#\+BIBLIOGRAPHY:.*\n?", "", text)
    text = re.sub(r"(?im)^\s*#\+PRINT_BIBLIOGRAPHY:?\s*(?:.*)?\n?", "", text)
    text = re.sub(r"(?im)^\s*#\+LATEX:\s*\\printbibliography\s*$\n?", "", text)
    text = re.sub(r"(?ims)\n*^\*+\s+(Refer[êe]ncias|References|Bibliography)\s*$.*\Z", "", text)
    return text.rstrip() + "\n\n#+LATEX: \\printbibliography\n"


def prepare_citations_for_pdf_export(org_text: str, cfg: dict[str, Any]) -> str:
    """Barreira final antes de gravar/compilar o Org."""
    if should_use_latex_citations(cfg):
        org_text = convert_org_cites_to_latex(org_text)
        org_text = normalize_latex_bibliography_block(org_text)
    else:
        org_text = normalize_org_citations(org_text)
        org_text = normalize_bibliography_block(org_text)
    return org_text

def strip_code_fences(text: str) -> str:
    text = (text or "").strip()
    text = re.sub(r"^```(?:org|markdown|md|text)?\n", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\n```$", "", text)
    return text.strip()


def sanitize_generated_org_fragment(text: str) -> str:
    text = strip_code_fences(text)
    text = text.replace("\r\n", "\n")
    text = re.sub(r"(?im)^\s*#\+.*$\n?", "", text)
    text = re.sub(r"(?im)^\s*<empty citation>\s*$\n?", "", text)
    text = text.replace("[][]", "")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def strip_redundant_leading_heading(text: str, heading_level: int, heading_title: str) -> str:
    body = sanitize_generated_org_fragment(text)
    if not body:
        return body
    stars = "*" * max(int(heading_level or 1), 1)
    wanted = normalize_heading_loose(heading_title)
    pattern = re.compile(rf"(?ms)^\s*{re.escape(stars)}\s+(?P<title>.+?)\s*\n")
    while True:
        m = pattern.match(body)
        if not m:
            break
        found = normalize_heading_loose(m.group("title"))
        if found != wanted:
            break
        body = body[m.end():].lstrip("\n")
    return body.strip()


def _find_section_body_span(org_text: str, heading_level: int, heading_title: str) -> tuple[int, int]:
    stars = "*" * heading_level
    pattern = re.compile(rf"(?m)^{re.escape(stars)}\s+{re.escape(heading_title)}\s*$")
    m = pattern.search(org_text)
    if not m:
        raise ValueError(f"Seção não encontrada: nivel={heading_level} titulo={heading_title!r}")
    pos = m.end()
    if pos < len(org_text) and org_text[pos] == "\n":
        pos += 1
    drawer = re.compile(r"(?ms)^:PROPERTIES:\n.*?\n:END:\n")
    dm = drawer.match(org_text[pos:])
    if dm:
        pos += dm.end()
    while pos < len(org_text) and org_text[pos] == "\n":
        pos += 1
    next_heading = re.compile(rf"(?m)^(?:\*{{1,{heading_level}}})\s+")
    nm = next_heading.search(org_text, pos)
    end = nm.start() if nm else len(org_text)
    return pos, end


def replace_section_body(org_text: str, heading_level: int, heading_title: str, new_body: str) -> str:
    start, end = _find_section_body_span(org_text, heading_level, heading_title)
    body = strip_redundant_leading_heading(new_body, heading_level, heading_title)
    prefix = org_text[:start].rstrip() + "\n\n"
    suffix = org_text[end:].lstrip("\n")
    if body:
        return prefix + body + "\n\n" + suffix
    return prefix + suffix


def extract_section_body(org_text: str, heading_level: int, heading_title: str) -> str:
    start, end = _find_section_body_span(org_text, heading_level, heading_title)
    return org_text[start:end].strip()


def extract_subheading_titles(section_body: str) -> list[str]:
    return [m.group(1).strip() for m in re.finditer(r"(?m)^\*\*\s+(.+)$", section_body)]


def build_bib_meta_lookup(bib_entries: list[str]) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for entry in bib_entries:
        meta = parse_bib_entry_meta(entry)
        key = str(meta.get("key") or "").strip()
        if key:
            out[key] = meta
    return out


def chunk_keys_weighted(keys: list[str], weights: list[float]) -> list[list[str]]:
    if not weights:
        return []
    if not keys:
        return [[] for _ in weights]
    total = float(sum(weights)) or 1.0
    raw_sizes = [len(keys) * (w / total) for w in weights]
    sizes = [int(x) for x in raw_sizes]
    remainder = len(keys) - sum(sizes)
    order = sorted(range(len(weights)), key=lambda i: (raw_sizes[i] - sizes[i]), reverse=True)
    for i in range(remainder):
        sizes[order[i % len(order)]] += 1
    chunks: list[list[str]] = []
    cursor = 0
    for size in sizes:
        chunks.append(keys[cursor:cursor + size])
        cursor += size
    if cursor < len(keys):
        chunks[-1].extend(keys[cursor:])
    return chunks


def build_corpus_overview(selected_corpus_catalog: list[dict[str, Any]] | None, limit: int = 24) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for item in (selected_corpus_catalog or [])[:limit]:
        out.append({
            "bib_key": item.get("bib_key"),
            "title": item.get("title"),
            "year": item.get("year"),
            "venue": item.get("venue"),
            "keywords": item.get("keywords") or [],
        })
    return out


def summarize_docs_by_keys(base_docs: list[SourceDoc], keys: list[str], excerpt_chars: int = 1200, max_docs: int = 8) -> list[dict[str, Any]]:
    chosen: list[SourceDoc] = []
    key_set = set(keys)
    for doc in base_docs:
        if doc.bib_key and doc.bib_key in key_set:
            chosen.append(doc)
    if not chosen:
        chosen = [d for d in base_docs if d.kind.startswith("texto_selecionado")][:max_docs]
    return summarize_docs(chosen[:max_docs], excerpt_chars=excerpt_chars)


def build_dissertation_front_sections_prompt(
    cfg: dict[str, Any],
    context: DocumentContext,
    template_text: str,
    guidance_docs: list[SourceDoc],
    selected_corpus_catalog: list[dict[str, Any]] | None = None,
) -> str:
    atividade = cfg.get("atividade", {})
    documento = cfg.get("documento", {})
    return textwrap.dedent(
        f"""
        Gere os elementos pré-textuais de uma dissertação acadêmica em português, em formato estruturado.

        Regras:
        1. Não use placeholders genéricos como "[...]", "Autor da epígrafe" ou "A definir", salvo quando não houver dado objetivo para nomes de banca/orientação, os quais não fazem parte desta etapa.
        2. DEDICATÓRIA e AGRADECIMENTOS devem ser breves, elegantes e plausíveis.
        3. A EPÍGRAFE é opcional; se não houver boa opção, devolva epigrafe_texto e epigrafe_autor vazios.
        4. O RESUMO deve ter densidade acadêmica e refletir a dissertação de revisão de literatura.
        5. O ABSTRACT deve ser tradução acadêmica consistente do resumo.
        6. As listas devem ser enxutas e compatíveis com o documento.
        7. O glossário deve conter termos realmente úteis ao tema.

        Contexto consolidado:
        {json.dumps(asdict(context), ensure_ascii=False, indent=2)}

        Metadados do trabalho:
        {json.dumps({
            "aluno": atividade.get("aluno"),
            "curso": atividade.get("curso"),
            "disciplina": atividade.get("disciplina"),
            "professor": atividade.get("professor"),
            "institution_name": documento.get("institution_name"),
            "school_name": documento.get("school_name"),
        }, ensure_ascii=False, indent=2)}

        Estrutura do template:
        {shorten_text(template_text, 12000)}

        Catálogo do corpus selecionado:
        {json.dumps(build_corpus_overview(selected_corpus_catalog), ensure_ascii=False, indent=2)}

        Orientações e artefatos relevantes:
        {json.dumps(summarize_docs(guidance_docs, excerpt_chars=1800), ensure_ascii=False, indent=2)}
        """
    ).strip()


def build_dissertation_stage_plan(template_text: str, cfg: dict[str, Any], bib_keys: list[str]) -> list[dict[str, Any]]:
    targets = dissertation_generation_targets(cfg)
    citation_target_keys = list(cfg.get("__citation_target_keys__") or bib_keys)
    top_sections = [
        ("INTRODUÇÃO", "introducao", targets.get("introducao_min", 1200)),
        ("REFERENCIAL TEÓRICO", "referencial_teorico", targets.get("referencial_teorico_min", 4000)),
        ("METODOLOGIA", "metodologia", targets.get("metodologia_min", 1800)),
        ("RESULTADOS E DISCUSSÃO", "resultados_discussao", targets.get("resultados_discussao_min", 5000)),
        ("CONSIDERAÇÕES FINAIS", "conclusao", targets.get("conclusao_min", 1200)),
    ]
    plan: list[dict[str, Any]] = []
    subsection_counts: list[int] = []
    by_top: list[list[str]] = []
    for top_title, _, _ in top_sections:
        body = extract_section_body(template_text, 1, top_title)
        subs = extract_subheading_titles(body)
        if not subs:
            subs = [top_title]
        by_top.append(subs)
        subsection_counts.append(len(subs))
    key_chunks = chunk_keys_weighted(citation_target_keys, [max(c, 1) for c in subsection_counts])
    chunk_idx = 0
    for (top_title, bucket, bucket_min), subs in zip(top_sections, by_top):
        per_sub_min = max(220, int(bucket_min / max(len(subs), 1)))
        template_body = extract_section_body(template_text, 1, top_title)
        siblings = subs[:]
        for sub_title in subs:
            mandatory_keys = key_chunks[chunk_idx] if chunk_idx < len(key_chunks) else []
            chunk_idx += 1
            plan.append({
                "top_title": top_title,
                "sub_title": sub_title,
                "bucket": bucket,
                "min_words": per_sub_min,
                "mandatory_keys": mandatory_keys,
                "all_keys": bib_keys,
                "citation_target_keys": citation_target_keys,
                "template_body": template_body,
                "sibling_subtitles": siblings,
            })
    return plan


def build_dissertation_subsection_prompt(
    cfg: dict[str, Any],
    context: DocumentContext,
    stage: dict[str, Any],
    bib_meta_lookup: dict[str, dict[str, Any]],
    base_docs: list[SourceDoc],
    guidance_docs: list[SourceDoc],
    selected_corpus_catalog: list[dict[str, Any]] | None = None,
    previous_excerpt: str = "",
) -> str:
    top_title = stage["top_title"]
    sub_title = stage["sub_title"]
    mandatory_keys = stage.get("mandatory_keys", [])
    all_keys = stage.get("all_keys", [])
    min_words = int(stage.get("min_words", 320))
    sibling_subtitles = stage.get("sibling_subtitles", [])
    focus_map = {
        "INTRODUÇÃO": "Contextualize o problema, delimite o objeto, explicite problema, objetivos e justificativa com densidade de dissertação.",
        "REFERENCIAL TEÓRICO": "Aprofunde conceitos, compare autores, explicite convergências, divergências e lacunas, evitando resumo superficial.",
        "METODOLOGIA": "Explique em detalhe o desenho da revisão, o corpus, critérios, bases, limites e racionalidade analítica.",
        "RESULTADOS E DISCUSSÃO": "Faça análise densa do corpus, compare abordagens, desenvolva implicações para a APF e articule achados com a literatura.",
        "CONSIDERAÇÕES FINAIS": "Sintetize criticamente, explicite limites e desdobre agenda futura sem mera repetição mecânica.",
    }
    bib_subset = [bib_meta_lookup[k] for k in mandatory_keys if k in bib_meta_lookup]
    docs_subset = summarize_docs_by_keys(base_docs, mandatory_keys, excerpt_chars=1000, max_docs=8)
    return textwrap.dedent(
        f"""
        Escreva APENAS o corpo em Org-mode da subseção ** {sub_title} que pertence à seção * {top_title}.

        Regras obrigatórias:
        1. NÃO repita o título de nível 1.
        2. Retorne o texto iniciando com a heading de nível 2 correspondente, isto é: ** {sub_title}
        3. Use exclusivamente citações nativas do Org Cite, como [cite:@chave].
        4. Não invente chaves bibliográficas.
        5. Desenvolva a subseção com densidade compatível com dissertação, evitando síntese curta.
        6. Meta mínima de extensão desta subseção: cerca de {min_words} palavras.
        7. Nesta subseção, use de forma orgânica e explícita estas chaves obrigatórias, cada uma ao menos uma vez, salvo incompatibilidade temática absoluta: {json.dumps(mandatory_keys, ensure_ascii=False)}.
        8. O projeto global exige que TODOS os PDFs/textos do diretório *_fulltext_cache/ sejam interpretados pela IA e distribuídos ao longo da dissertação. Aqui, sua responsabilidade principal é usar os documentos resumidos e metadados desta etapa para construir argumento substantivo, comparativo e citado, não apenas mencionar autores.
        9. Subtítulos irmãos desta seção para manter coerência global: {json.dumps(sibling_subtitles, ensure_ascii=False)}.
        10. {focus_map.get(top_title, '')}
        11. Evite listas excessivas; priorize parágrafos analíticos e comparativos.

        Contexto consolidado:
        {json.dumps(asdict(context), ensure_ascii=False, indent=2)}

        Catálogo do corpus selecionado:
        {json.dumps(build_corpus_overview(selected_corpus_catalog), ensure_ascii=False, indent=2)}

        Metadados bibliográficos prioritários para esta subseção:
        {json.dumps(bib_subset, ensure_ascii=False, indent=2)}

        Documentos resumidos mais relevantes para esta subseção:
        {json.dumps(docs_subset, ensure_ascii=False, indent=2)}

        Chaves globais disponíveis em toda a dissertação:
        {json.dumps(all_keys, ensure_ascii=False, indent=2)}

        Estrutura da seção no template:
        {shorten_text(stage.get('template_body', ''), 6000)}

        Último contexto já redigido para manter continuidade:
        {shorten_text(previous_excerpt, 5000)}

        Orientações e artefatos relevantes:
        {json.dumps(summarize_docs(guidance_docs[:6], excerpt_chars=1400), ensure_ascii=False, indent=2)}
        """
    ).strip()


def build_dissertation_subsection_repair_prompt(
    cfg: dict[str, Any],
    context: DocumentContext,
    stage: dict[str, Any],
    current_body: str,
    missing_keys: list[str],
    bib_meta_lookup: dict[str, dict[str, Any]],
    base_docs: list[SourceDoc],
) -> str:
    docs_subset = summarize_docs_by_keys(base_docs, missing_keys or stage.get("mandatory_keys", []), excerpt_chars=1000, max_docs=8)
    metas = [bib_meta_lookup[k] for k in missing_keys if k in bib_meta_lookup]
    return textwrap.dedent(
        f"""
        Reescreva a subseção abaixo em Org-mode, preservando sua linha argumentativa, mas corrigindo duas insuficiências: falta de cobertura bibliográfica obrigatória e/ou extensão abaixo do desejado.

        Subseção: ** {stage['sub_title']}
        Seção: * {stage['top_title']}
        Meta mínima aproximada de palavras: {stage.get('min_words', 320)}
        Chaves obrigatórias que ainda precisam aparecer organicamente: {json.dumps(missing_keys, ensure_ascii=False)}
        Use apenas Org Cite ([cite:@chave]).
        Retorne o texto completo da subseção, começando com ** {stage['sub_title']}.

        Contexto:
        {json.dumps(asdict(context), ensure_ascii=False, indent=2)}

        Metadados das chaves faltantes:
        {json.dumps(metas, ensure_ascii=False, indent=2)}

        Documentos resumidos para apoiar a revisão:
        {json.dumps(docs_subset, ensure_ascii=False, indent=2)}

        Subseção atual:
        {shorten_text(current_body, 12000)}
        """
    ).strip()


def maybe_repair_dissertation_subsection(
    client: OpenAI,
    model: str,
    cfg: dict[str, Any],
    context: DocumentContext,
    stage: dict[str, Any],
    body: str,
    bib_meta_lookup: dict[str, dict[str, Any]],
    base_docs: list[SourceDoc],
) -> tuple[str, str | None, dict[str, Any]]:
    cited = set(extract_cited_keys_from_org(body))
    mandatory = [k for k in stage.get("mandatory_keys", []) if k]
    missing = [k for k in mandatory if k not in cited]
    words = count_words(body)
    min_words = int(stage.get("min_words", 320))
    needs_repair = bool(missing or words < int(min_words * 0.75))
    diagnostics = {
        "top_title": stage.get("top_title"),
        "sub_title": stage.get("sub_title"),
        "word_count": words,
        "min_words": min_words,
        "mandatory_keys": mandatory,
        "missing_keys": missing,
        "needs_repair": needs_repair,
    }
    if not needs_repair:
        return body, None, diagnostics
    prompt = build_dissertation_subsection_repair_prompt(cfg, context, stage, body, missing, bib_meta_lookup, base_docs)
    resp = client.responses.parse(
        model=model,
        input=[{"role": "user", "content": prompt}],
        text_format=DissertationStageBodyOutput,
    )
    if resp.output_parsed is None:
        raise RuntimeError(f"A IA não retornou a subseção reparada para {stage.get('sub_title')}.")
    repaired = sanitize_generated_org_fragment(resp.output_parsed.org_body)
    return repaired, prompt, diagnostics


def build_missing_bibliography_patch_prompt(
    cfg: dict[str, Any],
    context: DocumentContext,
    missing_keys: list[str],
    bib_meta_lookup: dict[str, dict[str, Any]],
    base_docs: list[SourceDoc],
    current_org: str,
) -> str:
    docs_subset = summarize_docs_by_keys(base_docs, missing_keys, excerpt_chars=1000, max_docs=12)
    metas = [bib_meta_lookup[k] for k in missing_keys if k in bib_meta_lookup]
    return textwrap.dedent(
        f"""
        O documento acadêmico abaixo ainda não utilizou todas as chaves bibliográficas exigidas dos artigos com PDF no diretório *_fulltext_cache/. Gere APENAS uma nova subseção de Org-mode para ser inserida em * RESULTADOS E DISCUSSÃO, com heading de nível 2, intitulada ** Síntese integradora ampliada do corpus.

        Regras:
        1. Use organicamente TODAS estas chaves faltantes ao menos uma vez: {json.dumps(missing_keys, ensure_ascii=False)}.
        2. Use apenas citações Org Cite ([cite:@chave]).
        3. A subseção deve ser analítica, comparativa e aprofundada, não uma lista mecânica de referências.
        4. Interprete os documentos resumidos correspondentes e conecte as referências faltantes aos eixos de governança, ética, compliance, risco, capacidades estatais e implicações para a APF, conforme pertinente.
        5. Não trate a citação como adereço: para cada chave faltante, incorpore pelo menos uma contribuição substantiva do texto ao argumento.

        Contexto:
        {json.dumps(asdict(context), ensure_ascii=False, indent=2)}

        Metadados bibliográficos das chaves faltantes:
        {json.dumps(metas, ensure_ascii=False, indent=2)}

        Documentos resumidos correspondentes:
        {json.dumps(docs_subset, ensure_ascii=False, indent=2)}

        Documento atual:
        {shorten_text(current_org, 20000)}
        """
    ).strip()


def build_dissertation_org_from_stages(
    client: OpenAI,
    model: str,
    cfg: dict[str, Any],
    context: DocumentContext,
    template_text: str,
    base_docs: list[SourceDoc],
    guidance_docs: list[SourceDoc],
    bib_filename: str,
    bib_entries: list[str],
    bib_keys: list[str],
    style: str,
    selected_corpus_catalog: list[dict[str, Any]] | None = None,
) -> tuple[str, str, dict[str, Any]]:
    prompt_parts: list[str] = []
    stage_diagnostics: list[dict[str, Any]] = []
    working_org = template_text

    front_prompt = build_dissertation_front_sections_prompt(cfg, context, template_text, guidance_docs, selected_corpus_catalog=selected_corpus_catalog)
    prompt_parts.append("===== front_sections =====\n" + front_prompt + "\n")
    front_resp = client.responses.parse(
        model=model,
        input=[{"role": "user", "content": front_prompt}],
        text_format=DissertationFrontSectionsOutput,
    )
    if front_resp.output_parsed is None:
        raise RuntimeError("A IA não retornou os elementos pré-textuais da dissertação.")
    front = front_resp.output_parsed

    if front.dedicatoria.strip():
        working_org = replace_section_body(working_org, 1, "DEDICATÓRIA", front.dedicatoria.strip())
    if front.agradecimentos.strip():
        working_org = replace_section_body(working_org, 1, "AGRADECIMENTOS", front.agradecimentos.strip())
    if front.epigrafe_texto.strip():
        ep = front.epigrafe_texto.strip()
        if front.epigrafe_autor.strip():
            ep += "\n" + front.epigrafe_autor.strip()
        working_org = replace_section_body(working_org, 1, "EPÍGRAFE", ep)
    else:
        working_org = replace_section_body(working_org, 1, "EPÍGRAFE", "")
    resumo_body = front.resumo.strip()
    if front.palavras_chave:
        resumo_body += "\n\nPalavras-chave: " + "; ".join([p.strip() for p in front.palavras_chave if str(p).strip()]) + "."
    working_org = replace_section_body(working_org, 1, "RESUMO", resumo_body)
    abstract_body = front.abstract.strip()
    if front.keywords:
        abstract_body += "\n\nKeywords: " + "; ".join([p.strip() for p in front.keywords if str(p).strip()]) + "."
    working_org = replace_section_body(working_org, 1, "ABSTRACT", abstract_body)
    working_org = replace_section_body(working_org, 1, "LISTA DE ILUSTRAÇÕES", "\n".join(f"- {x}" for x in front.lista_ilustracoes if str(x).strip()))
    working_org = replace_section_body(working_org, 1, "LISTA DE TABELAS", "\n".join(f"- {x}" for x in front.lista_tabelas if str(x).strip()))
    working_org = replace_section_body(working_org, 1, "LISTA DE ABREVIATURAS E SIGLAS", "\n".join(f"- {x}" for x in front.siglas if str(x).strip()))
    working_org = replace_section_body(working_org, 1, "LISTA DE SÍMBOLOS", "\n".join(f"- {x}" for x in front.simbolos if str(x).strip()))
    if front.glossario:
        working_org = replace_section_body(working_org, 1, "GLOSSÁRIO", "\n".join(f"- {x}" for x in front.glossario if str(x).strip()))

    bib_meta_lookup = build_bib_meta_lookup(bib_entries)
    plan = build_dissertation_stage_plan(template_text, cfg, bib_keys)
    previous_excerpt = ""
    for idx, stage in enumerate(plan, start=1):
        print(f"[5/6] Gerando subseção {idx}/{len(plan)}: {stage['top_title']} :: {stage['sub_title']}...", flush=True)
        stage_prompt = build_dissertation_subsection_prompt(
            cfg=cfg,
            context=context,
            stage=stage,
            bib_meta_lookup=bib_meta_lookup,
            base_docs=base_docs,
            guidance_docs=guidance_docs,
            selected_corpus_catalog=selected_corpus_catalog,
            previous_excerpt=previous_excerpt,
        )
        prompt_parts.append(f"===== stage {idx} {stage['top_title']} / {stage['sub_title']} =====\n" + stage_prompt + "\n")
        stage_resp = client.responses.parse(
            model=model,
            input=[{"role": "user", "content": stage_prompt}],
            text_format=DissertationStageBodyOutput,
        )
        if stage_resp.output_parsed is None:
            raise RuntimeError(f"A IA não retornou a subseção {stage['sub_title']!r}.")
        stage_body = sanitize_generated_org_fragment(stage_resp.output_parsed.org_body)
        stage_body, repair_prompt, diagnostics = maybe_repair_dissertation_subsection(
            client=client,
            model=model,
            cfg=cfg,
            context=context,
            stage=stage,
            body=stage_body,
            bib_meta_lookup=bib_meta_lookup,
            base_docs=base_docs,
        )
        stage_diagnostics.append(diagnostics)
        if repair_prompt:
            prompt_parts.append(f"===== stage {idx} repair {stage['top_title']} / {stage['sub_title']} =====\n" + repair_prompt + "\n")
        working_org = replace_section_body(working_org, 2, stage["sub_title"], stage_body)
        previous_excerpt = stage_body

    cited_after = set(extract_cited_keys_from_org(working_org))
    citation_target_keys = list(cfg.get("__citation_target_keys__") or bib_keys)
    missing_global = [k for k in citation_target_keys if k and k not in cited_after]
    coverage_info = {
        "total_bib_keys": len(bib_keys),
        "citation_target_source": cfg.get("__citation_target_source__", "bib_keys"),
        "citation_target_total": len(citation_target_keys),
        "citation_target_keys": citation_target_keys,
        "cited_after_stages": len(cited_after),
        "missing_after_stages": missing_global,
    }
    if missing_global:
        patch_prompt = build_missing_bibliography_patch_prompt(cfg, context, missing_global, bib_meta_lookup, base_docs, working_org)
        prompt_parts.append("===== missing_bibliography_patch =====\n" + patch_prompt + "\n")
        patch_resp = client.responses.parse(
            model=model,
            input=[{"role": "user", "content": patch_prompt}],
            text_format=DissertationStageBodyOutput,
        )
        if patch_resp.output_parsed is None:
            raise RuntimeError("A IA não retornou a subseção integradora do corpus faltante.")
        patch_body = sanitize_generated_org_fragment(patch_resp.output_parsed.org_body)
        current_results = extract_section_body(working_org, 1, "RESULTADOS E DISCUSSÃO")
        current_results = (current_results.rstrip() + "\n\n" + patch_body.strip()).strip()
        working_org = replace_section_body(working_org, 1, "RESULTADOS E DISCUSSÃO", current_results)
        coverage_info["applied_patch_for_missing_keys"] = True
        coverage_info["missing_before_patch"] = missing_global

    working_org = ensure_document_class(working_org, "dissertacao")
    working_org = apply_citation_style(working_org, bib_filename, style)
    working_org = ensure_cover_command(working_org)
    working_org = cleanup_generated_org(working_org)
    # A bibliografia precisa ser a última normalização: cleanup_generated_org()
    # remove diretivas órfãs, e normalize_bibliography_block() recria uma única
    # diretiva canônica #+PRINT_BIBLIOGRAPHY: ao final do arquivo.
    working_org = normalize_bibliography_block(working_org)
    return working_org, "\n".join(prompt_parts), {"staged_generation": True, "section_diagnostics": stage_diagnostics, "coverage": coverage_info}


def generate_paper_org(client: OpenAI, model: str, cfg: dict[str, Any], context: DocumentContext, template_text: str, base_docs: list[SourceDoc], guidance_docs: list[SourceDoc], bib_filename: str, bib_entries: list[str], bib_keys: list[str], style: str, selected_corpus_catalog: list[dict[str, Any]] | None = None) -> tuple[str, str]:
    doc_type = normalize_document_type(cfg.get("documento", {}).get("tipo_documento"))
    staged_generation = bool(cfg.get("documento", {}).get("geracao_em_etapas", doc_type == "dissertacao"))
    if doc_type == "dissertacao" and staged_generation:
        org_text, prompt_audit, staged_diagnostics = build_dissertation_org_from_stages(
            client=client,
            model=model,
            cfg=cfg,
            context=context,
            template_text=template_text,
            base_docs=base_docs,
            guidance_docs=guidance_docs,
            bib_filename=bib_filename,
            bib_entries=bib_entries,
            bib_keys=bib_keys,
            style=style,
            selected_corpus_catalog=selected_corpus_catalog,
        )
        cfg["__staged_generation_diagnostics__"] = staged_diagnostics
        return org_text, prompt_audit
    prompt = build_paper_prompt(cfg, context, template_text, base_docs, guidance_docs, bib_keys, bib_entries, style, selected_corpus_catalog=selected_corpus_catalog)
    resp = client.responses.create(model=model, input=prompt)
    org_text = resp.output_text.strip()
    org_text = ensure_document_class(org_text, doc_type)
    org_text = apply_citation_style(org_text, bib_filename, style)
    org_text = ensure_cover_command(org_text)
    org_text = cleanup_generated_org(org_text)
    # A bibliografia precisa ser recriada depois da limpeza final do Org.
    org_text = normalize_bibliography_block(org_text)
    return org_text, prompt

def build_expansion_prompt(cfg: dict[str, Any], context: DocumentContext, current_org: str, base_docs: list[SourceDoc], guidance_docs: list[SourceDoc], diagnostics: dict[str, Any], bib_keys: list[str], selected_corpus_catalog: list[dict[str, Any]] | None = None) -> str:
    targets = diagnostics.get("targets", {})
    reasons = diagnostics.get("reasons", [])
    return textwrap.dedent(
        f"""
        Reescreva e expanda o documento acadêmico em Org-mode abaixo, preservando o cabeçalho técnico e a estrutura geral, mas aumentando substancialmente a densidade analítica e a extensão do texto.

        Contexto consolidado:
        {json.dumps(asdict(context), ensure_ascii=False, indent=2)}

        Diagnóstico do rascunho atual:
        {json.dumps(diagnostics, ensure_ascii=False, indent=2)}

        Regras obrigatórias:
        1. Preserve o formato Org-mode, a estrutura do template e o uso de Org Cite.
        2. Não invente chaves; use apenas estas chaves bibliográficas disponíveis: {json.dumps(bib_keys, ensure_ascii=False)}.
        3. Corrija compressões excessivas e desenvolva melhor as seções sinalizadas no diagnóstico.
        4. Para dissertação, o texto precisa ter fôlego analítico, comparando autores, desenvolvendo conceitos, explicitando lacunas e aprofundando implicações para a APF.
        5. Aprofunde especialmente as seções insuficientes segundo as metas mínimas abaixo:
        {json.dumps(targets, ensure_ascii=False, indent=2)}
        6. Razões da expansão necessária: {json.dumps(reasons, ensure_ascii=False)}.
        7. Use de modo mais denso o corpus formalmente selecionado e sua bibliografia.
        8. Não reduza o texto já existente; preserve o que estiver bom e amplie o que estiver curto.

        Catálogo do corpus selecionado:
        {shorten_text(json.dumps(selected_corpus_catalog or [], ensure_ascii=False, indent=2), 30000)}

        Textos-base resumidos:
        {json.dumps(summarize_docs(base_docs, excerpt_chars=3000), ensure_ascii=False, indent=2)}

        Orientações e artefatos da pesquisa:
        {json.dumps(summarize_docs(guidance_docs, excerpt_chars=2600), ensure_ascii=False, indent=2)}

        Documento atual a ser expandido:
        {shorten_text(current_org, 50000)}

        Retorne apenas o .org final completo e expandido.
        """
    ).strip()


def maybe_expand_dissertation_org(client: OpenAI, model: str, cfg: dict[str, Any], context: DocumentContext, org_text: str, base_docs: list[SourceDoc], guidance_docs: list[SourceDoc], bib_filename: str, bib_entries: list[str], bib_keys: list[str], style: str, selected_corpus_catalog: list[dict[str, Any]] | None = None) -> tuple[str, str | None, dict[str, Any] | None]:
    documento = cfg.get("documento", {})
    doc_type_local = normalize_document_type(documento.get("tipo_documento"))
    if doc_type_local != "dissertacao":
        return org_text, None, {}
    if bool(documento.get("geracao_em_etapas", True)):
        diagnostics = cfg.get("__staged_generation_diagnostics__") or {"staged_generation": True, "skipped_whole_document_expansion": True}
        return org_text, None, diagnostics
    selected_keys = sorted({d.bib_key for d in base_docs if d.kind.startswith("texto_selecionado") and d.bib_key})
    should_expand, diagnostics = should_expand_dissertation(cfg, org_text, selected_keys)
    if not should_expand:
        return org_text, None, diagnostics
    prompt = build_expansion_prompt(cfg, context, org_text, base_docs, guidance_docs, diagnostics, bib_keys, selected_corpus_catalog=selected_corpus_catalog)
    resp = client.responses.create(model=model, input=prompt)
    expanded = resp.output_text.strip()
    expanded = ensure_document_class(expanded, doc_type_local)
    expanded = apply_citation_style(expanded, bib_filename, style)
    expanded = ensure_cover_command(expanded)
    expanded = cleanup_generated_org(expanded)
    # A bibliografia precisa ser recriada depois da limpeza final do Org.
    expanded = normalize_bibliography_block(expanded)
    return expanded, prompt, diagnostics

def build_document_output_dir(cfg: dict[str, Any], research_root: Path) -> Path:
    documento = cfg.get("documento", {})
    doc_type = normalize_document_type(documento.get("tipo_documento"))
    default_suffix = DEFAULT_DOC_PREFIX_BY_TYPE.get(doc_type, "documento")
    prefix = (documento.get("prefixo") or ((cfg.get("saida", {}).get("prefixo") or "atividade") + f"_{default_suffix}")).strip()
    output_dir_raw = documento.get("output_dir")
    if output_dir_raw:
        base = resolve_configured_path(output_dir_raw, cfg)
        if base is None:
            base = research_root
    else:
        base = research_root
    create_subdir = bool(documento.get("criar_subdiretorio", True))
    if not create_subdir:
        return base
    if base.name == prefix:
        return base
    return base / prefix

def build_bundle_dir(cfg: dict[str, Any], research_root: Path) -> Path:
    pipeline = cfg.get("pipeline", {})
    bundle_dir_raw = pipeline.get("bundle_dir")
    if bundle_dir_raw:
        resolved = resolve_configured_path(bundle_dir_raw, cfg)
        if resolved is not None:
            return resolved
    return research_root / "documento_bundle"

def copy_if_exists(src: Path | None, dest: Path) -> str | None:
    if src is None or not src.exists():
        return None
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)
    return str(dest)

def build_bundle(cfg: dict[str, Any], research_paths: ResearchPaths, documento_context: DocumentContext, debug_json: dict[str, Any], bundle_dir: Path) -> Path:
    bundle_dir.mkdir(parents=True, exist_ok=True)
    refs_dir = bundle_dir / "referencias"
    base_dir = bundle_dir / "textos_base"
    orient_dir = bundle_dir / "orientacoes"

    manifest: dict[str, Any] = {
        "created_at": datetime.now().isoformat(),
        "mode_origem": cfg.get("atividade", {}).get("modo"),
        "research_root": str(research_paths.root_dir),
        "documento_context": asdict(documento_context),
        "artifacts": {},
    }

    artifacts = manifest["artifacts"]
    artifacts["research_org"] = copy_if_exists(research_paths.org_path, orient_dir / (research_paths.org_path.name if research_paths.org_path else "pesquisa.org"))
    artifacts["research_debug_json"] = copy_if_exists(research_paths.debug_path, orient_dir / (research_paths.debug_path.name if research_paths.debug_path else "pesquisa_debug.json"))
    artifacts["research_pdf"] = copy_if_exists(research_paths.pdf_path, orient_dir / (research_paths.pdf_path.name if research_paths.pdf_path else "pesquisa.pdf"))
    artifacts["research_prisma_svg"] = copy_if_exists(research_paths.prisma_svg_path, orient_dir / (research_paths.prisma_svg_path.name if research_paths.prisma_svg_path else "prisma.svg"))
    artifacts["research_prisma_pdf"] = copy_if_exists(research_paths.prisma_pdf_path, orient_dir / (research_paths.prisma_pdf_path.name if research_paths.prisma_pdf_path else "prisma.pdf"))
    artifacts["research_bib"] = copy_if_exists(research_paths.bib_path, refs_dir / (research_paths.bib_path.name if research_paths.bib_path else "references.bib"))
    artifacts["research_config"] = copy_if_exists(research_paths.config_path, orient_dir / (research_paths.config_path.name if research_paths.config_path else "config.toml"))

    copied_fulltexts: list[str] = []
    for pdf in (research_paths.selected_fulltext_paths or research_paths.fulltext_paths):
        target = base_dir / pdf.name
        copy_if_exists(pdf, target)
        copied_fulltexts.append(str(target))
    artifacts["fulltexts"] = copied_fulltexts
    artifacts["selected_entries_json"] = None
    selected_entries_path = bundle_dir / "selected_entries.json"
    if research_paths.selected_entries:
        write_text(selected_entries_path, json.dumps(research_paths.selected_entries, ensure_ascii=False, indent=2))
        artifacts["selected_entries_json"] = str(selected_entries_path)

    context_path = bundle_dir / "contexto_paper.json"
    write_text(context_path, json.dumps({
        "documento_context": asdict(documento_context),
        "debug_json_excerpt": debug_json,
        "config_excerpt": {
            "atividade": cfg.get("atividade", {}),
            "pesquisa": cfg.get("pesquisa", {}),
            "bibliografia": cfg.get("bibliografia", {}),
        },
    }, ensure_ascii=False, indent=2, default=str))
    manifest["context_json"] = str(context_path)

    manifest_path = bundle_dir / "manifest.json"
    write_text(manifest_path, json.dumps(manifest, ensure_ascii=False, indent=2))
    return manifest_path

def update_bundle_with_documento(
    manifest_path: Path | None,
    *,
    org_path: Path,
    bib_path: Path,
    context_json_path: Path,
    prompt_audit_path: Path,
    provenance_path: Path | None = None,
    extra_docs: list[SourceDoc] | None = None,
    pdf_path: Path | None = None,
) -> None:
    if manifest_path is None or not manifest_path.exists():
        return

    bundle_dir = manifest_path.parent
    documento_dir = bundle_dir / "paper"
    documento_dir.mkdir(parents=True, exist_ok=True)

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    artifacts = manifest.setdefault("artifacts", {})

    def copy_local(src: Path, key: str) -> str:
        dest = documento_dir / src.name
        shutil.copy2(src, dest)
        artifacts[key] = str(dest)
        return str(dest)

    copy_local(org_path, "documento_org")
    copy_local(bib_path, "documento_bib")
    copy_local(context_json_path, "documento_context_json")
    copy_local(prompt_audit_path, "documento_prompt_audit")
    if provenance_path and provenance_path.exists():
        copy_local(provenance_path, "documento_provenance_json")
    if pdf_path and pdf_path.exists():
        copy_local(pdf_path, "documento_pdf")

    copied_extras: list[str] = []
    for doc in extra_docs or []:
        p = Path(doc.path)
        if p.exists() and p.is_file():
            dest = documento_dir / "extras" / p.name
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(p, dest)
            copied_extras.append(str(dest))
    artifacts["documento_extra_articles"] = copied_extras

    manifest["updated_at"] = datetime.now().isoformat()
    write_text(manifest_path, json.dumps(manifest, ensure_ascii=False, indent=2))


def assemble_delivery_package(
    cfg: dict[str, Any],
    *,
    research_paths: ResearchPaths,
    documento_output_dir: Path,
    documento_prefix: str,
    manifest_path: Path | None,
    org_path: Path,
    bib_path: Path,
    context_json_path: Path,
    prompt_audit_path: Path,
    provenance_path: Path,
    reference_usage_json_path: Path,
    reference_usage_md_path: Path,
    section_limits_json_path: Path,
    documento_pdf_path: Path | None = None,
) -> Path | None:
    entrega = cfg.get("entrega", {})
    if not bool(entrega.get("gerar_pacote_final", True)):
        return None

    package_dir = documento_output_dir / "entrega_final"
    package_dir.mkdir(parents=True, exist_ok=True)

    def copy_optional(src: Path | None, subname: str | None = None) -> str | None:
        if src is None or not src.exists():
            return None
        dest = package_dir / (subname or src.name)
        shutil.copy2(src, dest)
        return str(dest)

    copied = {
        "pesquisa_org": copy_optional(research_paths.org_path),
        "pesquisa_bib": copy_optional(research_paths.bib_path),
        "pesquisa_debug": copy_optional(research_paths.debug_path),
        "pesquisa_pdf": copy_optional(research_paths.pdf_path),
        "documento_org": copy_optional(org_path),
        "documento_bib": copy_optional(bib_path),
        "documento_contexto": copy_optional(context_json_path),
        "documento_prompts": copy_optional(prompt_audit_path),
        "documento_proveniencia": copy_optional(provenance_path),
        "reference_usage_json": copy_optional(reference_usage_json_path),
        "reference_usage_md": copy_optional(reference_usage_md_path),
        "section_limits_json": copy_optional(section_limits_json_path),
        "bundle_manifest": copy_optional(manifest_path),
    }
    if documento_pdf_path and bool(entrega.get("incluir_documento_pdf", True)):
        copied["documento_pdf"] = copy_optional(documento_pdf_path)
    if research_paths.prisma_svg_path and bool(entrega.get("incluir_prisma_svg", True)):
        copied["pesquisa_prisma_svg"] = copy_optional(research_paths.prisma_svg_path)
    if research_paths.prisma_pdf_path and bool(entrega.get("incluir_prisma_pdf", True)):
        copied["pesquisa_prisma_pdf"] = copy_optional(research_paths.prisma_pdf_path)

    readme = package_dir / "README_entrega.md"
    lines = [
        f"# Pacote final de entrega — {documento_prefix}",
        "",
        "Este diretório reúne os principais artefatos finais do pipeline integrado.",
        "",
        "## Itens copiados",
    ]
    lines.extend(f"- **{k}**: `{v}`" for k, v in copied.items() if v)
    lines.append("")
    write_text(readme, "\n".join(lines))
    return package_dir



def extract_citation_style_from_org(org_text: str) -> str:
    m = re.search(r"(?im)^\s*#\+CITE_EXPORT:\s*biblatex(?:\s+([A-Za-z0-9_\-]+))?", org_text)
    if m and m.group(1):
        return m.group(1).strip().lower()
    return DEFAULT_STYLE


def escape_bfseries_line_text(line: str) -> str:
    """Escapa caracteres especiais em linhas LaTeX cruas de título visual.

    Corrige casos como:
      {\Large\bfseries Corpus local — atividade_x\par}
    sem alterar comandos LaTeX estruturais no restante do documento.
    """
    if r"\bfseries" not in line or r"\par" not in line:
        return line
    pattern = r"(\{\\(?:Large|large|LARGE|huge|Huge|normalsize)?\\bfseries\s+)(.*?)(\\par\})"
    def repl(m: re.Match) -> str:
        return m.group(1) + latex_escape_text(m.group(2)) + m.group(3)
    return re.sub(pattern, repl, line)


def postprocess_generated_tex_for_bibliography(tex_path: Path, bib_filename: str, style: str) -> None:
    tex = tex_path.read_text(encoding="utf-8", errors="ignore")
    bib_name = Path(str(bib_filename)).name
    style = (style or DEFAULT_STYLE).strip().lower() or DEFAULT_STYLE

    tex = re.sub(r"(?im)^\\PassOptionsToPackage\{.*\}\{biblatex\}\s*$\n?", "", tex)
    tex = re.sub(r"(?im)^\\ExecuteBibliographyOptions\{.*\}\s*$\n?", "", tex)
    tex = re.sub(r"(?ms)^sortcites=true,.*?(?=\\begin\{document\})", "", tex)
    tex = re.sub(r"(?ms)^pdfauthor=.*?(?=\\usepackage(?:\[[^\]]*\])?\{biblatex\}|\\begin\{document\})", "", tex)
    tex = re.sub(r"(?ms)^\\hypersetup\{.*?\}\s*", "", tex)
    tex = tex.replace("[][]", "")
    tex = tex.replace("<empty citation>", "")
    tex = re.sub(r"(?m)^\s*[,;:]+\s*$\n?", "", tex)
    tex = re.sub(r"(?m)^\s*P\d+\\?\}\s*$\n?", "", tex)
    tex = re.sub(r"\\dataaprovacao\{[^}]*__[^}]*\}", r"\\dataaprovacao{A definir}", tex)
    tex = re.sub(r"\\dataaprovacao\{\s*\}", r"\\dataaprovacao{A definir}", tex)
    tex = re.sub(r"\\autor\{\s*Seu nome\s*\}", rf"\\autor{{{DEFAULT_AUTHOR}}}", tex)
    tex = re.sub(r"\\author\{\s*Seu nome\s*\}", rf"\\author{{{DEFAULT_AUTHOR}}}", tex)

    if re.search(r"\\usepackage(?:\[[^\]]*\])?\{biblatex\}", tex):
        tex = re.sub(
            r"\\usepackage(?:\[[^\]]*\])?\{biblatex\}",
            rf"\\usepackage[backend=biber,style={style}]{{biblatex}}",
            tex,
            count=1,
        )
    else:
        tex = tex.replace(r"\begin{document}", f"\\usepackage[backend=biber,style={style}]{{biblatex}}\n\\addbibresource{{{bib_name}}}\n\\begin{{document}}", 1)

    tex = re.sub(r"\\addbibresource\{[^}]+\}", rf"\\addbibresource{{{bib_name}}}", tex)
    if not re.search(r"\\addbibresource\{[^}]+\}", tex):
        tex = tex.replace(
            rf"\usepackage[backend=biber,style={style}]{{biblatex}}",
            f"\\usepackage[backend=biber,style={style}]{{biblatex}}\n\\addbibresource{{{bib_name}}}",
            1,
        )

    tex = tex.replace(r"\maketitle", "")

    # Ãltima barreira contra placeholders de template que escapem do .org.
    for raw_key in template_placeholder_values({}, "").keys():
        tex = tex.replace("{{" + raw_key + "}}", "")
        tex = tex.replace("{" + raw_key + "}", "{}")
        tex = tex.replace(raw_key, raw_key.replace("_", " ").title())
    tex = re.sub(r"\{\{[^{}]{1,80}\}\}", "", tex)

    cleaned_lines = []
    seen_addbib = False
    for line in tex.splitlines():
        stripped = line.strip()
        if stripped.startswith("sortcites=true") or stripped.startswith("pdfauthor=") or stripped.startswith("pdftitle=") or stripped.startswith("pdfkeywords=") or stripped.startswith("pdfsubject=") or stripped.startswith("pdfcreator=") or stripped.startswith("pdflang="):
            continue
        if stripped in {",", ";", ":", "P26}", "P26\\}"}:
            continue
        if re.fullmatch(r"P\d+\\?\}", stripped):
            continue
        if stripped.startswith(r"\addbibresource{"):
            if seen_addbib:
                continue
            line = rf"\addbibresource{{{bib_name}}}"
            seen_addbib = True
        line = escape_bfseries_line_text(line)
        cleaned_lines.append(line)

    tex = "\n".join(cleaned_lines).strip() + "\n"
    tex_path.write_text(tex, encoding="utf-8")

def build_latex_env(latex_extra_path: Path | None = None) -> dict[str, str]:
    env = os.environ.copy()
    if latex_extra_path:
        resolved = latex_extra_path.expanduser().resolve()
        latex_dir = resolved.parent if resolved.is_file() else resolved
        texinputs_prefix = f"{latex_dir.as_posix()}//:"
        env["TEXINPUTS"] = texinputs_prefix + env.get("TEXINPUTS", "")
        env["BIBINPUTS"] = texinputs_prefix + env.get("BIBINPUTS", "")
        env["BSTINPUTS"] = texinputs_prefix + env.get("BSTINPUTS", "")
    return env

def clean_latex_intermediate_files(tex_path: Path) -> None:
    stem = tex_path.with_suffix("")
    suffixes = [".aux", ".bbl", ".bcf", ".blg", ".log", ".out", ".run.xml", ".toc", ".lof", ".lot", ".fls", ".fdb_latexmk"]
    for suffix in suffixes:
        path = stem.with_suffix(suffix)
        try:
            if path.exists():
                path.unlink()
        except Exception:
            pass


def run_compile_sequence(org_path: Path, *, emacs_init: Path | None = None, academic_writing: Path | None = None, latex_extra_path: Path | None = None) -> Path:
    emacs = ensure_command("emacs")
    lualatex = ensure_command("lualatex")
    org_text = org_path.read_text(encoding="utf-8", errors="ignore")
    uses_citations = org_uses_citation_pipeline(org_text)
    style = extract_citation_style_from_org(org_text)
    export_el = org_path.parent / f"{org_path.stem}_export_pdf.el"
    tex_path = org_path.with_suffix('.tex')
    bib_path = org_path.with_suffix('.bib')

    init_loader = ""
    if academic_writing is not None:
        init_loader = f'(load-file "{academic_writing.as_posix()}")\n'

    cite_requirements = ""
    cite_guard = ""
    export_call = "(org-latex-export-to-latex)"

    if uses_citations:
        ensure_command("biber")
        cite_requirements = "(ignore-errors (require 'oc))\n(ignore-errors (require 'oc-biblatex))\n"
        cite_guard = textwrap.dedent("""
        (setq org-cite-insert-processor 'basic
              org-cite-follow-processor 'basic
              org-cite-activate-processor 'basic
              org-cite-export-processors '((latex biblatex) (t basic)))
        (setq-local org-cite-export-processors '((latex biblatex) (t basic)))
        """).strip() + "\n"
    else:
        cite_guard = textwrap.dedent("""
        (setq org-cite-insert-processor 'basic
              org-cite-follow-processor 'basic
              org-cite-activate-processor 'basic
              org-cite-export-processors '((latex basic) (t basic)))
        (setq-local org-cite-export-processors '((latex basic) (t basic)))
        (setq-local org-cite-global-bibliography nil)
        (when (boundp 'org-export-with-cite-processors)
          (setq org-export-with-cite-processors nil)
          (setq-local org-export-with-cite-processors nil))
        (when (boundp 'org-export-process-citations)
          (setq org-export-process-citations nil)
          (setq-local org-export-process-citations nil))
        """).strip() + "\n"
        export_call = "(let ((org-cite-export-processors '((latex basic) (t basic))) (org-export-with-cite-processors nil) (org-export-process-citations nil)) (org-latex-export-to-latex nil nil nil nil '(:with-cite-processors nil)))"

    export_code = textwrap.dedent(f"""
    {init_loader}(require 'org)
    (require 'ox)
    (require 'ox-latex)
    {cite_requirements}(find-file "{org_path.as_posix()}")
    (setq-local org-export-use-babel nil)
    (setq-local org-confirm-babel-evaluate nil)
    {cite_guard}{export_call}
    """).strip() + "\n"
    write_text(export_el, export_code)

    cmd = [emacs, "--batch", "-Q"]
    if emacs_init is not None:
        cmd.extend(["-l", str(emacs_init)])
    cmd.extend(["-l", str(export_el)])
    env = build_latex_env(latex_extra_path)
    proc = subprocess.run(cmd, cwd=str(org_path.parent), capture_output=True, text=True, env=env)
    debug_print("Exportando TEX:", cmd)
    if proc.returncode != 0:
        raise RuntimeError(f"Falha ao exportar TEX via Emacs batch.\nSTDOUT:\n{proc.stdout}\nSTDERR:\n{proc.stderr}")
    if not tex_path.exists():
        raise RuntimeError(f"Exportação concluída sem erro, mas o TEX não foi encontrado: {tex_path}")

    postprocess_generated_tex_for_bibliography(tex_path, bib_path.name, style)
    clean_latex_intermediate_files(tex_path)

    commands = [[lualatex, "-interaction=nonstopmode", "-file-line-error", tex_path.name]]
    if uses_citations:
        biber = ensure_command("biber")
        commands.append([biber, tex_path.stem])
        commands.append([lualatex, "-interaction=nonstopmode", "-file-line-error", tex_path.name])
        commands.append([lualatex, "-interaction=nonstopmode", "-file-line-error", tex_path.name])
    else:
        commands.append([lualatex, "-interaction=nonstopmode", "-file-line-error", tex_path.name])

    combined_stdout = []
    combined_stderr = []
    for one in commands:
        proc = subprocess.run(one, cwd=str(org_path.parent), capture_output=True, text=True, env=env)
        debug_print("Compilando PDF:", one)
        combined_stdout.append(proc.stdout)
        combined_stderr.append(proc.stderr)
        if proc.returncode != 0:
            raise RuntimeError(
                "Falha ao compilar PDF do TEX pós-processado.\n"
                f"Comando: {' '.join(one)}\nSTDOUT:\n{proc.stdout}\nSTDERR:\n{proc.stderr}"
            )

    pdf_path = org_path.with_suffix(".pdf")
    if not pdf_path.exists():
        raise RuntimeError(
            "Compilação concluída sem erro, mas o PDF não foi encontrado.\n"
            f"STDOUT final:\n{''.join(combined_stdout)[-4000:]}\nSTDERR final:\n{''.join(combined_stderr)[-4000:]}"
        )
    return pdf_path

def build_provenance_payload(
    cfg: dict[str, Any],
    research_paths: ResearchPaths,
    documento_context: DocumentContext,
    base_docs: list[SourceDoc],
    guidance_docs: list[SourceDoc],
    extra_docs: list[SourceDoc],
    bib_keys: list[str],
    bib_entries: list[str],
    context_origin_info: dict[str, Any],
    template_path: Path | None,
    manifest_path: Path | None = None,
) -> dict[str, Any]:
    documento = cfg.get("documento", {})
    return {
        "generated_at": datetime.now().isoformat(),
        "pipeline": {
            "executar_pesquisa": bool(cfg.get("pipeline", {}).get("executar_pesquisa", True)),
            "executar_documento": bool(cfg.get("pipeline", {}).get("executar_documento", True)),
            "pesquisa_dir_existente": str(cfg.get("pipeline", {}).get("pesquisa_dir_existente") or ""),
        },
        "documento_controls": {
            "modo_escrita": str(documento.get("modo_escrita") or "novo"),
            "perfil_redacao": str(documento.get("perfil_redacao") or "academico_equilibrado"),
            "usar_bib_da_pesquisa": bool(documento.get("usar_bib_da_pesquisa", True)),
            "incluir_artigos_extras_no_bib": bool(documento.get("incluir_artigos_extras_no_bib", True)),
            "permitir_busca_correlata_extra": bool(documento.get("permitir_busca_correlata_extra", False)),
            "priorizar_citacoes_dos_selecionados": bool(documento.get("priorizar_citacoes_dos_selecionados", True)),
            "extras_so_complementam": bool(documento.get("extras_so_complementam", True)),
            "preservar_estrutura_do_org_anterior": bool(documento.get("preservar_estrutura_do_org_anterior", False)),
            "usar_contexto_consolidado_da_pesquisa": bool(documento.get("usar_contexto_consolidado_da_pesquisa", True)),
            "reformular_tema_recorte_objetivo": bool(documento.get("reformular_tema_recorte_objetivo", False)),
            "limites_palavras": {
                "total": documento.get("limite_palavras_total"),
                "introducao": documento.get("limite_palavras_introducao"),
                "revisao": documento.get("limite_palavras_revisao"),
                "conclusao": documento.get("limite_palavras_conclusao"),
            },
        },
        "context_origin": context_origin_info,
        "documento_context_final": asdict(documento_context),
        "template_path": str(template_path) if template_path else None,
        "bundle_manifest": str(manifest_path) if manifest_path else None,
        "research_artifacts": {
            "root_dir": str(research_paths.root_dir),
            "org_path": str(research_paths.org_path) if research_paths.org_path else None,
            "bib_path": str(research_paths.bib_path) if research_paths.bib_path else None,
            "debug_path": str(research_paths.debug_path) if research_paths.debug_path else None,
            "pdf_path": str(research_paths.pdf_path) if research_paths.pdf_path else None,
            "fulltext_cache_dir": str(research_paths.fulltext_cache_dir) if research_paths.fulltext_cache_dir else None,
        },
        "selected_entries_count": len(research_paths.selected_entries),
        "selected_fulltext_paths_used": [d.path for d in base_docs if d.kind.startswith("texto_selecionado")],
        "extra_article_paths_used": [d.path for d in extra_docs],
        "guidance_sources_used": [d.path for d in guidance_docs],
        "bib_keys_count": len(bib_keys),
        "bib_entries_count": len(bib_entries),
    }

def build_dry_run_report(cfg: dict[str, Any], research_cfg: dict[str, Any], research_script: Path, research_paths: ResearchPaths) -> dict[str, Any]:
    documento = cfg.get("documento", {})
    return {
        "generated_at": datetime.now().isoformat(),
        "dry_run": True,
        "pipeline": {**cfg.get("pipeline", {}), "executar_bundle": bool(cfg.get("pipeline", {}).get("executar_bundle", cfg.get("pipeline", {}).get("criar_bundle", True)))},
        "research_script_exists": research_script.exists(),
        "research_output_dir": str(research_paths.root_dir),
        "research_artifacts_exist": {
            "org": bool(research_paths.org_path and research_paths.org_path.exists()),
            "bib": bool(research_paths.bib_path and research_paths.bib_path.exists()),
            "debug": bool(research_paths.debug_path and research_paths.debug_path.exists()),
            "pdf": bool(research_paths.pdf_path and research_paths.pdf_path.exists()),
            "fulltext_cache_dir": bool(research_paths.fulltext_cache_dir and research_paths.fulltext_cache_dir.exists()),
        },
        "counts": {
            "selected_entries": len(research_paths.selected_entries),
            "selected_fulltexts": len(research_paths.selected_fulltext_paths),
            "fulltexts_total": len(research_paths.fulltext_paths),
        },
        "paper_plan": {
            "modo_escrita": str(documento.get("modo_escrita") or "novo"),
            "perfil_redacao": str(documento.get("perfil_redacao") or "academico_equilibrado"),
            "usar_artigos_selecionados_pesquisa": bool(documento.get("usar_artigos_selecionados_pesquisa", True)),
            "usar_bib_da_pesquisa": bool(documento.get("usar_bib_da_pesquisa", True)),
            "incluir_artigos_extras_no_bib": bool(documento.get("incluir_artigos_extras_no_bib", True)),
            "permitir_busca_correlata_extra": bool(documento.get("permitir_busca_correlata_extra", False)),
            "priorizar_citacoes_dos_selecionados": bool(documento.get("priorizar_citacoes_dos_selecionados", True)),
            "extras_so_complementam": bool(documento.get("extras_so_complementam", True)),
            "preservar_estrutura_do_org_anterior": bool(documento.get("preservar_estrutura_do_org_anterior", False)),
            "usar_contexto_consolidado_da_pesquisa": bool(documento.get("usar_contexto_consolidado_da_pesquisa", True)),
            "reformular_tema_recorte_objetivo": bool(documento.get("reformular_tema_recorte_objetivo", False)),
            "limites_palavras": {
                "total": documento.get("limite_palavras_total"),
                "introducao": documento.get("limite_palavras_introducao"),
                "revisao": documento.get("limite_palavras_revisao"),
                "conclusao": documento.get("limite_palavras_conclusao"),
            },
            "artigos_extras_paths": cfg.get("documento", {}).get("artigos_extras_paths", []),
        },
        "research_config_excerpt": {
            "atividade": research_cfg.get("atividade", {}),
            "pesquisa": research_cfg.get("pesquisa", {}),
            "saida": research_cfg.get("saida", {}),
        },
    }

# ---------------------------------------------------------------------------
# Config / CLI
# ---------------------------------------------------------------------------
def load_config(path: Path) -> dict[str, Any]:
    with open(path, "rb") as fh:
        return tomllib.load(fh)

def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Pipeline integrado de pesquisa + documento usando TOML unificado.")
    p.add_argument("--config", required=True, help="Arquivo TOML unificado do pipeline.")
    p.add_argument("--model", default=None, help="Override do modelo OpenAI para a etapa do documento.")
    p.add_argument("--debug", action="store_true")
    return p.parse_args()

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

# ---------------------------------------------------------------------------
# Mock run: gera toda a cadeia de artefatos com conteúdo sintético
# ---------------------------------------------------------------------------
def _mock_lipsum(topic: str, idx: int = 1) -> str:
    base = [
        f"Texto simulado para validação estrutural do pipeline sobre {topic}.",
        "Este conteúdo foi gerado em mock_run para testar a escrita de arquivos, a montagem do pacote e a compilação opcional do PDF.",
        "O objetivo é preservar uma semântica acadêmica mínima, com encadeamento coerente entre problema, método, resultados e implicações.",
        "As passagens incluem citações fictícias, tabelas de exemplo e referências simuladas para testar bibliografia, sumário e renderização."
    ]
    return " ".join(base + [f"Parágrafo sintético {idx}."])

def _mock_write_pdf(path: Path, title: str) -> None:
    # PDF mínimo e válido para testes de existência/empacotamento.
    payload = f"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>
endobj
4 0 obj
<< /Length 95 >>
stream
BT
/F1 18 Tf
72 780 Td
({title[:60].replace('(', '[').replace(')', ']')}) Tj
0 -28 Td
/F1 11 Tf
(Mock PDF gerado em mock_run para validar a cadeia de artefatos.) Tj
ET
endstream
endobj
5 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
xref
0 6
0000000000 65535 f\x20
0000000010 00000 n\x20
0000000063 00000 n\x20
0000000122 00000 n\x20
0000000251 00000 n\x20
0000000397 00000 n\x20
trailer
<< /Size 6 /Root 1 0 R >>
startxref
467
%%EOF
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload.encode("latin-1", errors="ignore"))

def _mock_make_bib_entries(n: int, topic: str) -> tuple[list[str], list[str], list[str]]:
    keys: list[str] = []
    entries: list[str] = []
    titles: list[str] = []
    for i in range(1, n + 1):
        key = f"mock{i:02d}"
        title = f"Mock study {i:02d} on {topic}"
        year = 2018 + (i % 8)
        keys.append(key)
        titles.append(title)
        entries.append(textwrap.dedent(f"""
        @article{{{key},
          author = {{Autor Mock {i} and Coautor Simulado {i}}},
          title = {{{title}}},
          year = {{{year}}},
          journaltitle = {{Journal of Synthetic Public Sector AI Studies}},
          volume = {{{1 + (i % 9)}}},
          number = {{{1 + (i % 4)}}},
          pages = {{{10+i}-{18+i}}},
          doi = {{10.9999/mock.{i:02d}}},
          url = {{https://example.org/mock/{i:02d}}}
        }}
        """).strip())
    return entries, keys, titles

def _mock_make_prisma_svg(path: Path, total: int, selected: int) -> None:
    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" width="900" height="520" viewBox="0 0 900 520">
  <style>
    .b {{ fill:#f8fafc; stroke:#334155; stroke-width:2; rx:16; ry:16; }}
    .t {{ font: 18px Arial, sans-serif; fill:#0f172a; }}
    .s {{ font: 15px Arial, sans-serif; fill:#334155; }}
    .a {{ stroke:#475569; stroke-width:3; marker-end:url(#m); fill:none; }}
  </style>
  <defs><marker id="m" markerWidth="10" markerHeight="7" refX="8" refY="3.5" orient="auto"><polygon points="0 0, 10 3.5, 0 7" fill="#475569"/></marker></defs>
  <rect x="320" y="35" width="260" height="80" class="b"/><text x="350" y="70" class="t">Registros identificados</text><text x="350" y="98" class="s">n = {total}</text>
  <rect x="320" y="165" width="260" height="80" class="b"/><text x="350" y="200" class="t">Após deduplicação</text><text x="350" y="228" class="s">n = {max(total-8, selected)}</text>
  <rect x="320" y="295" width="260" height="80" class="b"/><text x="350" y="330" class="t">Triagem de títulos/resumos</text><text x="350" y="358" class="s">n = {max(total-15, selected)}</text>
  <rect x="320" y="425" width="260" height="80" class="b"/><text x="350" y="460" class="t">Estudos selecionados</text><text x="350" y="488" class="s">n = {selected}</text>
  <line x1="450" y1="115" x2="450" y2="165" class="a"/><line x1="450" y1="245" x2="450" y2="295" class="a"/><line x1="450" y1="375" x2="450" y2="425" class="a"/>
</svg>'''
    write_text(path, svg)

def _mock_make_selected_entries(root_dir: Path, prefixo: str, titles: list[str], keys: list[str]) -> tuple[list[dict[str, Any]], list[Path]]:
    cache_dir = root_dir / f"{prefixo}_fulltext_cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    entries: list[dict[str, Any]] = []
    pdfs: list[Path] = []
    for i, (title, key) in enumerate(zip(titles, keys), start=1):
        pdf_path = cache_dir / f"{key}.pdf"
        _mock_write_pdf(pdf_path, title)
        pdfs.append(pdf_path)
        entries.append({
            "title": title,
            "abstract": _mock_lipsum(title, i),
            "year": str(2018 + (i % 8)),
            "authors": [f"Autor Mock {i}", f"Coautor Simulado {i}"],
            "source": "mock_run",
            "doi": f"10.9999/mock.{i:02d}",
            "url": f"https://example.org/mock/{i:02d}",
            "downloaded_pdf_path": str(pdf_path),
            "bib_key": key,
            "reason_selected": "Seleção sintética para validação completa da cadeia."
        })
    return entries, pdfs

def _mock_build_research_org(cfg: dict[str, Any], prefixo: str, titles: list[str], keys: list[str]) -> str:
    pesquisa = cfg.get("pesquisa", {})
    topic = pesquisa.get("tema") or "tema de pesquisa"
    rows = ["| ID | Título | Ano | Eixo | Chave |", "|-"]
    for i, (title, key) in enumerate(zip(titles, keys), start=1):
        eixo = ["governança", "ética", "compliance", "riscos"][i % 4]
        rows.append(f"| {i} | {title} | {2018 + (i % 8)} | {eixo} | {key} |")
    cites = " ".join(f"[cite:@{k}]" for k in keys[: min(6, len(keys))])
    rows_text = "\n".join(rows)
    selected_blocks = []
    for i, title in enumerate(titles[: min(10, len(titles))], start=1):
        selected_blocks.append(
            f"** Texto selecionado {i}\n{title}\n{_mock_lipsum(title, i+2)}\n"
        )
    selected_text = "\n".join(selected_blocks)
    return textwrap.dedent(f"""
    #+TITLE: {prefixo}
    #+AUTHOR: {cfg.get("atividade", {}).get("aluno") or DEFAULT_AUTHOR}
    #+LANGUAGE: pt_BR
    #+OPTIONS: toc:t num:t
    #+CITE_EXPORT: biblatex apa
    #+BIBLIOGRAPHY: {prefixo}.bib

    * Introdução
    {_mock_lipsum(topic, 1)}

    * Estratégia de busca
    {_mock_lipsum(topic, 2)}

    * Panorama preliminar
    {cites}

    * Tabela sintética de estudos selecionados
    {rows_text}

    * Estudos selecionados
    {selected_text}

    * Referências
    #+PRINT_BIBLIOGRAPHY:
    """).strip() + "\n"

def _mock_build_document_org(cfg: dict[str, Any], documento_prefix: str, doc_type: str, titles: list[str], keys: list[str]) -> str:
    pesquisa = cfg.get("pesquisa", {})
    tema = pesquisa.get("tema") or "tema"
    atividade = cfg.get("atividade", {})
    class_name = "fgv-dissertacao" if doc_type == "dissertacao" else "fgv-paper"
    cite_block = " ".join(f"[cite:@{k}]" for k in keys[: min(10, len(keys))])
    table = "\n".join([
        "| Eixo | Descrição | Intensidade | Referência |",
        "|-",
        f"| Governança | Estruturas decisórias e responsabilização | Alta | {keys[0] if keys else 'mock01'} |",
        f"| Ética | Transparência, vieses e supervisão humana | Alta | {keys[1] if len(keys)>1 else 'mock02'} |",
        f"| Compliance | Conformidade regulatória e controles | Média | {keys[2] if len(keys)>2 else 'mock03'} |",
        f"| Gestão de riscos | Monitoramento e mitigação | Alta | {keys[3] if len(keys)>3 else 'mock04'} |",
    ])
    corpo = textwrap.dedent(f"""
    #+TITLE: {documento_prefix}
    #+AUTHOR: {atividade.get("aluno") or DEFAULT_AUTHOR}
    #+LANGUAGE: pt_BR
    #+OPTIONS: toc:t num:t
    #+LATEX_COMPILER: lualatex
    #+LATEX_CLASS: {class_name}
    #+CITE_EXPORT: biblatex apa
    #+BIBLIOGRAPHY: {documento_prefix}.bib

    * RESUMO
    {_mock_lipsum(tema, 1)}

    * ABSTRACT
    {_mock_lipsum(tema, 2)}

    * INTRODUÇÃO
    {_mock_lipsum(tema, 3)}

    * REFERENCIAL TEÓRICO
    {cite_block}

    * METODOLOGIA
    {_mock_lipsum(tema, 4)}

    * RESULTADOS E DISCUSSÃO
    {_mock_lipsum(tema, 5)}

    ** Tabela sintética de eixos analíticos
    {table}

    ** Discussão integrada
    {' '.join(f"[cite:@{k}]" for k in keys[: min(8, len(keys))])}

    * CONSIDERAÇÕES FINAIS
    {_mock_lipsum(tema, 6)}

    * REFERÊNCIAS
    #+PRINT_BIBLIOGRAPHY:
    """).strip() + "\n"
    return corpo

def run_mock_pipeline(
    cfg: dict[str, Any],
    *,
    research_script: Path,
    research_output_dir: Path,
    temp_cfg_path: Path,
    executar_pesquisa: bool,
    executar_documento: bool,
    executar_bundle: bool,
) -> int:
    controle = cfg.get("controle", {})
    documento = cfg.get("documento", {})
    mock_seed = int(controle.get("mock_seed", 42) or 42)
    mock_registros = int(controle.get("mock_quantidade_registros", cfg.get("busca", {}).get("quantidade_selecionados", 12)) or 12)
    mock_registros = max(4, mock_registros)
    mock_gerar_pdf = bool(controle.get("mock_gerar_pdf", True))
    random.seed(mock_seed)

    research_output_dir.mkdir(parents=True, exist_ok=True)
    prefixo = (cfg.get("saida", {}).get("prefixo") or "atividade").strip()
    topic = str(cfg.get("pesquisa", {}).get("tema") or "tema de pesquisa").strip() or "tema de pesquisa"

    bib_entries, bib_keys, titles = _mock_make_bib_entries(mock_registros, shorten_text(topic, 60))
    selected_entries, selected_pdfs = _mock_make_selected_entries(research_output_dir, prefixo, titles, bib_keys)

    if executar_pesquisa or not detect_research_paths(cfg).org_path:
        org_path = research_output_dir / f"{prefixo}.org"
        bib_path = research_output_dir / f"{prefixo}.bib"
        debug_path = research_output_dir / f"{prefixo}_debug.json"
        pdf_path = research_output_dir / f"{prefixo}.pdf"
        prisma_svg_path = research_output_dir / f"{prefixo}_prisma.svg"
        prisma_pdf_path = research_output_dir / f"{prefixo}_prisma.pdf"

        write_text(org_path, _mock_build_research_org(cfg, prefixo, titles, bib_keys))
        write_text(bib_path, "\n\n".join(bib_entries).strip() + "\n")
        write_text(debug_path, json.dumps({
            "mock_run": True,
            "generated_at": datetime.now().isoformat(),
            "proposal": {
                "tema": cfg.get("pesquisa", {}).get("tema"),
                "recorte": cfg.get("pesquisa", {}).get("recorte"),
                "objetivo": cfg.get("pesquisa", {}).get("objetivo"),
                "pergunta_pesquisa": f"Como {topic.lower()} é tratado na literatura do setor público?",
                "hipotese": "A adoção responsável depende de governança, salvaguardas éticas e gestão de riscos."
            },
            "selected_all": selected_entries,
            "selected_count": len(selected_entries),
            "mock_seed": mock_seed,
        }, ensure_ascii=False, indent=2))
        _mock_make_prisma_svg(prisma_svg_path, total=max(mock_registros + 18, mock_registros), selected=mock_registros)
        if mock_gerar_pdf:
            _mock_write_pdf(pdf_path, prefixo)
            _mock_write_pdf(prisma_pdf_path, f"{prefixo} prisma")

    research_paths = detect_research_paths(cfg)
    source_label = "corpus local" if is_local_documents_mode(cfg) else "pesquisa"
    print(f"[2/6] Carregando contexto e artefatos do {source_label}...", flush=True)
    documento_context, debug_json = build_document_context(cfg, research_paths)

    manifest_path: Path | None = None
    if executar_bundle:
        print("[3/6] Montando bundle intermediário...", flush=True)
        manifest_path = build_bundle(cfg, research_paths, documento_context, debug_json, build_bundle_dir(cfg, research_paths.root_dir))

    if not executar_documento:
        report_path = research_output_dir / "pipeline_mock_run_report.json"
        write_text(report_path, json.dumps({
            "generated_at": datetime.now().isoformat(),
            "mock_run": True,
            "document_stage": False,
            "research_root": str(research_output_dir),
            "manifest_path": str(manifest_path) if manifest_path else None,
            "artifacts": {
                "research_org": str(research_paths.org_path) if research_paths.org_path else None,
                "research_bib": str(research_paths.bib_path) if research_paths.bib_path else None,
                "research_debug": str(research_paths.debug_path) if research_paths.debug_path else None,
            }
        }, ensure_ascii=False, indent=2))
        print("Mock run concluído. Artefatos sintéticos de pesquisa foram gerados.")
        print(f"- Relatório: {report_path}")
        return 0

    documento_output_dir = build_document_output_dir(cfg, research_paths.root_dir)
    documento_output_dir.mkdir(parents=True, exist_ok=True)
    doc_type = normalize_document_type(documento.get("tipo_documento"))
    documento_prefix = (documento.get("prefixo") or ((cfg.get("saida", {}).get("prefixo") or "atividade") + f"_{DEFAULT_DOC_PREFIX_BY_TYPE.get(doc_type, 'documento')}")).strip()
    documento_bib_name = f"{documento_prefix}.bib"

    base_docs: list[SourceDoc] = []
    for entry, pdf, key in zip(selected_entries, selected_pdfs, bib_keys):
        base_docs.append(SourceDoc(
            path=str(pdf),
            kind="texto_selecionado_mock",
            label=str(entry.get("title")),
            extracted_text=_mock_lipsum(str(entry.get("title")), 7),
            summary="Documento sintético para validação estrutural.",
            bib_key=key,
            metadata={"mock_run": True, "entry": entry},
        ))

    guidance_docs = collect_guidance_docs(cfg, research_paths)
    extra_docs: list[SourceDoc] = []

    org_text = _mock_build_document_org(cfg, documento_prefix, doc_type, titles, bib_keys)
    prompt_text = textwrap.dedent(f"""    MOCK RUN — nenhum prompt real enviado à OpenAI.
    Tema: {cfg.get("pesquisa", {}).get("tema")}
    Tipo de documento: {doc_type}
    Quantidade de referências sintéticas: {len(bib_keys)}
    """)

    org_path = documento_output_dir / f"{documento_prefix}.org"
    bib_path = documento_output_dir / documento_bib_name
    context_json_path = documento_output_dir / f"{documento_prefix}_contexto.json"
    prompt_audit_path = documento_output_dir / f"{documento_prefix}_prompts_auditoria.txt"
    provenance_path = documento_output_dir / f"{documento_prefix}_proveniencia.json"
    reference_usage_json_path = documento_output_dir / f"{documento_prefix}_uso_referencias.json"
    reference_usage_md_path = documento_output_dir / f"{documento_prefix}_uso_referencias.md"
    section_limits_json_path = documento_output_dir / f"{documento_prefix}_limites_secoes.json"
    documento_pdf_path = documento_output_dir / f"{documento_prefix}.pdf" if mock_gerar_pdf else None

    write_text(org_path, org_text)
    write_text(bib_path, "\n\n".join(bib_entries).strip() + "\n")
    write_text(context_json_path, json.dumps({
        "generated_at": datetime.now().isoformat(),
        "mock_run": True,
        "documento_context": asdict(documento_context),
        "research_paths": json_safe(asdict(research_paths)),
        "config_excerpt": {
            "atividade": cfg.get("atividade", {}),
            "pesquisa": cfg.get("pesquisa", {}),
            "documento": cfg.get("documento", {}),
        },
        "bundle_manifest": str(manifest_path) if manifest_path else None,
        "context_origin_info": {"used": False, "source": "mock_run"},
        "bib_keys": bib_keys,
        "base_docs": [asdict(d) for d in base_docs],
        "guidance_docs": [asdict(d) for d in guidance_docs],
        "selected_entries": selected_entries,
        "selected_fulltext_paths": [str(p) for p in selected_pdfs],
        "extra_article_paths": [],
    }, ensure_ascii=False, indent=2, default=str))
    write_text(prompt_audit_path, prompt_text)
    write_text(provenance_path, json.dumps(
        build_provenance_payload(
            cfg=cfg,
            research_paths=research_paths,
            documento_context=documento_context,
            base_docs=base_docs,
            guidance_docs=guidance_docs,
            extra_docs=extra_docs,
            bib_keys=bib_keys,
            bib_entries=bib_entries,
            context_origin_info={"used": False, "source": "mock_run"},
            template_path=Path(str(documento.get("template_org") or "")) if documento.get("template_org") else None,
            manifest_path=manifest_path,
        ),
        ensure_ascii=False,
        indent=2,
        default=str,
    ))
    reference_usage = build_reference_usage_map(org_text, base_docs, extra_docs, bib_keys, required_keys=cfg.get("__citation_target_keys__"))
    write_text(reference_usage_json_path, json.dumps(reference_usage, ensure_ascii=False, indent=2))
    write_text(reference_usage_md_path, render_reference_usage_markdown(reference_usage))
    write_text(section_limits_json_path, json.dumps({
        "mock_run": True,
        "word_counts": count_org_words_per_top_section(org_text),
        "configured_limits": {
            "total": documento.get("limite_palavras_total"),
            "introducao": documento.get("limite_palavras_introducao"),
            "revisao": documento.get("limite_palavras_revisao"),
            "conclusao": documento.get("limite_palavras_conclusao"),
        }
    }, ensure_ascii=False, indent=2))

    if documento_pdf_path:
        _mock_write_pdf(documento_pdf_path, documento_prefix)

    if manifest_path:
        update_bundle_with_documento(
            manifest_path,
            org_path=org_path,
            bib_path=bib_path,
            context_json_path=context_json_path,
            prompt_audit_path=prompt_audit_path,
            provenance_path=provenance_path,
            extra_docs=extra_docs,
            pdf_path=documento_pdf_path,
        )

    package_dir = assemble_delivery_package(
        cfg,
        research_paths=research_paths,
        documento_output_dir=documento_output_dir,
        documento_prefix=documento_prefix,
        manifest_path=manifest_path,
        org_path=org_path,
        bib_path=bib_path,
        context_json_path=context_json_path,
        prompt_audit_path=prompt_audit_path,
        provenance_path=provenance_path,
        reference_usage_json_path=reference_usage_json_path,
        reference_usage_md_path=reference_usage_md_path,
        section_limits_json_path=section_limits_json_path,
        documento_pdf_path=documento_pdf_path,
    )

    report_path = research_output_dir / "pipeline_mock_run_report.json"
    write_text(report_path, json.dumps({
        "generated_at": datetime.now().isoformat(),
        "mock_run": True,
        "mock_seed": mock_seed,
        "mock_quantidade_registros": mock_registros,
        "mock_gerar_pdf": mock_gerar_pdf,
        "research_root": str(research_output_dir),
        "document_root": str(documento_output_dir),
        "manifest_path": str(manifest_path) if manifest_path else None,
        "delivery_package": str(package_dir) if package_dir else None,
        "artifacts": {
            "research_org": str(research_paths.org_path) if research_paths.org_path else None,
            "research_bib": str(research_paths.bib_path) if research_paths.bib_path else None,
            "research_debug": str(research_paths.debug_path) if research_paths.debug_path else None,
            "research_pdf": str(research_paths.pdf_path) if research_paths.pdf_path else None,
            "research_prisma_svg": str(research_paths.prisma_svg_path) if research_paths.prisma_svg_path else None,
            "document_org": str(org_path),
            "document_bib": str(bib_path),
            "document_context": str(context_json_path),
            "document_prompt_audit": str(prompt_audit_path),
            "document_provenance": str(provenance_path),
            "reference_usage_json": str(reference_usage_json_path),
            "reference_usage_md": str(reference_usage_md_path),
            "section_limits_json": str(section_limits_json_path),
            "document_pdf": str(documento_pdf_path) if documento_pdf_path else None,
        }
    }, ensure_ascii=False, indent=2))

    print("Mock run concluído. Toda a cadeia de artefatos sintéticos foi gerada.")
    print(f"- Relatório: {report_path}")
    mock_source_prefix = "Corpus local" if is_local_documents_mode(cfg) else "Pesquisa"
    print(f"- {mock_source_prefix} ORG: {research_paths.org_path}")
    print(f"- Documento ORG: {org_path}")
    print(f"- Documento BIB: {bib_path}")
    if documento_pdf_path:
        print(f"- Documento PDF sintético: {documento_pdf_path}")
    if package_dir:
        print(f"- Pacote final de entrega: {package_dir}")
    return 0






# ---------------------------------------------------------------------------
# Modo derivação de dissertação: reorientação temática/empírica segura
# ---------------------------------------------------------------------------
def derivation_config(cfg: dict[str, Any]) -> dict[str, Any]:
    section = cfg.get("derivacao", {}) if isinstance(cfg.get("derivacao"), dict) else {}
    return section if isinstance(section, dict) else {}


def should_run_derivation_mode(cfg: dict[str, Any]) -> bool:
    """Retorna True quando o modo de derivação/reorientação está ativo.

    Esse modo usa uma dissertação .org existente como matriz intelectual e gera
    outro documento, preservando o original. É indicado para casos como:

      dissertação sobre IA/ESG no governo federal
      -> nova dissertação derivada sobre IA aplicada ao ATESTMED
    """
    d = derivation_config(cfg)
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    return bool(
        d.get("ativo")
        or d.get("enabled")
        or d.get("modo_derivacao")
        or d.get("modo_derivacao_dissertacao")
        or documento.get("modo_derivacao")
        or documento.get("modo_derivacao_dissertacao")
    )


def resolve_derivation_base_org_path(cfg: dict[str, Any]) -> Path:
    """Resolve o .org base da dissertação que será usada como matriz."""
    d = derivation_config(cfg)
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    keys = (
        "documento_org_base",
        "documento_org_existente",
        "org_base",
        "org_existente",
        "arquivo_org_base",
        "documento_base",
    )
    for key in keys:
        raw = d.get(key) or documento.get(key)
        path = resolve_configured_path(raw, cfg) if raw else None
        if path and path.exists() and path.is_file():
            return path
    raise RuntimeError(
        "Modo derivação ativo, mas o .org base não foi localizado. "
        "Informe [derivacao].documento_org_base ou [documento].documento_org_base."
    )


def resolve_derivation_base_bib_path(cfg: dict[str, Any], base_org_path: Path, base_org_text: str | None = None) -> Path | None:
    """Resolve o .bib base da dissertação original."""
    d = derivation_config(cfg)
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    for key in ("bib_base", "bibliografia_base", "referencias_base", "arquivo_bib_base"):
        raw = d.get(key) or documento.get(key)
        path = resolve_configured_path(raw, cfg) if raw else None
        if path and path.exists() and path.is_file() and path.suffix.lower() == ".bib":
            return path
    for candidate in _org_bibliography_paths(base_org_path, base_org_text):
        if candidate.exists() and candidate.is_file():
            return candidate
    return None


def resolve_derivation_output_org_path(cfg: dict[str, Any], base_org_path: Path) -> Path:
    """Define o caminho do .org derivado, sem sobrescrever o original por padrão."""
    d = derivation_config(cfg)
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    raw_output = (
        d.get("documento_org_saida")
        or d.get("org_saida")
        or d.get("output_org_path")
        or d.get("arquivo_org_saida")
        or documento.get("documento_org_saida")
        or documento.get("org_saida")
        or documento.get("output_org_path")
    )
    if raw_output:
        path = safe_resolve_user_path(str(raw_output), base_dir=base_org_path.parent)
        if path is None:
            path = Path(os.path.expanduser(str(raw_output))).resolve()
        if path.exists() and path.is_dir():
            stem = str(d.get("prefixo_saida") or f"{base_org_path.stem}{str(d.get('sufixo_saida') or '_derivada')}").strip()
            return path / f"{stem}{base_org_path.suffix}"
        if not path.suffix:
            path = path.with_suffix(base_org_path.suffix)
        return path

    output_dir = None
    if d.get("output_dir") or documento.get("output_dir_derivacao"):
        output_dir = resolve_configured_path(d.get("output_dir") or documento.get("output_dir_derivacao"), cfg)
    if output_dir is None:
        output_dir = base_org_path.parent

    prefix = str(d.get("prefixo_saida") or documento.get("prefixo_saida_derivacao") or "").strip()
    if not prefix:
        suffix = str(d.get("sufixo_saida") or documento.get("sufixo_saida_derivacao") or "_derivada").strip()
        if not suffix:
            suffix = "_derivada"
        prefix = base_org_path.stem if base_org_path.stem.endswith(suffix) else f"{base_org_path.stem}{suffix}"

    return output_dir / f"{prefix}{base_org_path.suffix}"


def _section_values_as_list(section: dict[str, Any], *keys: str) -> list[str]:
    values: list[str] = []
    for key in keys:
        values.extend(_ensure_list_of_strings(section.get(key)))
    seen: set[str] = set()
    out: list[str] = []
    for value in values:
        if value not in seen:
            seen.add(value)
            out.append(value)
    return out


def collect_derivation_docs(cfg: dict[str, Any], base_org_path: Path) -> tuple[list[SourceDoc], list[SourceDoc]]:
    """Coleta orientações e dados locais específicos da derivação.

    Retorna duas listas:
    - guidance_docs: orientações, prompts e diretrizes;
    - local_data_docs: dados/textos locais do novo objeto empírico.
    """
    d = derivation_config(cfg)
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    config_base = get_config_base_dir(cfg)

    guidance_docs: list[SourceDoc] = []
    local_data_docs: list[SourceDoc] = []

    guidance_values = _section_values_as_list(
        d,
        "orientacoes_paths",
        "orientacao_paths",
        "orientacao_path",
        "orientacoes",
        "orientacao_inline",
        "prompt_orientacao",
        "prompt_reorientacao",
    )
    guidance_values.extend(_section_values_as_list(documento, "orientacoes_derivacao_paths", "orientacao_derivacao_inline"))

    for idx, raw in enumerate(guidance_values, start=1):
        path = safe_resolve_user_path(raw, base_dir=config_base)
        if path and path.exists() and path.is_file():
            try:
                guidance_docs.append(SourceDoc(
                    path=str(path),
                    kind="orientacao_derivacao",
                    label=path.name,
                    extracted_text=read_text_file(path, 40000),
                ))
            except Exception as exc:
                debug_print(f"Falha ao ler orientação de derivação {path}: {exc}")
        else:
            guidance_docs.append(SourceDoc(
                path=f"inline:orientacao_derivacao_{idx}",
                kind="orientacao_derivacao_inline",
                label=f"orientacao_derivacao_{idx}",
                extracted_text=shorten_text(str(raw), 40000),
            ))

    data_values = _section_values_as_list(
        d,
        "dados_locais_paths",
        "dados_locais_path",
        "dados_paths",
        "dados_path",
        "contexto_local_paths",
        "contexto_local_path",
        "arquivos_dados",
        "dados_locais_inline",
        "contexto_local_inline",
        "dados_inline",
    )
    data_values.extend(_section_values_as_list(documento, "dados_locais_derivacao_paths", "dados_locais_derivacao_inline"))

    for idx, raw in enumerate(data_values, start=1):
        path = safe_resolve_user_path(raw, base_dir=config_base)
        if path and path.exists():
            files = collect_readable_files([str(path)], base_dir=config_base)
            for file_path in files:
                try:
                    local_data_docs.append(SourceDoc(
                        path=str(file_path),
                        kind="dados_locais_derivacao",
                        label=file_path.name,
                        extracted_text=read_text_file(file_path, 50000),
                    ))
                except Exception as exc:
                    debug_print(f"Falha ao ler dado local de derivação {file_path}: {exc}")
        else:
            local_data_docs.append(SourceDoc(
                path=f"inline:dados_locais_derivacao_{idx}",
                kind="dados_locais_derivacao_inline",
                label=f"dados_locais_derivacao_{idx}",
                extracted_text=shorten_text(str(raw), 50000),
            ))

    if not guidance_docs:
        guidance_docs.append(SourceDoc(
            path=str(base_org_path),
            kind="org_base_como_orientacao",
            label=base_org_path.name,
            extracted_text=shorten_text(base_org_path.read_text(encoding="utf-8", errors="ignore"), 50000),
        ))

    return guidance_docs, local_data_docs


def collect_derivation_extra_article_docs(cfg: dict[str, Any], max_files: int = 30) -> list[SourceDoc]:
    """Coleta novas fontes acadêmicas específicas da derivação."""
    d = derivation_config(cfg)
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    raw_items: list[Any] = []
    for key in ("artigos_extras_paths", "fontes_extras_paths", "novas_fontes_paths", "referencias_extras_paths"):
        raw_items.extend(_ensure_list_of_strings(d.get(key)))
    raw_items.extend(_ensure_list_of_strings(documento.get("artigos_extras_paths")))
    files = collect_readable_files(raw_items, get_config_base_dir(cfg))[:max_files]
    docs: list[SourceDoc] = []
    for path in files:
        try:
            docs.append(SourceDoc(
                path=str(path),
                kind="artigo_extra_derivacao",
                label=path.name,
                extracted_text=read_text_file(path, 35000),
                metadata={"source_path": str(path)},
            ))
        except Exception as exc:
            debug_print(f"Falha ao ler artigo extra de derivação {path}: {exc}")
    return docs


def build_derivation_context(cfg: dict[str, Any], base_org_text: str) -> DocumentContext:
    """Monta o novo contexto temático/empírico da dissertação derivada."""
    d = derivation_config(cfg)
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    pesquisa = cfg.get("pesquisa", {}) if isinstance(cfg.get("pesquisa"), dict) else {}

    def first(*values: Any) -> str:
        for value in values:
            text = str(value or "").strip()
            if text:
                return text
        return ""

    title_match = re.search(r"(?im)^#\+TITLE:\s*(.+)$", base_org_text or "")
    base_title = title_match.group(1).strip() if title_match else None
    palavras = d.get("palavras_chave") or documento.get("palavras_chave") or pesquisa.get("palavras_chave") or []
    if isinstance(palavras, str):
        palavras = [p.strip() for p in re.split(r"[,;]", palavras) if p.strip()]

    return DocumentContext(
        tema=first(d.get("novo_tema"), d.get("tema"), documento.get("novo_tema"), documento.get("tema"), pesquisa.get("tema")),
        recorte=first(d.get("novo_recorte"), d.get("recorte"), documento.get("novo_recorte"), documento.get("recorte"), pesquisa.get("recorte")),
        objetivo=first(d.get("novo_objetivo"), d.get("objetivo"), documento.get("novo_objetivo"), documento.get("objetivo"), pesquisa.get("objetivo")),
        pergunta_pesquisa=d.get("nova_pergunta_pesquisa") or d.get("pergunta_pesquisa") or documento.get("pergunta_pesquisa") or pesquisa.get("pergunta_pesquisa"),
        hipotese=d.get("nova_hipotese") or d.get("hipotese") or documento.get("hipotese") or pesquisa.get("hipotese"),
        palavras_chave=list(palavras or []),
        titulo_sugerido=first(d.get("novo_titulo"), d.get("titulo"), documento.get("titulo_trabalho"), documento.get("titulo"), base_title),
        tipo_estudo=first(d.get("tipo_estudo"), documento.get("tipo_estudo"), pesquisa.get("tipo_estudo"), "dissertacao_derivada"),
        idiomas=list(pesquisa.get("idiomas") or ["português"]),
        modo_origem="derivacao_dissertacao",
        titulo_trabalho_base=base_title,
    )


def build_derivation_prompt(
    cfg: dict[str, Any],
    context: DocumentContext,
    base_org_text: str,
    base_bib_entries: list[str],
    extra_bib_entries: list[str],
    bib_keys: list[str],
    guidance_docs: list[SourceDoc],
    local_data_docs: list[SourceDoc],
    extra_article_docs: list[SourceDoc],
    output_bib_name: str,
    style: str,
) -> str:
    """Prompt principal do modo derivação."""
    d = derivation_config(cfg)
    estrategia = str(d.get("estrategia") or "preservar_estrutura_e_adaptar_conteudo").strip()
    foco = str(d.get("foco") or d.get("descricao_reorientacao") or "").strip()
    preservacao = str(d.get("grau_preservacao") or "alto_para_teoria_baixo_para_objeto_empirico").strip()
    min_palavras = d.get("min_palavras_total") or cfg.get("documento", {}).get("min_palavras_total") or cfg.get("documento", {}).get("limite_palavras_total")

    return textwrap.dedent(f"""
    Gere uma NOVA dissertação em Org-mode a partir de uma dissertação base já existente.

    Esta é uma tarefa de DERIVAÇÃO/REORIENTAÇÃO TEMÁTICA E EMPÍRICA, não uma simples revisão gramatical.

    Objetivo do modo:
    - preservar a dissertação original intacta;
    - usar a dissertação original como matriz intelectual, estrutural e teórica;
    - criar outro documento completo, adequado ao novo tema, novo recorte, novo objetivo e novo objeto empírico;
    - reaproveitar criticamente o que continuar válido;
    - substituir ou reescrever o que estiver preso ao objeto anterior;
    - incorporar dados locais, prompts, orientações e novas fontes fornecidas.

    Estratégia configurada: {estrategia}
    Grau de preservação desejado: {preservacao}
    Foco específico informado: {foco or 'não informado'}

    Novo contexto da dissertação derivada:
    {json.dumps(asdict(context), ensure_ascii=False, indent=2)}

    Regras obrigatórias:
    1. Retorne apenas o conteúdo completo do novo arquivo .org.
    2. Preserve o cabeçalho técnico Org/LaTeX quando ele existir e ajuste título/metadados ao novo objeto.
    3. Não sobrescreva mentalmente a dissertação original; trate-a como matriz para um novo documento.
    4. Preserve a estrutura geral quando ela for útil, mas adapte títulos, introdução, metodologia, resultados, discussão e conclusão ao novo objeto.
    5. Reaproveite o referencial teórico sobre IA, governança, accountability, risco, ESG, capacidades estatais, administração pública e temas correlatos somente quando fizer sentido para o novo objeto.
    6. Reescreva fortemente trechos que dependam do objeto empírico anterior.
    7. Incorpore os dados locais e textos operacionais como evidência empírica/contextual, distinguindo-os das fontes acadêmicas.
    8. Use as novas fontes acadêmicas quando forem relevantes para fortalecer a transposição analítica.
    9. Use apenas citações Org Cite, como [cite:@chave] ou [cite/t:@chave].
    10. Não invente chaves bibliográficas; use somente estas chaves disponíveis: {json.dumps(bib_keys, ensure_ascii=False)}.
    11. Use o estilo bibliográfico final {style.upper()}.
    12. Não crie seção manual de referências fora do mecanismo #+PRINT_BIBLIOGRAPHY:.
    13. Garanta que o documento final aponte para a bibliografia de saída: {output_bib_name}.
    14. Ao final, inclua uma conclusão compatível com o novo objeto empírico e uma agenda de pesquisa/aprimoramento, quando pertinente.
    {f'15. Meta mínima aproximada de extensão textual: {min_palavras} palavras.' if min_palavras else ''}

    Diretriz de transposição:
    - Introdução: reorientar fortemente para o novo objeto.
    - Referencial teórico: preservar e adaptar, conectando ao novo objeto.
    - Metodologia: reescrever para explicar o uso da dissertação matriz, dados locais, documentos, prompts e fontes novas.
    - Resultados/discussão: reconstruir com centralidade no novo objeto empírico.
    - Considerações finais: reescrever para responder ao novo objetivo.

    Dissertação base em Org-mode:
    {shorten_text(base_org_text, 90000)}

    Entradas BibTeX da dissertação base:
    {shorten_text(json.dumps(base_bib_entries, ensure_ascii=False, indent=2), 30000)}

    Entradas BibTeX das novas fontes adicionadas:
    {shorten_text(json.dumps(extra_bib_entries, ensure_ascii=False, indent=2), 25000)}

    Orientações de reorientação/derivação:
    {json.dumps(summarize_docs(guidance_docs, excerpt_chars=5000), ensure_ascii=False, indent=2)}

    Dados locais, textos operacionais e contexto do novo objeto:
    {json.dumps(summarize_docs(local_data_docs, excerpt_chars=6000), ensure_ascii=False, indent=2)}

    Novas fontes acadêmicas resumidas:
    {json.dumps(summarize_docs(extra_article_docs, excerpt_chars=3500), ensure_ascii=False, indent=2)}
    """).strip()


def run_derivation_mode(cfg: dict[str, Any], model_override: str | None = None) -> int:
    """Executa o modo seguro de derivação/reorientação de dissertação."""
    if not should_run_derivation_mode(cfg):
        return 0

    d = derivation_config(cfg)
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    if should_generate_only_mindmap(cfg):
        raise RuntimeError("Não ative [derivacao].ativo e somente_mapa_mental ao mesmo tempo.")

    base_org_path = resolve_derivation_base_org_path(cfg)
    base_org_text = base_org_path.read_text(encoding="utf-8", errors="ignore")
    base_bib_path = resolve_derivation_base_bib_path(cfg, base_org_path, base_org_text)
    output_org_path = resolve_derivation_output_org_path(cfg, base_org_path)
    output_org_path.parent.mkdir(parents=True, exist_ok=True)
    output_bib_path = output_org_path.with_suffix(".bib")
    output_prefix = output_org_path.stem

    print("[1/4] Carregando dissertação base para derivação...", flush=True)
    print(f"- ORG base preservado: {base_org_path}", flush=True)
    print(f"- ORG derivado: {output_org_path}", flush=True)
    if base_bib_path:
        print(f"- BIB base: {base_bib_path}", flush=True)

    backup_dir = output_org_path.parent / "_derivacao_backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    base_backup = backup_file_if_exists(base_org_path, backup_dir, label="base_preservado")
    previous_output_backup = backup_file_if_exists(output_org_path, backup_dir, label="saida_anterior")

    base_bib_entries, base_bib_keys = parse_bib_entries(base_bib_path) if base_bib_path else ([], [])
    guidance_docs, local_data_docs = collect_derivation_docs(cfg, base_org_path)
    extra_article_docs = collect_derivation_extra_article_docs(cfg, max_files=int(d.get("max_artigos_extras") or 30))

    print("[2/4] Preparando fontes novas e bibliografia derivada...", flush=True)
    client, model = make_client(model_override or cfg.get("openai", {}).get("model"))
    extra_entries: list[str] = []
    extra_keys: list[str] = []
    if extra_article_docs and bool(d.get("incluir_artigos_extras_no_bib", documento.get("incluir_artigos_extras_no_bib", True))):
        extra_article_docs, extra_entries, extra_keys = build_bib_entries_for_extra_docs(client, model, extra_article_docs, base_bib_keys)

    all_bib_entries = base_bib_entries + extra_entries
    all_bib_keys = [k for k in base_bib_keys + extra_keys if k]
    write_text(output_bib_path, "\n\n".join(all_bib_entries).strip() + ("\n" if all_bib_entries else ""))

    context = build_derivation_context(cfg, base_org_text)
    style = str(d.get("estilo_citacao") or documento.get("estilo_citacao") or cfg.get("bibliografia", {}).get("estilo_citacao") or DEFAULT_STYLE)

    print("[3/4] Gerando dissertação derivada por IA...", flush=True)
    prompt = build_derivation_prompt(
        cfg=cfg,
        context=context,
        base_org_text=base_org_text,
        base_bib_entries=base_bib_entries,
        extra_bib_entries=extra_entries,
        bib_keys=all_bib_keys,
        guidance_docs=guidance_docs,
        local_data_docs=local_data_docs,
        extra_article_docs=extra_article_docs,
        output_bib_name=output_bib_path.name,
        style=style,
    )
    resp = client.responses.create(model=model, input=prompt)
    derived_org = resp.output_text.strip()
    derived_org = ensure_document_class(derived_org, "dissertacao")
    derived_org = apply_citation_style(derived_org, output_bib_path.name, style)
    derived_org = ensure_cover_command(derived_org)
    derived_org = cleanup_generated_org(derived_org)
    derived_org = normalize_bibliography_block(derived_org)
    final_title = str(context.titulo_sugerido or d.get("novo_titulo") or documento.get("titulo") or output_prefix).strip()
    derived_org = apply_dissertation_template_metadata(derived_org, cfg, final_title)
    derived_org = render_template_placeholders(derived_org, cfg, final_title)

    original_words = count_words(base_org_text)
    derived_words = count_words(derived_org)
    if original_words >= 1000 and derived_words < int(original_words * float(d.get("limiar_minimo_proporcao_palavras") or 0.35)):
        raise RuntimeError(
            "Proteção acionada: a dissertação derivada ficou pequena demais em relação à base "
            f"({derived_words}/{original_words} palavras). Nenhum arquivo derivado foi gravado."
        )

    write_text(output_org_path, derived_org)

    prompt_audit_path = output_org_path.with_name(f"{output_prefix}_prompts_auditoria.txt")
    audit_json_path = output_org_path.with_name(f"{output_prefix}_derivacao_auditoria.json")
    reference_usage_json_path = output_org_path.with_name(f"{output_prefix}_uso_referencias.json")
    reference_usage_md_path = output_org_path.with_name(f"{output_prefix}_uso_referencias.md")
    section_limits_json_path = output_org_path.with_name(f"{output_prefix}_limites_secoes.json")

    prompt_audit = "===== modo_derivacao_dissertacao =====\n" + prompt + "\n"
    write_text(prompt_audit_path, prompt_audit)

    reference_usage = build_reference_usage_map(derived_org, [], extra_article_docs, all_bib_keys, required_keys=all_bib_keys)
    write_text(reference_usage_json_path, json.dumps(reference_usage, ensure_ascii=False, indent=2))
    write_text(reference_usage_md_path, render_reference_usage_markdown(reference_usage))
    write_text(section_limits_json_path, json.dumps({
        "word_counts": count_org_words_per_top_section(derived_org),
        "base_word_count": original_words,
        "derived_word_count": derived_words,
    }, ensure_ascii=False, indent=2))

    audit_payload = {
        "generated_at": datetime.now().isoformat(),
        "mode": "derivacao_dissertacao",
        "base_org_path": str(base_org_path),
        "base_bib_path": str(base_bib_path) if base_bib_path else None,
        "output_org_path": str(output_org_path),
        "output_bib_path": str(output_bib_path),
        "base_backup": str(base_backup) if base_backup else None,
        "previous_output_backup": str(previous_output_backup) if previous_output_backup else None,
        "documento_context": asdict(context),
        "guidance_docs": [asdict(doc) for doc in guidance_docs],
        "local_data_docs": [asdict(doc) for doc in local_data_docs],
        "extra_article_docs": [asdict(doc) for doc in extra_article_docs],
        "base_bib_keys": base_bib_keys,
        "extra_bib_keys": extra_keys,
        "all_bib_keys": all_bib_keys,
        "word_counts": {
            "base": original_words,
            "derived": derived_words,
        },
        "reference_usage": reference_usage,
    }
    write_text(audit_json_path, json.dumps(audit_payload, ensure_ascii=False, indent=2, default=str))

    pdf_path: Path | None = None
    export_pdf = bool(d.get("exportar_pdf", d.get("recompilar_pdf", documento.get("exportar_pdf", cfg.get("saida", {}).get("exportar_pdf", False)))))
    if export_pdf:
        print("[4/4] Compilando PDF da dissertação derivada...", flush=True)
        latex = cfg.get("latex", {}) if isinstance(cfg.get("latex"), dict) else {}
        academic_writing = resolve_configured_path(latex.get("org_latex_class_init"), cfg) if latex.get("org_latex_class_init") else None
        emacs_init = resolve_configured_path(latex.get("emacs_init"), cfg) if latex.get("emacs_init") else None
        latex_extra_path = resolve_configured_path(latex.get("latex_extra_path"), cfg) if latex.get("latex_extra_path") else None
        try:
            pdf_path = run_compile_sequence(output_org_path, emacs_init=emacs_init, academic_writing=academic_writing, latex_extra_path=latex_extra_path)
        except Exception as exc:
            pdf_error_path = output_org_path.with_name(f"{output_prefix}_pdf_erro.txt")
            write_text(
                pdf_error_path,
                "A compilação do PDF da dissertação derivada falhou, mas o .org e o .bib foram preservados.\n\n"
                f"Documento ORG base: {base_org_path}\n"
                f"Documento ORG derivado: {output_org_path}\n"
                f"Documento BIB derivado: {output_bib_path}\n\n"
                f"Erro:\n{exc}\n",
            )
            print("Aviso: a compilação do PDF falhou, mas os artefatos derivados foram preservados.")
            print(f"- Log simplificado da falha: {pdf_error_path}")

    print("\nDissertação derivada gerada:")
    print(f"- ORG base preservado: {base_org_path}")
    print(f"- ORG derivado: {output_org_path}")
    print(f"- BIB derivado: {output_bib_path}")
    print(f"- Auditoria: {audit_json_path}")
    print(f"- Prompts: {prompt_audit_path}")
    print(f"- Uso de referências: {reference_usage_json_path}")
    if pdf_path:
        print(f"- PDF derivado: {pdf_path}")
    return 0

def should_generate_only_mindmap(cfg: dict[str, Any]) -> bool:
    """Modo incremental: gera apenas o mapa mental sobre artefatos já existentes."""
    mm = mindmap_config(cfg)
    return bool(
        should_generate_mindmap(cfg)
        and (
            mm.get("somente_mapa_mental")
            or mm.get("only_mindmap")
            or mm.get("gerar_apenas_mapa")
            or mm.get("modo") == "somente_mapa_mental"
        )
    )


def resolve_existing_document_org_path(cfg: dict[str, Any]) -> Path:
    """Resolve o .org já gerado que receberá o mapa mental.

    Ordem de prioridade:
    1. [mapa_mental].documento_org_existente
    2. [documento].documento_org_existente
    3. [documento].output_dir + [documento].prefixo + .org
    """
    mm = mindmap_config(cfg)
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    for raw in (mm.get("documento_org_existente"), documento.get("documento_org_existente")):
        path = resolve_configured_path(raw, cfg) if raw else None
        if path and path.exists() and path.is_file():
            return path

    # Fallback: deriva a partir de output_dir/prefixo do documento.
    prefix = str(documento.get("prefixo") or "").strip()
    if not prefix:
        doc_type = normalize_document_type(documento.get("tipo_documento"))
        prefix = (cfg.get("saida", {}).get("prefixo") or "atividade") + f"_{DEFAULT_DOC_PREFIX_BY_TYPE.get(doc_type, 'documento')}"
    output_dir = resolve_configured_path(documento.get("output_dir"), cfg) if documento.get("output_dir") else None
    if output_dir is None:
        output_dir = get_bundle_root_from_config(cfg) / "output" / "documento"
    create_subdir = bool(documento.get("criar_subdiretorio", True))
    doc_dir = output_dir / prefix if create_subdir and output_dir.name != prefix else output_dir
    candidate = doc_dir / f"{prefix}.org"
    if candidate.exists() and candidate.is_file():
        return candidate
    raise RuntimeError(
        "Modo somente_mapa_mental ativo, mas não foi possível localizar o .org existente. "
        "Informe [mapa_mental].documento_org_existente no TOML."
    )


def resolve_existing_context_json_path(cfg: dict[str, Any], org_path: Path) -> Path | None:
    mm = mindmap_config(cfg)
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    for raw in (mm.get("contexto_json_existente"), documento.get("contexto_json_existente")):
        path = resolve_configured_path(raw, cfg) if raw else None
        if path and path.exists() and path.is_file():
            return path
    candidate = org_path.with_name(f"{org_path.stem}_contexto.json")
    return candidate if candidate.exists() and candidate.is_file() else None


def _source_doc_from_dict(raw: Any) -> SourceDoc | None:
    if not isinstance(raw, dict):
        return None
    try:
        return SourceDoc(
            path=str(raw.get("path") or raw.get("source_path") or ""),
            kind=str(raw.get("kind") or "documento_existente"),
            label=str(raw.get("label") or Path(str(raw.get("path") or "documento")).name),
            extracted_text=str(raw.get("extracted_text") or raw.get("summary") or raw.get("excerpt") or ""),
            summary=raw.get("summary"),
            bib_key=raw.get("bib_key"),
            metadata=raw.get("metadata") if isinstance(raw.get("metadata"), dict) else {},
        )
    except Exception:
        return None


def load_docs_from_fulltext_cache_for_mindmap(cfg: dict[str, Any], org_path: Path) -> list[SourceDoc]:
    """Fallback quando o contexto.json não existir ou não trouxer base_docs."""
    mm = mindmap_config(cfg)
    raw_dir = mm.get("fulltext_cache_dir") or mm.get("cache_dir") or mm.get("documentos_dir")
    cache_dir = resolve_configured_path(raw_dir, cfg) if raw_dir else None
    if cache_dir is None:
        # tenta localizar um *_fulltext_cache próximo ao documento final
        roots = [org_path.parent, org_path.parent.parent, get_bundle_root_from_config(cfg) / "output" / "corpus_local"]
        for root in roots:
            try:
                matches = sorted(root.rglob("*_fulltext_cache")) if root.exists() else []
            except Exception:
                matches = []
            if matches:
                cache_dir = matches[0]
                break
    docs: list[SourceDoc] = []
    if cache_dir and cache_dir.exists() and cache_dir.is_dir():
        max_docs = int(mm.get("max_docs_contexto") or 20)
        max_chars = int(mm.get("max_caracteres_por_texto") or 18000)
        for path in sorted(p for p in cache_dir.iterdir() if p.is_file() and p.suffix.lower() in READABLE_SUFFIXES)[:max_docs]:
            try:
                docs.append(SourceDoc(
                    path=str(path),
                    kind="texto_fulltext_cache_existente",
                    label=path.stem.replace("_", " "),
                    extracted_text=read_text_file(path, max_chars=max_chars),
                ))
            except Exception as exc:
                debug_print(f"Falha ao ler documento para mapa mental {path}: {exc}")
    return docs


def build_context_from_existing_artifacts(cfg: dict[str, Any], org_path: Path, context_json_path: Path | None) -> tuple[DocumentContext, list[SourceDoc], list[SourceDoc], list[dict[str, Any]], dict[str, Any]]:
    payload: dict[str, Any] = {}
    if context_json_path:
        try:
            payload = json.loads(context_json_path.read_text(encoding="utf-8"))
        except Exception as exc:
            debug_print(f"Falha ao ler contexto existente {context_json_path}: {exc}")
            payload = {}

    ctx_raw = payload.get("documento_context") if isinstance(payload.get("documento_context"), dict) else {}
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}
    pesquisa = cfg.get("pesquisa", {}) if isinstance(cfg.get("pesquisa"), dict) else {}
    atividade = cfg.get("atividade", {}) if isinstance(cfg.get("atividade"), dict) else {}

    context = DocumentContext(
        tema=str(ctx_raw.get("tema") or documento.get("tema") or pesquisa.get("tema") or atividade.get("tema") or ""),
        recorte=str(ctx_raw.get("recorte") or documento.get("recorte") or pesquisa.get("recorte") or ""),
        objetivo=str(ctx_raw.get("objetivo") or documento.get("objetivo") or pesquisa.get("objetivo") or ""),
        pergunta_pesquisa=ctx_raw.get("pergunta_pesquisa") or documento.get("pergunta_pesquisa") or pesquisa.get("pergunta_pesquisa"),
        hipotese=ctx_raw.get("hipotese") or documento.get("hipotese") or pesquisa.get("hipotese"),
        palavras_chave=list(ctx_raw.get("palavras_chave") or pesquisa.get("palavras_chave") or []),
        titulo_sugerido=ctx_raw.get("titulo_sugerido") or documento.get("titulo") or documento.get("titulo_trabalho") or atividade.get("titulo"),
        tipo_estudo=ctx_raw.get("tipo_estudo") or pesquisa.get("tipo_estudo"),
        idiomas=list(ctx_raw.get("idiomas") or pesquisa.get("idiomas") or ["português"]),
        modo_origem=str(ctx_raw.get("modo_origem") or "documento_existente"),
        titulo_trabalho_base=ctx_raw.get("titulo_trabalho_base") or documento.get("titulo_trabalho"),
    )

    base_docs = [d for d in (_source_doc_from_dict(x) for x in (payload.get("base_docs") or [])) if d]
    guidance_docs = [d for d in (_source_doc_from_dict(x) for x in (payload.get("guidance_docs") or [])) if d]
    selected_corpus_catalog = payload.get("selected_corpus_catalog") if isinstance(payload.get("selected_corpus_catalog"), list) else []

    if not base_docs:
        base_docs = load_docs_from_fulltext_cache_for_mindmap(cfg, org_path)
    if not guidance_docs:
        guidance_docs = collect_guidance_docs(cfg, ResearchPaths(root_dir=org_path.parent, org_path=org_path, bib_path=org_path.with_suffix(".bib") if org_path.with_suffix(".bib").exists() else None))

    return context, base_docs, guidance_docs, selected_corpus_catalog, payload


def resolve_mindmap_output_org_path(cfg: dict[str, Any], source_org_path: Path) -> Path:
    """Define onde o .org com mapa mental será gravado.

    Por segurança, no modo somente_mapa_mental o padrão NÃO é mais sobrescrever
    o .org base informado em [mapa_mental].documento_org_existente. O padrão é
    criar/atualizar uma cópia derivada no mesmo diretório:

        arquivo.org -> arquivo_com_mapa.org

    Para manter o comportamento antigo, o usuário precisa optar explicitamente:

        [mapa_mental]
        sobrescrever_org_existente = true

    Também é possível informar um caminho de saída explícito:

        [mapa_mental]
        documento_org_saida = "/caminho/arquivo_com_mapa.org"
    """
    mm = mindmap_config(cfg)
    documento = cfg.get("documento", {}) if isinstance(cfg.get("documento"), dict) else {}

    raw_output = (
        mm.get("documento_org_saida")
        or mm.get("org_saida")
        or mm.get("output_org_path")
        or mm.get("arquivo_org_saida")
        or documento.get("documento_org_saida")
        or documento.get("org_saida")
        or documento.get("output_org_path")
    )
    if raw_output:
        path = safe_resolve_user_path(str(raw_output), base_dir=source_org_path.parent)
        if path is None:
            path = Path(os.path.expanduser(str(raw_output))).resolve()
        if path.exists() and path.is_dir():
            return path / f"{source_org_path.stem}_com_mapa{source_org_path.suffix}"
        if not path.suffix:
            path = path.with_suffix(source_org_path.suffix)
        return path

    overwrite = bool(
        mm.get("sobrescrever_org_existente")
        or mm.get("atualizar_org_existente")
        or mm.get("overwrite_org")
        or documento.get("sobrescrever_org_existente")
        or documento.get("atualizar_org_existente")
        or documento.get("overwrite_org")
    )
    if overwrite:
        return source_org_path

    suffix = str(mm.get("sufixo_org_saida") or documento.get("sufixo_org_saida") or "_com_mapa").strip()
    if not suffix:
        suffix = "_com_mapa"
    if source_org_path.stem.endswith(suffix):
        return source_org_path
    return source_org_path.with_name(f"{source_org_path.stem}{suffix}{source_org_path.suffix}")


def backup_file_if_exists(path: Path, backup_dir: Path, label: str = "backup") -> Path | None:
    """Cria backup de um arquivo existente sem interromper o fluxo."""
    if not path.exists() or not path.is_file():
        return None
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / f"{path.stem}_{label}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{path.suffix}"
    shutil.copy2(path, backup_path)
    return backup_path


def _org_bibliography_paths(org_path: Path, org_text: str | None = None) -> list[Path]:
    """Localiza arquivos .bib associados a um .org.

    A função considera, nesta ordem:
    1. o .bib com o mesmo stem do .org;
    2. caminhos indicados em linhas #+BIBLIOGRAPHY:.

    Os caminhos relativos são resolvidos em relação ao diretório do próprio .org.
    """
    candidates: list[Path] = []
    seen: set[str] = set()

    def add(raw_path: Any) -> None:
        raw = str(raw_path or "").strip()
        if not raw:
            return
        raw = raw.removeprefix("file:").strip()
        raw = raw.strip("'\"")
        if not raw:
            return
        path = Path(os.path.expanduser(raw))
        if not path.is_absolute():
            path = (org_path.parent / path).resolve()
        else:
            path = path.resolve()
        if path.suffix.lower() != ".bib":
            return
        try:
            key = str(path.resolve())
        except Exception:
            key = str(path)
        if key in seen:
            return
        seen.add(key)
        candidates.append(path)

    add(org_path.with_suffix(".bib"))

    text = org_text
    if text is None and org_path.exists() and org_path.is_file():
        try:
            text = org_path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            text = ""

    for match in re.finditer(r"(?im)^\s*#\+BIBLIOGRAPHY:\s+(.+?)\s*$", text or ""):
        raw_line = match.group(1).strip()
        try:
            parts = shlex.split(raw_line)
        except Exception:
            parts = raw_line.split()
        for part in parts:
            add(part)

    return candidates


def copy_bibliography_for_mindmap_output(
    source_org_path: Path,
    output_org_path: Path,
    source_org_text: str | None = None,
    output_org_text: str | None = None,
) -> list[Path]:
    """Copia bibliografia necessária quando o modo seguro cria *_com_mapa.org.

    Ao recompilar um .org derivado, algumas etapas do fluxo LaTeX/Biber passam
    a procurar um .bib com o mesmo stem do .org de saída. Exemplo:

        atividade.org -> atividade.bib
        atividade_com_mapa.org -> atividade_com_mapa.bib

    Esta rotina preserva o .org base e copia o .bib original para o stem do
    arquivo derivado, além de copiar bibliografias referenciadas no #+BIBLIOGRAPHY
    para o diretório de saída quando necessário.
    """
    copied: list[Path] = []
    try:
        if source_org_path.resolve() == output_org_path.resolve():
            return copied
    except Exception:
        if str(source_org_path) == str(output_org_path):
            return copied

    output_org_path.parent.mkdir(parents=True, exist_ok=True)
    source_bibs = [p for p in _org_bibliography_paths(source_org_path, source_org_text) if p.exists() and p.is_file()]

    output_stem_bib = output_org_path.with_suffix(".bib")
    if source_bibs:
        primary_bib = source_bibs[0]
        try:
            same_file = primary_bib.resolve() == output_stem_bib.resolve()
        except Exception:
            same_file = str(primary_bib) == str(output_stem_bib)
        if not same_file:
            shutil.copy2(primary_bib, output_stem_bib)
            copied.append(output_stem_bib)

    for bib in source_bibs:
        dest = output_org_path.parent / bib.name
        try:
            same_file = bib.resolve() == dest.resolve()
        except Exception:
            same_file = str(bib) == str(dest)
        if same_file or dest.exists():
            continue
        shutil.copy2(bib, dest)
        copied.append(dest)

    if source_bibs:
        primary_bib = source_bibs[0]
        for out_bib in _org_bibliography_paths(output_org_path, output_org_text):
            if out_bib.exists():
                continue
            out_bib.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(primary_bib, out_bib)
            copied.append(out_bib)

    unique: list[Path] = []
    seen: set[str] = set()
    for path in copied:
        try:
            key = str(path.resolve())
        except Exception:
            key = str(path)
        if key not in seen:
            seen.add(key)
            unique.append(path)
    return unique

def run_only_mindmap_on_existing_document(cfg: dict[str, Any], model_override: str | None = None) -> int:
    """Gera mapa mental usando artefatos já produzidos, sem reescrever o documento-base.

    Implementação segura:
    - lê o .org indicado em [mapa_mental].documento_org_existente;
    - preserva esse .org intacto por padrão;
    - grava o resultado em *_com_mapa.org, salvo se sobrescrever_org_existente=true;
    - recompila o PDF a partir do .org derivado.
    """
    if not should_generate_only_mindmap(cfg):
        return 0

    source_org_path = resolve_existing_document_org_path(cfg)
    output_org_path = resolve_mindmap_output_org_path(cfg, source_org_path)
    output_org_path.parent.mkdir(parents=True, exist_ok=True)

    context_json_path = resolve_existing_context_json_path(cfg, source_org_path)
    documento_output_dir = output_org_path.parent
    documento_prefix = output_org_path.stem

    original_org_text = source_org_path.read_text(encoding="utf-8", errors="ignore")
    org_text = original_org_text

    context, base_docs, guidance_docs, selected_corpus_catalog, context_payload = build_context_from_existing_artifacts(cfg, source_org_path, context_json_path)
    print("[1/3] Carregando documento existente para mapa mental...", flush=True)
    print(f"- Documento ORG base: {source_org_path}", flush=True)
    if output_org_path != source_org_path:
        print(f"- Documento ORG com mapa mental: {output_org_path}", flush=True)
        print("- O ORG base será preservado sem sobrescrita.", flush=True)
    else:
        print("- Atenção: sobrescrita do ORG base ativada explicitamente.", flush=True)
    if context_json_path:
        print(f"- Contexto existente: {context_json_path}", flush=True)
    if base_docs:
        print(f"- Textos de contexto disponíveis: {len(base_docs)}", flush=True)

    print("[2/3] Gerando mapa mental PlantUML sem reescrever o documento-base...", flush=True)
    client, model = make_client(model_override or cfg.get("openai", {}).get("model"))
    org_text, mindmap_info, mindmap_prompt_text = maybe_generate_activity_mindmap(
        client=client,
        model=model,
        cfg=cfg,
        context=context,
        org_text=org_text,
        base_docs=base_docs,
        guidance_docs=guidance_docs,
        selected_corpus_catalog=selected_corpus_catalog,
        documento_output_dir=documento_output_dir,
        documento_prefix=documento_prefix,
    )

    backup_dir = source_org_path.parent / "_mindmap_backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    source_backup_path = backup_dir / f"{source_org_path.stem}_base_{datetime.now().strftime('%Y%m%d_%H%M%S')}{source_org_path.suffix}"
    write_text(source_backup_path, original_org_text)

    previous_output_backup = None
    if output_org_path.exists():
        previous_output_backup = backup_file_if_exists(output_org_path, backup_dir, label="saida_anterior")

    original_words = count_words(original_org_text)
    new_words = count_words(org_text)
    if original_words >= 300 and new_words < int(original_words * 0.60):
        raise RuntimeError(
            "Proteção acionada: a inserção do mapa mental reduziria o documento de "
            f"{original_words} para {new_words} palavras. Nenhum .org de saída foi gravado. "
            f"Backup do ORG base criado em: {source_backup_path}"
        )

    write_text(output_org_path, org_text)
    copied_bib_paths = copy_bibliography_for_mindmap_output(
        source_org_path=source_org_path,
        output_org_path=output_org_path,
        source_org_text=original_org_text,
        output_org_text=org_text,
    )
    print(f"- Backup do ORG base: {source_backup_path}", flush=True)
    for copied_bib in copied_bib_paths:
        print(f"- Bibliografia copiada para recompilação: {copied_bib}", flush=True)
    if previous_output_backup:
        print(f"- Backup da saída anterior: {previous_output_backup}", flush=True)

    source_prompt_audit_path = source_org_path.with_name(f"{source_org_path.stem}_prompts_auditoria.txt")
    prompt_audit_path = output_org_path.with_name(f"{documento_prefix}_prompts_auditoria.txt")
    audit_append = "\n\n===== only_mindmap_existing_document =====\n"
    audit_append += f"Documento ORG base: {source_org_path}\n"
    audit_append += f"Documento ORG de saída: {output_org_path}\n"
    audit_append += (mindmap_prompt_text or "") + "\n"
    audit_append += "\n===== mindmap_info =====\n" + json.dumps(mindmap_info or {}, ensure_ascii=False, indent=2) + "\n"
    if prompt_audit_path.exists():
        prompt_audit_path.write_text(prompt_audit_path.read_text(encoding="utf-8", errors="ignore") + audit_append, encoding="utf-8")
    elif source_prompt_audit_path.exists() and source_prompt_audit_path != prompt_audit_path:
        base_audit = source_prompt_audit_path.read_text(encoding="utf-8", errors="ignore")
        write_text(prompt_audit_path, base_audit.rstrip() + audit_append)
    else:
        write_text(prompt_audit_path, audit_append.lstrip())

    output_context_json_path = output_org_path.with_name(f"{documento_prefix}_contexto.json")
    context_payload["mindmap_info"] = mindmap_info
    context_payload["mindmap_updated_at"] = datetime.now().isoformat()
    context_payload["source_org_path"] = str(source_org_path)
    context_payload["output_org_path"] = str(output_org_path)
    write_text(output_context_json_path, json.dumps(context_payload, ensure_ascii=False, indent=2, default=str))

    pdf_path: Path | None = None
    mm = mindmap_config(cfg)
    recompi = bool(mm.get("recompilar_pdf", mm.get("exportar_pdf", cfg.get("documento", {}).get("exportar_pdf", False))))
    if recompi:
        print("[3/3] Recompilando PDF do documento com mapa mental...", flush=True)
        latex = cfg.get("latex", {}) if isinstance(cfg.get("latex"), dict) else {}
        academic_writing = resolve_configured_path(latex.get("org_latex_class_init"), cfg) if latex.get("org_latex_class_init") else None
        emacs_init = resolve_configured_path(latex.get("emacs_init"), cfg) if latex.get("emacs_init") else None
        latex_extra_path = resolve_configured_path(latex.get("latex_extra_path"), cfg) if latex.get("latex_extra_path") else None
        try:
            pdf_path = run_compile_sequence(output_org_path, emacs_init=emacs_init, academic_writing=academic_writing, latex_extra_path=latex_extra_path)
        except Exception as exc:
            pdf_error_path = output_org_path.with_name(f"{documento_prefix}_mapa_mental_pdf_erro.txt")
            write_text(pdf_error_path, f"A recompilação do PDF após inserir o mapa mental falhou, mas o .org de saída foi preservado.\n\nDocumento ORG base: {source_org_path}\nDocumento ORG de saída: {output_org_path}\n\nErro:\n{exc}\n")
            print("Aviso: a recompilação do PDF falhou, mas o .org com o mapa mental foi preservado.")
            print(f"- Log simplificado da falha: {pdf_error_path}")

    print("\nMapa mental gerado sobre arquivos existentes:")
    print(f"- Documento ORG base preservado: {source_org_path}")
    print(f"- Documento ORG com mapa mental: {output_org_path}")
    if mindmap_info:
        print(f"- Mapa mental PlantUML: {mindmap_info.get('puml_path')}")
        if mindmap_info.get("image_path"):
            print(f"- Mapa mental imagem: {mindmap_info.get('image_path')}")
        if mindmap_info.get("render_error"):
            print(f"- Aviso de renderização: {mindmap_info.get('render_error')}")
    print(f"- Auditoria de prompts: {prompt_audit_path}")
    print(f"- Contexto de saída: {output_context_json_path}")
    if pdf_path:
        print(f"- PDF recompilado: {pdf_path}")
    return 0

def main() -> int:
    global DEBUG
    load_env()
    args = parse_args()
    DEBUG = bool(args.debug)

    cfg_path = Path(args.config).expanduser().resolve()
    cfg = load_config(cfg_path)
    cfg["__config_dir__"] = str(cfg_path.parent)
    pipeline = cfg.get("pipeline", {})

    # Modo derivação: cria uma nova dissertação a partir de um .org base,
    # novas orientações, dados locais e fontes extras, preservando o original.
    if should_run_derivation_mode(cfg):
        return run_derivation_mode(cfg, model_override=args.model)

    # Modo incremental: usa o .org/contexto já gerados e produz apenas o mapa mental.
    # Não prepara corpus local, não reescreve atividade/dissertação e não altera bibliografia.
    if should_generate_only_mindmap(cfg):
        return run_only_mindmap_on_existing_document(cfg, model_override=args.model)

    executar_pesquisa = bool(pipeline.get("executar_pesquisa", True))
    executar_documento = bool(pipeline.get("executar_documento", True))
    executar_bundle = bool(pipeline.get("executar_bundle", pipeline.get("criar_bundle", True)))
    dry_run = bool(cfg.get("controle", {}).get("dry_run", False))
    mock_run = bool(cfg.get("controle", {}).get("mock_run", False))

    local_research_paths: ResearchPaths | None = None
    if is_local_documents_mode(cfg):
        print("[1/6] Preparando corpus local de documentos/ZIP...", flush=True)
        local_research_paths = prepare_local_corpus(cfg)
        pipeline = cfg.get("pipeline", {})
        executar_pesquisa = False

    research_script = resolve_configured_path(pipeline.get("script_pesquisa") or DEFAULT_RESEARCH_SCRIPT, cfg)
    if research_script is None:
        research_script = Path(DEFAULT_RESEARCH_SCRIPT).expanduser().resolve()

    # prepara TOML temporário só com as seções aceitas pelo gerador de pesquisa
    research_output_dir = detect_research_output_dir(cfg)
    temp_cfg_path = research_output_dir / "pipeline_research_config.toml"
    research_output_dir.mkdir(parents=True, exist_ok=True)
    research_cfg = filter_research_config(cfg, research_output_dir)
    write_text(temp_cfg_path, dict_to_toml(research_cfg))

    research_paths = local_research_paths or detect_research_paths(cfg)

    if dry_run:
        report = build_dry_run_report(cfg, research_cfg, research_script, research_paths)
        dry_run_path = research_output_dir / "pipeline_dry_run_report.json"
        write_text(dry_run_path, json.dumps(report, ensure_ascii=False, indent=2))
        print("Dry run concluído. Nenhuma chamada à OpenAI ou execução pesada foi realizada.")
        print(f"- Relatório: {dry_run_path}")
        return 0

    if mock_run:
        return run_mock_pipeline(
            cfg,
            research_script=research_script,
            research_output_dir=research_output_dir,
            temp_cfg_path=temp_cfg_path,
            executar_pesquisa=executar_pesquisa,
            executar_documento=executar_documento,
            executar_bundle=executar_bundle,
        )

    if executar_pesquisa:
        print("[1/6] Executando etapa de pesquisa...", flush=True)
        run_research_stage(temp_cfg_path, research_script)
        research_paths = detect_research_paths(cfg)
    if research_paths.org_path is None:
        raise RuntimeError(f"Não foi possível localizar o .org da pesquisa em {research_paths.root_dir}.")

    source_label = "corpus local" if is_local_documents_mode(cfg) else "pesquisa"
    print(f"[2/6] Carregando contexto e artefatos do {source_label}...", flush=True)
    documento_context, debug_json = build_document_context(cfg, research_paths)

    manifest_path: Path | None = None
    if executar_bundle:
        print("[3/6] Montando bundle intermediário...", flush=True)
        manifest_path = build_bundle(cfg, research_paths, documento_context, debug_json, build_bundle_dir(cfg, research_paths.root_dir))

    if not executar_documento:
        if is_local_documents_mode(cfg):
            print("Corpus local preparado. Etapa do documento desativada no TOML. O .org intermediário foi preservado.")
        else:
            print("Pesquisa concluída. Etapa do documento desativada no TOML. O .org da pesquisa foi preservado.")
        if manifest_path:
            print(f"Bundle: {manifest_path}")
        return 0

    documento = cfg.get("documento", {})
    doc_type = normalize_document_type(documento.get("tipo_documento"))
    doc_label = document_type_label(doc_type)
    doc_label_with_article = document_type_label(doc_type, article=True)

    print(f"[4/6] Inicializando cliente OpenAI para geração de {doc_label_with_article}...", flush=True)
    client, model = make_client(args.model or cfg.get("openai", {}).get("model"))

    local_bib_revision_info: dict[str, Any] | None = None
    if is_local_documents_mode(cfg):
        local_cfg = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
        if bool(local_cfg.get("gerar_bib_revisado_ia", local_cfg.get("gerar_bib_revisado_por_ia", False))):
            print("[4/6] Gerando .bib revisado por IA para o corpus local...", flush=True)
            local_bib_revision_info = maybe_generate_revised_local_bib(client, model, cfg, research_paths)
            if local_bib_revision_info and local_bib_revision_info.get("generated"):
                print(f"[4/6] .bib revisado gerado: {local_bib_revision_info.get('bib_path')}", flush=True)
            elif local_bib_revision_info:
                print(f"[4/6] .bib revisado não gerado: {local_bib_revision_info.get('reason')}", flush=True)

    # documento output
    documento_output_dir = build_document_output_dir(cfg, research_paths.root_dir)
    documento_output_dir.mkdir(parents=True, exist_ok=True)
    documento_prefix = (documento.get("prefixo") or ((cfg.get("saida", {}).get("prefixo") or "atividade") + f"_{DEFAULT_DOC_PREFIX_BY_TYPE.get(doc_type, 'documento')}" )).strip()
    documento_bib_name = f"{documento_prefix}.bib"

    # template do documento
    # prioridade: template_paper.org; fallback organizacional: template_research.org
    template_path = None
    doc_type = normalize_document_type(documento.get("tipo_documento"))
    if documento.get("template_org"):
        template_path = resolve_configured_path(documento["template_org"], cfg)
    else:
        template_path = find_fallback_template(doc_type)

    template_text = read_template_raw(template_path) if template_path and template_path.exists() else build_default_document_template(doc_type)
    template_text = normalize_snippet_placeholders(template_text)

    # docs e bibliografia
    base_docs = collect_base_docs(cfg, research_paths)
    guidance_docs = collect_guidance_docs(cfg, research_paths)
    extra_docs = collect_extra_article_docs(cfg)
    if extra_docs:
        debug_print(f"Artigos extras carregados: {len(extra_docs)}")
    base_docs.extend(extra_docs)

    documento = cfg.get("documento", {})
    usar_bib_da_pesquisa = bool(documento.get("usar_bib_da_pesquisa", True))
    incluir_artigos_extras_no_bib = bool(documento.get("incluir_artigos_extras_no_bib", True))
    style = (documento.get("estilo_citacao") or cfg.get("bibliografia", {}).get("estilo_citacao") or DEFAULT_STYLE)

    bib_entries, bib_keys = (parse_bib_entries(research_paths.bib_path) if usar_bib_da_pesquisa else ([], []))
    if extra_docs and incluir_artigos_extras_no_bib:
        extra_docs, extra_bib_entries, extra_bib_keys = build_bib_entries_for_extra_docs(client, model, extra_docs, bib_keys)
        bib_entries.extend(extra_bib_entries)
        bib_keys.extend(extra_bib_keys)
    if not bib_entries:
        debug_print("Nenhum .bib disponível para o documento; bibliografia ficará vazia até revisão posterior.")

    bib_dedup_report: dict[str, Any] | None = None
    bib_alias_map: dict[str, str] = {}
    local_cfg_for_dedup = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    dedup_enabled = bool(local_cfg_for_dedup.get("deduplicar_bib", local_cfg_for_dedup.get("deduplicar_referencias", True)))
    if bib_entries and dedup_enabled:
        before_bib_count = len(bib_entries)
        bib_entries, bib_keys, bib_alias_map, bib_dedup_report = deduplicate_bib_entries(bib_entries, preferred_keys=bib_keys)
        apply_bib_key_aliases(cfg, research_paths, bib_alias_map)
        if before_bib_count != len(bib_entries):
            print(f"[4/6] Referências deduplicadas: {before_bib_count} -> {len(bib_entries)} entrada(s).", flush=True)
        if usar_bib_da_pesquisa and research_paths.bib_path:
            write_text(research_paths.bib_path, "\n\n".join(bib_entries).strip() + ("\n" if bib_entries else ""))
            if bib_dedup_report and bool(local_cfg_for_dedup.get("salvar_diagnostico_metadados", True)):
                dedup_path = research_paths.bib_path.with_name(research_paths.bib_path.stem + "_deduplicacao_bib.json")
                write_text(dedup_path, json.dumps(bib_dedup_report, ensure_ascii=False, indent=2, default=str))

    assign_bib_keys_to_selected_docs(base_docs, research_paths, bib_entries)
    apply_local_revised_bib_keys_to_docs(cfg, base_docs)
    fulltext_citation_keys, unresolved_fulltext_pdfs = build_fulltext_cache_citation_keys(research_paths, bib_entries)
    if (cfg.get("__local_external_bib_path__") or cfg.get("__local_revised_bib_path__")) and bib_keys:
        # Em documentos locais, se o usuário informou uma bibliografia própria
        # ou o pipeline gerou um .bib revisado por IA, ela passa a ser a fonte
        # de verdade das chaves disponíveis. Isso evita citações sintéticas
        # derivadas de nomes de arquivos e impede referências indefinidas.
        fulltext_citation_keys = list(dict.fromkeys(bib_keys))
        unresolved_fulltext_pdfs = []
        cfg["__citation_target_keys__"] = fulltext_citation_keys
        cfg["__citation_target_source__"] = "local_revised_bib" if cfg.get("__local_revised_bib_path__") else "local_external_bib"
    elif fulltext_citation_keys:
        cfg["__citation_target_keys__"] = fulltext_citation_keys
        cfg["__citation_target_source__"] = "fulltext_cache"
    else:
        cfg["__citation_target_keys__"] = sorted({d.bib_key for d in base_docs if d.kind.startswith("texto_selecionado") and d.bib_key})
        cfg["__citation_target_source__"] = "selected_base_docs"
    if unresolved_fulltext_pdfs:
        debug_print(f"PDFs do fulltext_cache sem chave BibTeX mapeada: {len(unresolved_fulltext_pdfs)}")
    selected_corpus_catalog = build_selected_corpus_catalog(research_paths, bib_entries, limit=max(40, len(fulltext_citation_keys) + 10))

    documento_context, inferred_context_info = infer_missing_document_context(
        client=client,
        model=model,
        cfg=cfg,
        context=documento_context,
        base_docs=base_docs,
        guidance_docs=guidance_docs,
        selected_corpus_catalog=selected_corpus_catalog,
    )
    if inferred_context_info.get("used"):
        print("[4/6] Campos vazios de título/tema/recorte/objetivo inferidos por IA a partir do corpus e das orientações.", flush=True)

    documento_context, context_origin_info = maybe_rewrite_document_context(client, model, cfg, documento_context, guidance_docs)
    context_origin_info = {
        "missing_fields_inference": inferred_context_info,
        "rewrite_or_override": context_origin_info,
    }

    print(f"[5/6] Gerando ORG de {doc_label_with_article}...", flush=True)
    org_text, prompt_text = generate_paper_org(
        client=client,
        model=model,
        cfg=cfg,
        context=documento_context,
        template_text=template_text,
        base_docs=base_docs,
        guidance_docs=guidance_docs,
        bib_filename=documento_bib_name,
        bib_entries=bib_entries,
        bib_keys=bib_keys,
        style=style,
        selected_corpus_catalog=selected_corpus_catalog,
    )
    expansion_prompt_text = None
    expansion_diagnostics = None
    org_text, expansion_prompt_text, expansion_diagnostics = maybe_expand_dissertation_org(
        client=client,
        model=model,
        cfg=cfg,
        context=documento_context,
        org_text=org_text,
        base_docs=base_docs,
        guidance_docs=guidance_docs,
        bib_filename=documento_bib_name,
        bib_entries=bib_entries,
        bib_keys=bib_keys,
        style=style,
        selected_corpus_catalog=selected_corpus_catalog,
    )

    front = infer_final_front_matter(client, model, documento_context, org_text)
    atividade = cfg.get("atividade", {})
    local_cfg = cfg.get("documentos_locais", {}) if isinstance(cfg.get("documentos_locais"), dict) else {}
    configured_title = str(documento.get("titulo_trabalho") or documento.get("titulo") or atividade.get("titulo_trabalho") or atividade.get("titulo") or "").strip()
    context_title = str(documento_context.titulo_sugerido or "").strip()
    if normalize_document_type(documento.get("tipo_documento")) != "dissertacao" and configured_title and not is_empty_context_field(configured_title):
        final_title = configured_title
    elif context_title and not is_empty_context_field(context_title):
        final_title = context_title
    elif cfg.get("__local_corpus_prepared__") and not str(front.title or "").strip():
        final_title = pretty_title_from_prefix(documento_prefix)
    else:
        final_title = front.title.strip()
    org_text = apply_final_front_matter(
        org_text,
        title=final_title,
        author=str(atividade.get("aluno") or DEFAULT_AUTHOR),
        documento_type=front.documento_type.strip(),
        cover_note=front.cover_note.strip(),
        institution_name=str(documento.get("institution_name") or DEFAULT_INSTITUTION),
        course_name=str(atividade.get("curso") or ""),
        discipline_name=str(atividade.get("disciplina") or ""),
        professor_name=str(atividade.get("professor") or ""),
        city_name=str(atividade.get("polo") or "Brasília"),
    )
    if bib_alias_map:
        org_text = replace_org_citation_keys(org_text, bib_alias_map)

    if normalize_document_type(documento.get("tipo_documento")) != "dissertacao":
        org_text = force_activity_template_visible_title(org_text, final_title)
    if normalize_document_type(documento.get("tipo_documento")) == "dissertacao":
        org_text = apply_dissertation_template_metadata(org_text, cfg, final_title)
    org_text = render_template_placeholders(org_text, cfg, final_title)
    org_text = normalize_snippet_placeholders(org_text)
    org_text = apply_paper_template_metadata(org_text, cfg, final_title)

    mindmap_info: dict[str, Any] | None = None
    mindmap_prompt_text: str | None = None
    if should_generate_mindmap(cfg):
        print("[5/6] Gerando mapa mental PlantUML da atividade...", flush=True)
        org_text, mindmap_info, mindmap_prompt_text = maybe_generate_activity_mindmap(
            client=client,
            model=model,
            cfg=cfg,
            context=documento_context,
            org_text=org_text,
            base_docs=base_docs,
            guidance_docs=guidance_docs,
            selected_corpus_catalog=selected_corpus_catalog,
            documento_output_dir=documento_output_dir,
            documento_prefix=documento_prefix,
        )
        if mindmap_info and mindmap_info.get("generated"):
            if mindmap_info.get("rendered"):
                print(f"[5/6] Mapa mental gerado: {mindmap_info.get('image_path')}", flush=True)
            else:
                print(f"[5/6] Mapa mental PlantUML gerado: {mindmap_info.get('puml_path')}", flush=True)
                if mindmap_info.get("render_error"):
                    print(f"Aviso: {mindmap_info.get('render_error')}", flush=True)

    # Barreira final contra capa duplicada, vazamentos técnicos e citações quebradas no PDF.
    org_text = remove_technical_leaks_from_org(org_text, cfg)
    org_text = prepare_citations_for_pdf_export(org_text, cfg)
    org_text = polish_narrative_latex_citations(org_text, cfg)
    validate_final_org_or_raise(org_text, bib_keys, cfg)

    # grava artefatos do documento
    org_path = documento_output_dir / f"{documento_prefix}.org"
    bib_path = documento_output_dir / documento_bib_name
    context_json_path = documento_output_dir / f"{documento_prefix}_contexto.json"
    prompt_audit_path = documento_output_dir / f"{documento_prefix}_prompts_auditoria.txt"
    provenance_path = documento_output_dir / f"{documento_prefix}_proveniencia.json"

    write_text(org_path, org_text)
    write_text(bib_path, "\n\n".join(bib_entries).strip() + ("\n" if bib_entries else ""))
    write_text(context_json_path, json.dumps({
        "generated_at": datetime.now().isoformat(),
        "documento_context": asdict(documento_context),
        "research_paths": json_safe(asdict(research_paths)),
        "config_excerpt": {
            "atividade": cfg.get("atividade", {}),
            "pesquisa": cfg.get("pesquisa", {}),
            "documento": cfg.get("documento", {}),
        },
        "bundle_manifest": str(manifest_path) if manifest_path else None,
        "context_origin_info": context_origin_info,
        "local_bib_revision_info": local_bib_revision_info,
        "bib_dedup_report": bib_dedup_report,
        "mindmap_info": mindmap_info,
        "bib_keys": bib_keys,
        "base_docs": [asdict(d) for d in base_docs],
        "guidance_docs": [asdict(d) for d in guidance_docs],
        "selected_entries": research_paths.selected_entries,
        "selected_fulltext_paths": [str(p) for p in research_paths.selected_fulltext_paths],
        "selected_corpus_catalog": selected_corpus_catalog,
        "citation_target_source": cfg.get("__citation_target_source__"),
        "citation_target_keys": cfg.get("__citation_target_keys__", []),
        "unresolved_fulltext_pdfs": unresolved_fulltext_pdfs,
        "expansion_diagnostics": expansion_diagnostics,
        "extra_article_paths": [d.path for d in extra_docs],
    }, ensure_ascii=False, indent=2))
    prompt_audit_content = "===== generate_paper_org =====\n" + prompt_text + "\n"
    if expansion_prompt_text:
        prompt_audit_content += "\n===== maybe_expand_dissertation_org =====\n" + expansion_prompt_text + "\n"
        prompt_audit_content += "\n===== diagnostics =====\n" + json.dumps(expansion_diagnostics or {}, ensure_ascii=False, indent=2) + "\n"
    if mindmap_prompt_text:
        prompt_audit_content += "\n===== generate_activity_mindmap =====\n" + mindmap_prompt_text + "\n"
        prompt_audit_content += "\n===== mindmap_info =====\n" + json.dumps(mindmap_info or {}, ensure_ascii=False, indent=2) + "\n"
    write_text(prompt_audit_path, prompt_audit_content)
    write_text(provenance_path, json.dumps(
        build_provenance_payload(
            cfg=cfg,
            research_paths=research_paths,
            documento_context=documento_context,
            base_docs=base_docs,
            guidance_docs=guidance_docs,
            extra_docs=extra_docs,
            bib_keys=bib_keys,
            bib_entries=bib_entries,
            context_origin_info=context_origin_info,
            template_path=template_path,
            manifest_path=manifest_path,
        ),
        ensure_ascii=False,
        indent=2,
        default=str,
    ))

    reference_usage = build_reference_usage_map(org_text, base_docs, extra_docs, bib_keys, required_keys=cfg.get("__citation_target_keys__"))
    reference_usage_json_path = documento_output_dir / f"{documento_prefix}_uso_referencias.json"
    reference_usage_md_path = documento_output_dir / f"{documento_prefix}_uso_referencias.md"
    write_text(reference_usage_json_path, json.dumps(reference_usage, ensure_ascii=False, indent=2))
    write_text(reference_usage_md_path, render_reference_usage_markdown(reference_usage))

    section_limits_json_path = documento_output_dir / f"{documento_prefix}_limites_secoes.json"
    write_text(section_limits_json_path, json.dumps({
        "word_counts": count_org_words_per_top_section(org_text),
        "configured_limits": {
            "total": documento.get("limite_palavras_total"),
            "introducao": documento.get("limite_palavras_introducao"),
            "revisao": documento.get("limite_palavras_revisao"),
            "conclusao": documento.get("limite_palavras_conclusao"),
        }
    }, ensure_ascii=False, indent=2))

    # compila PDF do documento se pedido
    documento_pdf_path = None
    if bool(documento.get("exportar_pdf", cfg.get("saida", {}).get("exportar_pdf", False))):
        print(f"[6/6] Compilando PDF de {doc_label_with_article}...", flush=True)
        latex = cfg.get("latex", {})
        academic_writing = resolve_configured_path(latex.get("org_latex_class_init"), cfg) if latex.get("org_latex_class_init") else None
        emacs_init = resolve_configured_path(latex.get("emacs_init"), cfg) if latex.get("emacs_init") else None
        latex_extra_path = resolve_configured_path(latex.get("latex_extra_path"), cfg) if latex.get("latex_extra_path") else None
        try:
            documento_pdf_path = run_compile_sequence(org_path, emacs_init=emacs_init, academic_writing=academic_writing, latex_extra_path=latex_extra_path)
        except Exception as exc:
            pdf_error_path = documento_output_dir / f"{documento_prefix}_pdf_erro.txt"
            write_text(
                pdf_error_path,
                f"A compilação do PDF de {doc_label_with_article} falhou, mas o .org foi preservado.\n\n"
                f"Documento ORG: {org_path}\n"
                f"Documento BIB: {bib_path}\n\n"
                f"Erro:\n{exc}\n",
            )
            print(f"Aviso: a compilação do PDF de {doc_label_with_article} falhou, mas os artefatos textuais foram preservados.")
            print(f"- Documento ORG preservado em: {org_path}")
            print(f"- Log simplificado da falha do PDF: {pdf_error_path}")
            documento_pdf_path = None

    print("\nArquivos gerados:")
    local_mode = is_local_documents_mode(cfg)
    source_prefix = "Corpus local" if local_mode else "Pesquisa"
    used_docs_label = "Documentos locais usados no documento" if local_mode else "PDFs selecionados usados no documento"

    print(f"- {source_prefix} ORG: {research_paths.org_path}")
    if research_paths.bib_path:
        print(f"- {source_prefix} BIB: {research_paths.bib_path}")
    if research_paths.debug_path:
        print(f"- {source_prefix} DEBUG: {research_paths.debug_path}")
    if research_paths.pdf_path:
        print(f"- {source_prefix} PDF: {research_paths.pdf_path}")
    if research_paths.fulltext_cache_dir:
        print(f"- {source_prefix} FULLTEXT CACHE: {research_paths.fulltext_cache_dir}")
    if research_paths.selected_fulltext_paths:
        print(f"- {used_docs_label}: {len(research_paths.selected_fulltext_paths)}")
    if extra_docs:
        print(f"- Artigos extras usados no documento: {len(extra_docs)}")
    print(f"- Documento ORG: {org_path}")
    print(f"- Documento BIB: {bib_path}")
    print(f"- Documento CONTEXTO: {context_json_path}")
    print(f"- Documento PROVENIÊNCIA: {provenance_path}")
    print(f"- Documento PROMPTS: {prompt_audit_path}")
    print(f"- Diretório do documento: {documento_output_dir}")
    print(f"- Uso de referências (JSON): {reference_usage_json_path}")
    print(f"- Uso de referências (MD): {reference_usage_md_path}")
    print(f"- Contagem/limites por seção: {section_limits_json_path}")
    if mindmap_info and mindmap_info.get("generated"):
        print(f"- Mapa mental PlantUML: {mindmap_info.get('puml_path')}")
        if mindmap_info.get("image_path"):
            print(f"- Mapa mental imagem: {mindmap_info.get('image_path')}")
    if manifest_path:
        update_bundle_with_documento(
            manifest_path,
            org_path=org_path,
            bib_path=bib_path,
            context_json_path=context_json_path,
            prompt_audit_path=prompt_audit_path,
            provenance_path=provenance_path,
            extra_docs=extra_docs,
            pdf_path=documento_pdf_path,
        )

    package_dir = assemble_delivery_package(
        cfg,
        research_paths=research_paths,
        documento_output_dir=documento_output_dir,
        documento_prefix=documento_prefix,
        manifest_path=manifest_path,
        org_path=org_path,
        bib_path=bib_path,
        context_json_path=context_json_path,
        prompt_audit_path=prompt_audit_path,
        provenance_path=provenance_path,
        reference_usage_json_path=reference_usage_json_path,
        reference_usage_md_path=reference_usage_md_path,
        section_limits_json_path=section_limits_json_path,
        documento_pdf_path=documento_pdf_path,
    )

    req_total = int(reference_usage.get("counts", {}).get("required_fulltext_cache_total", 0))
    req_cited = int(reference_usage.get("counts", {}).get("required_fulltext_cache_cited_total", 0))
    if req_total and req_cited < req_total:
        print(f"Aviso: o documento citou {req_cited}/{req_total} chave(s) dos PDFs presentes no fulltext_cache. Verifique o mapa de uso de referências para as chaves faltantes.")

    if documento_pdf_path:
        print(f"- Documento PDF: {documento_pdf_path}")
    if manifest_path:
        print(f"- Bundle MANIFEST: {manifest_path}")
    if package_dir:
        print(f"- Pacote final de entrega: {package_dir}")

    return 0

if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\nOperação cancelada pelo usuário.", file=sys.stderr)
        raise SystemExit(130)
    except Exception as exc:
        print(f"Erro: {exc}", file=sys.stderr)
        if DEBUG:
            traceback.print_exc()
        raise SystemExit(1)
